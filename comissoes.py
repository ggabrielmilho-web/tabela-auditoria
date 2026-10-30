# -*- coding: utf-8 -*-
"""Comissão de frete: a planilha do responsável conferida contra o SSW (30/10/2026).

Quem contrata o frete de terceiro (carreteiro) ganha comissão sobre o resultado de cada viagem. O
controle era uma planilha por responsável e mês, uma linha por CTRB (Cliente, Responsável, Data,
Origem, Destino, FRETE, ICMS, Frete s/ ICMS, Frete Motorista, Carga, Descarga, Pedágio, Outros
Custos, Total, Resultado, Margem, Comissão, CTRB, Motorista...). O diretor anexa a planilha e a tela
refaz cada linha pelo SSW, aponta onde ela diverge e guarda o fechamento aprovado.

A conta (regras do diretor, 30/10/2026):

- **Receita** = frete do CT-e sem ICMS e SEM o campo "Outros". O "Outros" é a descarga que a Rizza
  cobra do cliente e paga ao próprio carreteiro pelo evento 5100 — repasse, não frete vendido. O
  `receita_rateada` do Auditoria Receita soma a descarga; aqui ela sai na mesma proporção do rateio.
- **Custo** = frete do motorista + vale-pedágio SÓ quando a Rizza pagou o vale (há lançamento 5403 do
  CTRB no 477). Vale-pedágio de TAG do cliente aparece destacado no CTRB, mas não sai do caixa.
- **Comissão** = PERCENTUAL do resultado; negativa desconta.
- **Mês** = o da emissão do CTRB. CTRB de outro mês fica fora por padrão (alerta).

Casos que a conferência de out/26 mostrou e a tela trata:

- **CTRB refeito**: o da planilha não tem lançamento no 477 e existe outro do mesmo motorista, emitido
  até 3 dias depois, com o mesmo CT-e ou com a mesma rota e o mesmo valor. A conta usa o substituto.
- **Reemissão para o mesmo CT-e** (dois CTRBs do mesmo motorista no mesmo CT-e): o Auditoria Receita
  divide a receita entre os dois; a conta usa o CT-e inteiro.
- **Linha sem CTRB**: sugere o CTRB do mesmo motorista com data próxima (prefere o mesmo valor).
- **Fora da planilha**: CTRB do mês cuja observação diz "FRETE COMBINADO COM <responsável>" e que a
  planilha não traz — o diretor pode incluir.

Este módulo não fala Flask nem abre conexão: recebe a função de consulta ao BI e o cursor.
"""

import io
import json
import re
import threading
import time
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone

BRT = timezone(timedelta(hours=-3))

PERCENTUAL = 0.03      # fixo por enquanto (diretor, 30/10/2026)
TOLERANCIA = 1.00      # diferença de comissão abaixo disto conta como "confere"
DIAS_SUBSTITUTO = 3    # o CTRB refeito sai até 3 dias depois do original
DIAS_SEM_CTRB = 5      # janela para sugerir o CTRB de uma linha que veio sem número
MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
MESES_LONGO = ['janeiro', 'fevereiro', 'março', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro',
               'outubro', 'novembro', 'dezembro']


def hoje_brt():
    return datetime.now(BRT).date()


def rotulo_mes(mes, longo=False):
    """'2026-10' → 'out/26' (ou 'outubro/2026')."""
    try:
        a, m = int(mes[:4]), int(mes[5:7])
    except (TypeError, ValueError):
        return mes or ''
    return f'{MESES_LONGO[m - 1]}/{a}' if longo else f'{MESES[m - 1]}/{str(a)[2:]}'


# ── Normalização ─────────────────────────────────────────────────────────────

def _sem_acento(s):
    return ''.join(c for c in unicodedata.normalize('NFD', str(s or '')) if unicodedata.category(c) != 'Mn')


def nome_norm(s):
    """'José ' → 'JOSE'; 'JOAO DA  SILVA' → 'JOAO DA SILVA'."""
    return re.sub(r'\s+', ' ', _sem_acento(s).upper()).strip()


def _lev(a, b):
    if a == b:
        return 0
    ant = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        atual = [i]
        for j, cb in enumerate(b, 1):
            atual.append(min(ant[j] + 1, atual[j - 1] + 1, ant[j - 1] + (ca != cb)))
        ant = atual
    return ant[-1]


def mesmo_nome(a, b):
    """O nome digitado na observação do CTRB erra letra ("MARCLO" por "MARCELO"): mesmo começo e até 2 letras."""
    a, b = nome_norm(a), nome_norm(b)
    if not a or not b:
        return False
    return a == b or (a[:3] == b[:3] and _lev(a, b) <= 2)


def chave_ctrb(v):
    """O CTRB como o SSW grava ('UDI000123-4') a partir do que foi digitado: 'UDI-000123-4',
    'UDI000123-4', 'UDI-UDI 000123-4', 'UDI 123-4'. None se não parecer CTRB."""
    s = nome_norm(v)
    m = re.search(r'([A-Z]{3})[\s\-.]*(\d{1,6})\s*-\s*(\d)\s*$', s) or re.search(r'([A-Z]{3})[\s\-.]*(\d{6})(\d)\s*$', s)
    if not m:
        return None
    return f'{m.group(1)}{int(m.group(2)):06d}-{m.group(3)}'


_RE_COMBINADO = re.compile(r'COMBI[A-Z]*\s+(?:C\s*/\s*|COM\s+|C\.\s*)([A-Z]{3,})')


def responsavel_da_obs(obs):
    """'FRETE COMBINADO C/MARCELO R$ 5.300,00…' → 'MARCELO' (None se a observação não diz)."""
    m = _RE_COMBINADO.search(nome_norm(obs))
    return m.group(1) if m else None


def _num(v):
    if v is None or v == '':
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace('R$', '').replace(' ', '').strip()
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0


def _txt(v):
    s = re.sub(r'\s+', ' ', str(v if v is not None else '').replace('\xa0', ' ')).strip()
    return s or None


def _iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or '')
    if re.match(r'^\d{4}-\d{2}-\d{2}', s):
        return s[:10]
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{2,4})', s)
    if m:
        a = int(m.group(3))
        return f'{a + 2000 if a < 100 else a:04d}-{int(m.group(2)):02d}-{int(m.group(1)):02d}'
    return None


def _dias(a, b):
    try:
        return abs((date.fromisoformat(a[:10]) - date.fromisoformat(b[:10])).days)
    except (TypeError, ValueError):
        return 9999


# ── A planilha ───────────────────────────────────────────────────────────────

# cabeçalho normalizado → campo. A ordem importa: 'frete s/ icms' e 'frete motorista' antes de 'frete'.
_CABECALHO = [
    ('cliente', 'cliente'), ('respons', 'responsavel'), ('status', 'status'), ('veiculo', 'veiculo'),
    ('data', 'data'), ('origem', 'origem'), ('destino', 'destino'),
    ('frete s', 'frete_s_icms'), ('frete motorista', 'frete_motorista'), ('icms', 'icms'), ('frete', 'frete'),
    ('carga', 'carga'), ('descarga', 'descarga'), ('pedagio', 'pedagio'), ('outros', 'outros'), ('total', 'total'),
    ('resultado', 'resultado'), ('margem', 'margem'), ('comiss', 'comissao'), ('ctrb', 'ctrb'),
    ('motorista', 'motorista'), ('canhoto', 'canhoto'), ('pago', 'pago'),
]
_NUMERICOS = ('frete', 'icms', 'frete_s_icms', 'frete_motorista', 'carga', 'descarga', 'pedagio', 'outros',
              'total', 'resultado', 'comissao')


def _campo(cab):
    c = nome_norm(cab).lower()
    if not c:
        return None
    if c.startswith('descarga'):
        return 'descarga'
    for prefixo, campo in _CABECALHO:
        if c.startswith(prefixo):
            return campo
    return None


def ler_planilha(arq):
    """xlsx → {'linhas': [...], 'aba': nome}. Acha a linha de cabeçalho (a que tem CTRB e Comissão) e
    lê até o fim; a linha de total (sem responsável, sem CTRB e sem frete) fica de fora."""
    import openpyxl
    wb = openpyxl.load_workbook(arq, data_only=True, read_only=True)
    for ws in wb.worksheets:
        cols, linhas = None, []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if cols is None:
                mapa = {}
                for j, v in enumerate(row):
                    campo = _campo(v)
                    if campo and campo not in mapa.values():
                        mapa[j] = campo
                if 'ctrb' in mapa.values() and 'comissao' in mapa.values():
                    cols = mapa
                continue
            r = {campo: (row[j] if j < len(row) else None) for j, campo in cols.items()}
            if not (r.get('responsavel') or r.get('ctrb') or r.get('frete')):
                continue
            item = {'lin': i}
            for campo, v in r.items():
                if campo in _NUMERICOS:
                    item[campo] = round(_num(v), 2)
                elif campo == 'data':
                    item[campo] = _iso(v)
                elif campo == 'margem':
                    item[campo] = _num(v)
                else:
                    item[campo] = _txt(v)
            if not item.get('frete_s_icms') and item.get('frete'):
                item['frete_s_icms'] = round(item['frete'] - item.get('icms', 0), 2)
            linhas.append(item)
        if cols is not None and linhas:
            return {'linhas': linhas, 'aba': ws.title}
    raise ValueError('Não achei o cabeçalho com as colunas "CTRB" e "Comissão" — é a planilha de comissão?')


def resumo_planilha(linhas):
    """Responsáveis (com o nome como foi escrito e quantas linhas) e o mês mais frequente da coluna Data."""
    por = defaultdict(lambda: {'linhas': 0, 'comissao': 0.0, 'nomes': defaultdict(int)})
    meses = defaultdict(int)
    for l in linhas:
        n = l.get('responsavel') or ''
        k = nome_norm(n) or '(SEM RESPONSÁVEL)'
        por[k]['linhas'] += 1
        por[k]['comissao'] += l.get('comissao') or 0
        por[k]['nomes'][n or '(sem responsável)'] += 1
        if l.get('data'):
            meses[l['data'][:7]] += 1
    resp = [{'chave': k, 'nome': max(v['nomes'].items(), key=lambda x: x[1])[0], 'linhas': v['linhas'],
             'comissao': round(v['comissao'], 2)} for k, v in por.items()]
    resp.sort(key=lambda r: -r['linhas'])
    mes = max(meses.items(), key=lambda x: x[1])[0] if meses else hoje_brt().isoformat()[:7]
    return {'responsaveis': resp, 'mes': mes, 'meses': dict(meses)}


# ── O SSW, pelo Power BI ─────────────────────────────────────────────────────

_OSS = "'public ctrbs_oss'"
_AR = "'Auditoria Receita'"
_CE = "'public conhecimentos_emitidos'"
_DZ = "'public consulta_despesas_477'"


def _d(dt):
    return f'DATE({dt.year},{dt.month},{dt.day})'


def _lista(vals):
    return '{' + ','.join('"' + str(v).replace('"', '') + '"' for v in vals) + '}'


def _fatias(vals, n=250):
    vals = sorted(set(v for v in vals if v))
    return [vals[i:i + n] for i in range(0, len(vals), n)]


def janela(mes):
    """Do 1º dia de dois meses antes ao dia 20 do mês seguinte: cobre os CTRBs de fim de mês que a
    planilha traz, os substitutos e os pagamentos."""
    a, m = int(mes[:4]), int(mes[5:7])
    ini = date(a + (m - 3) // 12, (m - 3) % 12 + 1, 1)
    fim = date(a + m // 12, m % 12 + 1, 20)
    return ini, fim


_COLS_OSS = (f'"ctrb", {_OSS}[ctrb], "emissao", {_OSS}[emissao], "motorista", {_OSS}[motorista], '
             f'"origem", {_OSS}[cidade_uf_origem], "destino", {_OSS}[cidade_uf_destino], '
             f'"vap", {_OSS}[valor_a_pagar], "vped", {_OSS}[vale_pedagio], "dvped", {_OSS}[despesa_vale_pedagio], '
             f'"manifesto", {_OSS}[manifesto], "obs", {_OSS}[observacao]')
_COLS_AR = (f'"ctrb", {_AR}[CTRB], "tipo", {_AR}[Tipo Operacao], "data", {_AR}[data_ref_ctrc], "ctrc", {_AR}[CTRC], '
            f'"rec", {_AR}[receita_rateada], "vfsi", {_AR}[valor_frete_sem_icms], "fmt", {_AR}[frete_motorista_total], '
            f'"vped", {_AR}[vale_pedagio], "motorista", {_AR}[motorista], "cliente", {_AR}[cliente_pagador], '
            f'"origem", {_AR}[cidade_uf_origem], "destino", {_AR}[cidade_uf_destino]')
_COLS_CE = (f'"ctrc", {_CE}[serie_numero_ctrc], "vf", {_CE}[valor_frete], "vfsi", {_CE}[valor_frete_sem_icms], '
            f'"icms", {_CE}[valor_icms], "outros", {_CE}[outros], "cliente", {_CE}[cliente_pagador]')
_COLS_DZ = (f'"uni", {_DZ}[uni], "num", {_DZ}[numlancto], "parcela", {_DZ}[parcela], "evento", {_DZ}[evento], '
            f'"valor", {_DZ}[vlr_final], "sit", {_DZ}[sit_des], "pgto", {_DZ}[data_pgto], "hist", {_DZ}[historico_despesa]')

_RE_CTRB_HIST = re.compile(r'CTRB\s+([A-Z]{3})\s*(\d{1,6})-(\d)')


def buscar_ssw(dax, mes, ctrbs_planilha=()):
    """Tudo o que a conferência de um mês precisa, em ~6 consultas: os CTRBs (ctrbs_oss), o Auditoria
    Receita, os CT-e e os lançamentos do 477 que citam CTRB. `dax` levanta erro se o BI cortar."""
    ini, fim = janela(mes)
    oss = dax(f'EVALUATE SELECTCOLUMNS(FILTER({_OSS}, {_OSS}[emissao] >= {_d(ini)} && {_OSS}[emissao] < {_d(fim)}), {_COLS_OSS})')
    vistos = {chave_ctrb(r.get('ctrb')) for r in oss}
    for f in _fatias([k for k in ctrbs_planilha if k not in vistos]):
        oss += dax(f'EVALUATE SELECTCOLUMNS(FILTER({_OSS}, {_OSS}[ctrb] IN {_lista(f)}), {_COLS_OSS})')
    ar = dax(f'EVALUATE SELECTCOLUMNS(FILTER({_AR}, {_AR}[data_ref_ctrc] >= {_d(ini)} && {_AR}[data_ref_ctrc] < {_d(fim)}), {_COLS_AR})')
    vistos = {chave_ctrb(r.get('ctrb')) for r in ar}
    for f in _fatias([k for k in ctrbs_planilha if k not in vistos]):
        ar += dax(f'EVALUATE SELECTCOLUMNS(FILTER({_AR}, {_AR}[CTRB] IN {_lista(f)}), {_COLS_AR})')
    ctrcs = {c.strip() for r in ar for c in str(r.get('ctrc') or '').split(',') if c.strip()}
    ce = []
    for f in _fatias(ctrcs):
        ce += dax(f'EVALUATE SELECTCOLUMNS(FILTER({_CE}, {_CE}[serie_numero_ctrc] IN {_lista(f)}), {_COLS_CE})')
    dz = dax(f'EVALUATE SELECTCOLUMNS(FILTER({_DZ}, {_DZ}[inclusao] >= {_d(ini - timedelta(days=15))} && '
             f'SEARCH("CTRB", {_DZ}[historico_despesa], 1, 0) > 0), {_COLS_DZ})')
    return {'oss': oss, 'ar': ar, 'ce': ce, 'dz': dz, 'mes': mes,
            'lido_em': datetime.now(BRT).strftime('%d/%m/%Y %H:%M')}


_CACHE = {}
_CACHE_TRAVA = threading.Lock()
CACHE_S = 600


def dados_ssw(dax, mes, ctrbs_planilha=(), forcar=False):
    """`buscar_ssw` com cache de 10 min por mês (o BI recarrega ~8×/dia)."""
    chave = (mes, tuple(sorted(set(k for k in ctrbs_planilha if k))))
    with _CACHE_TRAVA:
        v = _CACHE.get(chave)
        if not forcar and v and time.time() - v[0] < CACHE_S:
            return v[1]
    dados = buscar_ssw(dax, mes, ctrbs_planilha)
    with _CACHE_TRAVA:
        _CACHE.clear() if len(_CACHE) > 20 else None
        _CACHE[chave] = (time.time(), dados)
    return dados


# ── A conferência ────────────────────────────────────────────────────────────

class _Base:
    """Os dados do SSW indexados para a conferência."""

    def __init__(self, ssw):
        self.oss = {}
        for r in ssw['oss']:
            k = chave_ctrb(r.get('ctrb'))
            if k:
                self.oss[k] = r
        self.ar = {}
        self.ar_por_ctrc = defaultdict(list)
        for r in ssw['ar']:
            k = chave_ctrb(r.get('ctrb'))
            if not k:
                continue
            r = dict(r, _ctrcs=[c.strip() for c in str(r.get('ctrc') or '').split(',') if c.strip()])
            self.ar[k] = r
            for c in r['_ctrcs']:
                self.ar_por_ctrc[c].append(k)
        self.ce = {str(r.get('ctrc') or '').strip(): r for r in ssw['ce']}
        self.pag = defaultdict(list)
        for r in ssw['dz']:
            m = _RE_CTRB_HIST.search(nome_norm(r.get('hist')))
            if m:
                self.pag[f'{m.group(1)}{int(m.group(2)):06d}-{m.group(3)}'].append(r)

    def emissao(self, k):
        r = self.oss.get(k) or {}
        return _iso(r.get('emissao')) or _iso((self.ar.get(k) or {}).get('data'))

    def motorista(self, k):
        return nome_norm((self.oss.get(k) or {}).get('motorista') or (self.ar.get(k) or {}).get('motorista'))

    def pagamento(self, k):
        """Situação do pagamento do frete (sem o vale-pedágio): sem lançamento, pendente, parcial ou pago."""
        ps = [p for p in self.pag.get(k, []) if str(p.get('evento')) != '5403']
        parc = [{'parcela': p.get('parcela'), 'valor': round(_num(p.get('valor')), 2), 'sit': p.get('sit'),
                 'data': _iso(p.get('pgto')), 'evento': p.get('evento')} for p in ps]
        parc.sort(key=lambda p: str(p['parcela']))
        if not ps:
            sit = 'sem'
        elif all(p.get('sit') == 'LIQU' for p in ps):
            sit = 'pago'
        elif any(p.get('sit') == 'LIQU' for p in ps):
            sit = 'parcial'
        else:
            sit = 'pendente'
        aberto = round(sum(p['valor'] for p in parc if p['sit'] != 'LIQU'), 2)
        return {'sit': sit, 'parcelas': parc, 'aberto': aberto}

    def vped_pago(self, k):
        return any(str(p.get('evento')) == '5403' for p in self.pag.get(k, []))

    def irmaos(self, k):
        """Outros CTRBs do MESMO motorista no mesmo CT-e (reemissão)."""
        a = self.ar.get(k)
        if not a:
            return []
        mot = self.motorista(k)
        out = set()
        for c in a['_ctrcs']:
            for o in self.ar_por_ctrc.get(c, []):
                if o != k and self.motorista(o) == mot:
                    out.add(o)
        return sorted(out)


def conta(base, k):
    """A comissão de um CTRB pela régua do diretor. None se o CTRB não tem CT-e no Auditoria Receita."""
    a = base.ar.get(k)
    if not a or a.get('rec') is None:
        return None
    irmaos = base.irmaos(k)
    mot = base.motorista(k)
    # reemissão: o BI divide a receita entre os CTRBs do mesmo motorista; a viagem é uma só
    receita_bi = _num(a.get('vfsi')) if irmaos else _num(a.get('rec'))
    # Quanto de cada CT-e o BI pôs neste CTRB: inteiro se o CT-e só está aqui (ou na reemissão); o que sobra
    # da receita rateada se divide entre os CT-e que também estão em CTRB de outro motorista. A descarga e o
    # ICMS saem CT-e por CT-e nessa parte — não pela proporção da viagem toda.
    ctes = [(c, base.ce.get(c) or {}) for c in a['_ctrcs']]
    divididos = {c for c, _ in ctes if not irmaos and any(o != k and base.motorista(o) != mot
                                                          for o in base.ar_por_ctrc.get(c, []))}
    inteiro = sum(_num(r.get('vfsi')) for c, r in ctes if c not in divididos)
    vfsi_div = sum(_num(r.get('vfsi')) for c, r in ctes if c in divididos)
    parte_div = min(1.0, max(0.0, (receita_bi - inteiro) / vfsi_div)) if vfsi_div else 0.0
    parte = {c: (parte_div if c in divididos else 1.0) for c, _ in ctes}
    outros = round(sum(_num(r.get('outros')) * parte[c] for c, r in ctes), 2)
    icms = round(sum(_num(r.get('icms')) * parte[c] for c, r in ctes), 2)
    receita = round(receita_bi - outros, 2)
    de_outros = []
    for c, r in ctes:
        if c in divididos:
            donos = [o for o in base.ar_por_ctrc.get(c, []) if o != k and base.motorista(o) != mot]
            de_outros.append({'cte': c, 'cliente': _txt(r.get('cliente')), 'aqui': round(_num(r.get('vfsi')) * parte[c], 2),
                              'total': round(_num(r.get('vfsi')), 2),
                              'ctrbs': [{'ctrb': o, 'motorista': base.motorista(o)} for o in donos]})
    vped = round(_num(a.get('vped')), 2)
    pago = base.vped_pago(k)
    frete_mot = round(_num(a.get('fmt')) - vped, 2)
    custo = round(frete_mot + (vped if pago else 0.0), 2)
    resultado = round(receita - custo, 2)
    return {'ctrb': k, 'receita': receita, 'descarga': outros, 'icms': icms,
            'frete_motorista': frete_mot, 'vale_pedagio': vped, 'vale_pago_rizza': pago, 'custo': custo,
            'resultado': resultado, 'comissao': round(resultado * PERCENTUAL, 2),
            'ctes': a['_ctrcs'], 'irmaos': irmaos, 'rateado': bool(de_outros), 'de_outros': de_outros,
            'cliente': a.get('cliente'), 'tipo': a.get('tipo')}


def _substituto(base, k):
    """O CTRB que refez `k`: mesmo motorista, emitido até DIAS_SUBSTITUTO depois, com o mesmo CT-e ou
    com a mesma rota e o mesmo valor, e que tenha lançamento no 477."""
    o = base.oss.get(k)
    if not o:
        return None
    mot, em = base.motorista(k), base.emissao(k)
    irmaos = set(base.irmaos(k))
    cands = []
    for c, r in base.oss.items():
        if c == k or base.motorista(c) != mot or not base.pag.get(c):
            continue
        ec = base.emissao(c)
        if not ec or not em or not (0 <= (date.fromisoformat(ec) - date.fromisoformat(em)).days <= DIAS_SUBSTITUTO):
            continue
        mesma_rota = (nome_norm(r.get('origem')) == nome_norm(o.get('origem'))
                      and nome_norm(r.get('destino')) == nome_norm(o.get('destino'))
                      and abs(_num(r.get('vap')) - _num(o.get('vap'))) < 0.01)
        if c in irmaos or mesma_rota:
            cands.append((0 if c in irmaos else 1, _dias(ec, em), c))
    return min(cands)[2] if cands else None


def _candidato_sem_ctrb(base, l, usados):
    """Para a linha sem número: o CTRB do mesmo motorista com data próxima (prefere o mesmo valor)."""
    mot, data = nome_norm(l.get('motorista')), l.get('data')
    if not mot or not data:
        return None
    cands = []
    for c, r in base.oss.items():
        if c in usados or base.motorista(c) != mot:
            continue
        d = _dias(base.emissao(c) or '', data)
        if d > DIAS_SEM_CTRB:
            continue
        cands.append((0 if abs(_num(r.get('vap')) - (l.get('frete_motorista') or 0)) < 0.01 else 1, d, c))
    return min(cands)[2] if cands else None


def _brl(v):
    s = f'{abs(v):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return ('−' if v < 0 else '') + 'R$ ' + s


def _motivos(l, c):
    """Por que a linha diverge, em português de gente."""
    out = []
    rec_pl = l.get('frete_s_icms') or 0
    if abs(rec_pl - c['receita']) > 1:
        if c['icms'] > 0 and abs(rec_pl - (c['receita'] + c['icms'])) <= 1:
            out.append(f"Frete lançado com ICMS: {_brl(rec_pl)}; sem ICMS é {_brl(c['receita'])}.")
        else:
            out.append(f"Frete na planilha {_brl(rec_pl)}; pelo SSW é {_brl(c['receita'])}.")
            for d in c.get('de_outros') or []:
                quem = ', '.join(f"{x['ctrb']} ({x['motorista'].title()})" for x in d['ctrbs'])
                out.append(f"O CT-e {d['cte']}{' de ' + d['cliente'] if d['cliente'] else ''} ({_brl(d['total'])}) também "
                           f"está no CTRB {quem}, de outro motorista; o BI pôs {_brl(d['aqui'])} dele nesta viagem. "
                           f"Se a carga não era deste caminhão, vale a planilha.")
    custo_pl = l.get('total') or 0
    if abs(custo_pl - c['custo']) > 1:
        outros_pl = (l.get('carga') or 0) + (l.get('descarga') or 0) + (l.get('outros') or 0)
        ped_pl = l.get('pedagio') or 0
        if c['vale_pago_rizza'] and c['vale_pedagio'] > 0 and abs(custo_pl + c['vale_pedagio'] - c['custo']) <= 1:
            out.append(f"Vale-pedágio de {_brl(c['vale_pedagio'])} foi pago pela Rizza e não foi descontado.")
        elif ped_pl > 0 and not c['vale_pago_rizza'] and abs(custo_pl - ped_pl - c['custo']) <= 1:
            out.append(f"A planilha desconta {_brl(ped_pl)} de pedágio, mas o vale é TAG do cliente (a Rizza não pagou).")
        else:
            fm_pl = l.get('frete_motorista') or 0
            if abs(fm_pl - c['frete_motorista']) > 1:
                out.append(f"Frete motorista na planilha {_brl(fm_pl)}; o CTRB é {_brl(c['frete_motorista'])}.")
            if outros_pl > 0:
                out.append(f"A planilha desconta {_brl(outros_pl)} de carga/descarga/outros que o CTRB não tem.")
            if c['vale_pago_rizza'] and c['vale_pedagio'] > 0 and abs(ped_pl - c['vale_pedagio']) > 1:
                out.append(f"Vale-pedágio pago pela Rizza: {_brl(c['vale_pedagio'])} (planilha: {_brl(ped_pl)}).")
    if not out:
        out.append(f"Resultado na planilha {_brl(l.get('resultado') or 0)}; conferido {_brl(c['resultado'])}.")
    return out


def conferir(linhas, responsavel, mes, ssw, comissionados=None):
    """As linhas da planilha de UM responsável conferidas contra o SSW.

    `comissionados` = {CTRB: 'out/26 · Fulano'} das comissões já aprovadas (outros fechamentos).
    Cada linha volta com `planilha`, `conferido` (ou None), `status` (confere/diverge/sem_base),
    `motivos`, `alertas`, `pagamento` e `padrao` (conferida/planilha/fora) — a escolha que vale se o
    diretor não mexer. `fora` = CTRBs do responsável no mês que a planilha não traz."""
    base = _Base(ssw)
    comissionados = comissionados or {}
    resp = nome_norm(responsavel)
    k_planilha = {chave_ctrb(l.get('ctrb')) for l in linhas} - {None}
    usados = set(k_planilha)
    saida = []
    for l in linhas:
        k = chave_ctrb(l.get('ctrb'))
        alertas, usado = [], k
        if not k:
            cand = _candidato_sem_ctrb(base, l, usados)
            alertas.append({'tipo': 'sem_ctrb', 'texto': 'Sem CTRB na planilha',
                            'detalhe': (f'Provável {cand}: mesmo motorista, emitido em {_ddmm(base.emissao(cand))}, '
                                        f'{_brl(_num(base.oss[cand].get("vap")))}.'
                                        if cand else 'Não achei CTRB do mesmo motorista com data próxima.'),
                            'sugestao': cand})
            usado = cand
        elif k not in base.oss and k not in base.ar:
            alertas.append({'tipo': 'nao_achado', 'texto': 'CTRB não existe no SSW',
                            'detalhe': f'"{l.get("ctrb")}" não aparece no SSW — provável erro de digitação.'})
            usado = None
        elif not base.pag.get(k):
            sub = _substituto(base, k)
            if sub:
                alertas.append({'tipo': 'substituido', 'texto': f'CTRB substituído por {sub}',
                                'detalhe': f'{k} não tem lançamento no 477; {sub} é do mesmo motorista, emitido em '
                                           f'{_ddmm(base.emissao(sub))}, e foi lançado. A conta usa o {sub}.'})
                usado = sub
            else:
                alertas.append({'tipo': 'pendente', 'texto': 'CTRB pendente',
                                'detalhe': f'{k} ainda não tem lançamento de pagamento no 477.'})
        if usado:
            usados.add(usado)
        c = conta(base, usado) if usado else None
        if k and k in base.oss and len(re.findall(r'[A-Z]{3}', nome_norm(l.get('ctrb')))) > 1:
            alertas.append({'tipo': 'digitacao', 'texto': 'Nº do CTRB digitado errado',
                            'detalhe': f'"{l["ctrb"]}" foi lido como {k}.'})
        em = base.emissao(usado) if usado else None
        if em and em[:7] != mes:
            alertas.append({'tipo': 'outro_mes', 'texto': f'CTRB de {rotulo_mes(em[:7])}',
                            'detalhe': f'Emitido em {_ddmm(em)}. Pela data do CTRB, é da comissão de '
                                       f'{rotulo_mes(em[:7], longo=True)} — fica fora deste mês, a não ser que '
                                       f'você inclua.'})
        ja = comissionados.get(usado) or (comissionados.get(k) if k else None)
        if ja:
            alertas.append({'tipo': 'ja_comissionado', 'texto': 'Já comissionado',
                            'detalhe': f'Este CTRB já está na comissão aprovada de {ja}.'})
        if c is None:
            status, motivos = 'sem_base', []
            if usado and not any(a['tipo'] in ('sem_ctrb', 'nao_achado') for a in alertas):
                alertas.append({'tipo': 'sem_receita', 'texto': 'Sem CT-e no BI',
                                'detalhe': f'{usado} não tem CT-e ligado no Auditoria Receita (diária, ou CT-e '
                                           f'ainda não emitido): vale o valor da planilha.'})
        elif abs((l.get('comissao') or 0) - c['comissao']) < TOLERANCIA:
            status, motivos = 'confere', []
        else:
            status, motivos = 'diverge', _motivos(l, c)
        if any(a['tipo'] in ('outro_mes', 'ja_comissionado') for a in alertas):
            padrao = 'fora'
        elif c is None:
            padrao = 'planilha'
        elif status == 'diverge' and c['rateado'] and abs((l.get('frete_s_icms') or 0) - c['receita']) > 1:
            # CT-e dividido com CTRB de outro motorista: o rateio do BI não diz qual parte é desta viagem
            padrao = 'planilha'
        else:
            padrao = 'conferida'
        o = base.oss.get(usado) or {}
        saida.append({
            'chave': f"L{l['lin']}", 'lin': l['lin'], 'planilha': l, 'ctrb': k, 'ctrb_usado': usado,
            'emissao': em, 'conferido': c, 'status': status, 'motivos': motivos, 'alertas': alertas,
            'pagamento': base.pagamento(usado) if usado else {'sit': 'sem', 'parcelas': [], 'aberto': 0},
            'obs': _txt(o.get('obs')), 'resp_ssw': responsavel_da_obs(o.get('obs')), 'padrao': padrao})
    # Fora da planilha: CTRB do mês combinado com o responsável que ninguém citou
    fora = []
    planilha_ctes = defaultdict(list)
    for s in saida:
        for ct in ((s['conferido'] or {}).get('ctes') or []):
            planilha_ctes[ct].append(s['ctrb_usado'])
    for k, r in base.oss.items():
        em = base.emissao(k)
        if k in usados or not em or em[:7] != mes:
            continue
        quem = responsavel_da_obs(r.get('obs'))
        if not quem or not mesmo_nome(quem, resp):
            continue
        c = conta(base, k)
        pg = base.pagamento(k)
        notas = []
        do_mesmo_cte = sorted({o for ct in (c or {}).get('ctes', []) for o in planilha_ctes.get(ct, []) if o})
        if do_mesmo_cte:
            notas.append(f"Mesmo CT-e do {', '.join(do_mesmo_cte)} da planilha — reemissão: não é outra viagem.")
            if pg['aberto'] > 0:
                notas.append(f"Tem {_brl(pg['aberto'])} a pagar no 477: confira com o financeiro antes de pagar.")
        elif c and c['irmaos']:
            notas.append(f"O mesmo CT-e tem outro CTRB do mesmo motorista ({', '.join(c['irmaos'])}).")
        if c is None:
            notas.append('Sem CT-e no Auditoria Receita.')
        elif c['resultado'] < 0:
            notas.append('Resultado negativo.')
        fora.append({'chave': f'X{k}', 'ctrb': k, 'emissao': em, 'motorista': _txt(r.get('motorista')),
                     'rota': f"{_txt(r.get('origem')) or '—'} → {_txt(r.get('destino')) or '—'}",
                     'valor': round(_num(r.get('vap')), 2), 'conferido': c, 'pagamento': pg, 'notas': notas,
                     'reemissao': bool(do_mesmo_cte), 'obs': _txt(r.get('obs')), 'padrao': 'fora'})
    fora.sort(key=lambda x: (x['reemissao'], x['emissao'] or ''))
    return {'linhas': saida, 'fora': fora, 'mes': mes, 'responsavel': responsavel, 'lido_em': ssw.get('lido_em'),
            'percentual': PERCENTUAL, 'tolerancia': TOLERANCIA}


def _ddmm(iso):
    return f'{iso[8:10]}/{iso[5:7]}' if iso else '—'


def valor_da_escolha(item, escolha):
    """Quanto a linha (ou o CTRB de fora) paga na escolha dada."""
    if escolha == 'fora':
        return 0.0
    if escolha == 'planilha':
        return round(((item.get('planilha') or {}).get('comissao')) or 0.0, 2)
    c = item.get('conferido')
    if c is None:
        return round(((item.get('planilha') or {}).get('comissao')) or 0.0, 2)
    return c['comissao']


def aplicar_decisoes(conf, decisoes):
    """Põe em cada item a escolha (a do diretor ou o padrão) e o valor a pagar; devolve os totais."""
    tot = {'planilha': 0.0, 'pagar': 0.0, 'confere': 0, 'diverge': 0, 'alertas': 0, 'fora_incluidos': 0}
    for it in conf['linhas']:
        d = decisoes.get(it['chave']) or {}
        it['escolha'] = d.get('escolha') or it['padrao']
        it['motivo_escolha'] = d.get('motivo')
        it['decidido_por'] = d.get('por')
        it['pagar'] = valor_da_escolha(it, it['escolha'])
        tot['planilha'] += (it['planilha'].get('comissao') or 0)
        tot['pagar'] += it['pagar']
        if it['status'] == 'confere':
            tot['confere'] += 1
        elif it['status'] == 'diverge':
            tot['diverge'] += 1
        if it['alertas']:
            tot['alertas'] += 1
    for it in conf['fora']:
        d = decisoes.get(it['chave']) or {}
        it['escolha'] = d.get('escolha') or it['padrao']
        it['motivo_escolha'] = d.get('motivo')
        it['pagar'] = valor_da_escolha(it, it['escolha'])
        tot['pagar'] += it['pagar']
        tot['fora_incluidos'] += 1 if it['escolha'] != 'fora' else 0
    tot['planilha'] = round(tot['planilha'], 2)
    tot['pagar'] = round(tot['pagar'], 2)
    tot['diferenca'] = round(tot['pagar'] - tot['planilha'], 2)
    tot['linhas'] = len(conf['linhas'])
    tot['fora'] = len(conf['fora'])
    conf['totais'] = tot
    return tot


# ── Banco ────────────────────────────────────────────────────────────────────

DDL = """
CREATE TABLE IF NOT EXISTS comissao_fechamento (
    id             SERIAL PRIMARY KEY,
    resp_chave     TEXT NOT NULL,                 -- nome normalizado (sem acento, maiúsculo)
    responsavel    TEXT NOT NULL,                 -- como a planilha escreve
    mes            TEXT NOT NULL,                 -- 'AAAA-MM', o da emissão dos CTRBs
    arquivo        TEXT,
    planilha       JSONB NOT NULL,                -- as linhas lidas da planilha (deste responsável)
    situacao       TEXT NOT NULL DEFAULT 'rascunho',   -- rascunho | aprovada
    conferencia    JSONB,                         -- foto da conferência no momento da aprovação
    total_planilha NUMERIC(14,2),
    total_pagar    NUMERIC(14,2),
    criado_por     TEXT,
    criado_em      TIMESTAMPTZ NOT NULL DEFAULT now(),
    aprovado_por   TEXT,
    aprovado_em    TIMESTAMPTZ,
    UNIQUE (resp_chave, mes)
);
CREATE TABLE IF NOT EXISTS comissao_decisao (
    fechamento_id  INT NOT NULL REFERENCES comissao_fechamento(id) ON DELETE CASCADE,
    chave          TEXT NOT NULL,                 -- 'L<linha da planilha>' ou 'X<CTRB de fora>'
    escolha        TEXT NOT NULL,                 -- conferida | planilha | fora | incluir
    motivo         TEXT,
    por            TEXT,
    em             TIMESTAMPTZ NOT NULL DEFAULT now(),
    PRIMARY KEY (fechamento_id, chave)
);
CREATE TABLE IF NOT EXISTS comissao_log (
    id             BIGSERIAL PRIMARY KEY,
    fechamento_id  INT,
    chave          TEXT,
    acao           TEXT NOT NULL,
    antes          TEXT,
    depois         TEXT,
    por            TEXT,
    em             TIMESTAMPTZ NOT NULL DEFAULT now()
);
CREATE INDEX IF NOT EXISTS comissao_log_fech ON comissao_log (fechamento_id, em);
"""

ESCOLHAS = ('conferida', 'planilha', 'fora', 'incluir')


def garantir_tabelas(cur):
    cur.execute(DDL)


def _log(cur, fid, chave, acao, antes, depois, por):
    cur.execute('INSERT INTO comissao_log (fechamento_id, chave, acao, antes, depois, por) VALUES (%s,%s,%s,%s,%s,%s)',
                (fid, chave, acao, antes, depois, por))


_SEL = """SELECT id, responsavel, resp_chave, mes, arquivo, situacao, total_planilha, total_pagar, criado_por,
                 to_char(criado_em AT TIME ZONE 'America/Sao_Paulo', 'DD/MM/YYYY HH24:MI'), aprovado_por,
                 to_char(aprovado_em AT TIME ZONE 'America/Sao_Paulo', 'DD/MM/YYYY HH24:MI'),
                 jsonb_array_length(planilha)
            FROM comissao_fechamento"""


def _fech(r):
    return {'id': r[0], 'responsavel': r[1], 'resp_chave': r[2], 'mes': r[3], 'mes_rotulo': rotulo_mes(r[3], True),
            'arquivo': r[4], 'situacao': r[5], 'total_planilha': float(r[6]) if r[6] is not None else None,
            'total_pagar': float(r[7]) if r[7] is not None else None, 'criado_por': r[8], 'criado_em': r[9],
            'aprovado_por': r[10], 'aprovado_em': r[11], 'linhas': r[12]}


def listar(cur):
    cur.execute(_SEL + ' ORDER BY mes DESC, responsavel')
    return [_fech(r) for r in cur.fetchall()]


def obter(cur, fid):
    cur.execute(_SEL + ' WHERE id = %s', (fid,))
    r = cur.fetchone()
    if not r:
        return None
    f = _fech(r)
    cur.execute('SELECT planilha, conferencia FROM comissao_fechamento WHERE id = %s', (fid,))
    f['planilha'], f['conferencia'] = cur.fetchone()
    return f


def existente(cur, resp_chave, mes):
    cur.execute(_SEL + ' WHERE resp_chave = %s AND mes = %s', (resp_chave, mes))
    r = cur.fetchone()
    return _fech(r) if r else None


def salvar_planilha(cur, responsavel, mes, arquivo, linhas, por):
    """Cria o fechamento (responsável, mês) ou troca a planilha de um rascunho. Aprovado não muda:
    tem de ser reaberto antes. Trocar a planilha apaga as decisões (as linhas podem ter mudado)."""
    rk = nome_norm(responsavel)
    atual = existente(cur, rk, mes)
    total = round(sum(l.get('comissao') or 0 for l in linhas), 2)
    if atual and atual['situacao'] == 'aprovada':
        raise ValueError(f'A comissão de {atual["mes_rotulo"]} de {atual["responsavel"]} já foi aprovada. '
                         f'Reabra antes de trocar a planilha.')
    if atual:
        cur.execute("""UPDATE comissao_fechamento SET planilha = %s, arquivo = %s, responsavel = %s,
                              total_planilha = %s, total_pagar = NULL, conferencia = NULL WHERE id = %s""",
                    (json.dumps(linhas, ensure_ascii=False), arquivo, responsavel, total, atual['id']))
        cur.execute('DELETE FROM comissao_decisao WHERE fechamento_id = %s', (atual['id'],))
        _log(cur, atual['id'], None, 'planilha trocada', str(atual['linhas']), str(len(linhas)), por)
        return atual['id'], True
    cur.execute("""INSERT INTO comissao_fechamento (resp_chave, responsavel, mes, arquivo, planilha, total_planilha, criado_por)
                   VALUES (%s,%s,%s,%s,%s,%s,%s) RETURNING id""",
                (rk, responsavel, mes, arquivo, json.dumps(linhas, ensure_ascii=False), total, por))
    fid = cur.fetchone()[0]
    _log(cur, fid, None, 'planilha anexada', None, f'{len(linhas)} linhas · {arquivo or ""}', por)
    return fid, False


def decisoes(cur, fid):
    cur.execute("""SELECT chave, escolha, motivo, por, to_char(em AT TIME ZONE 'America/Sao_Paulo', 'DD/MM HH24:MI')
                     FROM comissao_decisao WHERE fechamento_id = %s""", (fid,))
    return {k: {'escolha': e, 'motivo': m, 'por': p, 'em': w} for k, e, m, p, w in cur.fetchall()}


def decidir(cur, fid, chave, escolha, motivo, por, padrao):
    """Grava a escolha de uma linha. Voltar ao padrão sem motivo apaga a decisão."""
    cur.execute('SELECT escolha FROM comissao_decisao WHERE fechamento_id = %s AND chave = %s', (fid, chave))
    r = cur.fetchone()
    antes = r[0] if r else None
    motivo = _txt(motivo)
    if escolha == padrao and not motivo:
        cur.execute('DELETE FROM comissao_decisao WHERE fechamento_id = %s AND chave = %s', (fid, chave))
    else:
        cur.execute("""INSERT INTO comissao_decisao (fechamento_id, chave, escolha, motivo, por) VALUES (%s,%s,%s,%s,%s)
                       ON CONFLICT (fechamento_id, chave) DO UPDATE SET escolha = EXCLUDED.escolha,
                              motivo = EXCLUDED.motivo, por = EXCLUDED.por, em = now()""",
                    (fid, chave, escolha, motivo, por))
    _log(cur, fid, chave, 'decisão', antes or padrao, escolha + (f' — {motivo}' if motivo else ''), por)


def aprovar(cur, fid, conf, por):
    tot = conf['totais']
    cur.execute("""UPDATE comissao_fechamento SET situacao = 'aprovada', conferencia = %s, total_pagar = %s,
                          total_planilha = %s, aprovado_por = %s, aprovado_em = now()
                    WHERE id = %s AND situacao = 'rascunho'""",
                (json.dumps(conf, ensure_ascii=False), tot['pagar'], tot['planilha'], por, fid))
    if cur.rowcount != 1:
        raise ValueError('Esta comissão não está mais em rascunho.')
    _log(cur, fid, None, 'aprovada', None, f"{tot['pagar']:.2f}", por)


def reabrir(cur, fid, por):
    cur.execute("""UPDATE comissao_fechamento SET situacao = 'rascunho', conferencia = NULL, aprovado_por = NULL,
                          aprovado_em = NULL WHERE id = %s AND situacao = 'aprovada'""", (fid,))
    if cur.rowcount != 1:
        raise ValueError('Esta comissão não está aprovada.')
    _log(cur, fid, None, 'reaberta', None, None, por)


def excluir(cur, fid, por):
    cur.execute("DELETE FROM comissao_fechamento WHERE id = %s AND situacao = 'rascunho'", (fid,))
    if cur.rowcount != 1:
        raise ValueError('Só um rascunho pode ser excluído (reabra a aprovada antes).')
    _log(cur, fid, None, 'excluída', None, None, por)


def historico(cur, fid):
    cur.execute("""SELECT chave, acao, antes, depois, por, to_char(em AT TIME ZONE 'America/Sao_Paulo', 'DD/MM/YYYY HH24:MI')
                     FROM comissao_log WHERE fechamento_id = %s ORDER BY em DESC, id DESC LIMIT 300""", (fid,))
    return [{'chave': c, 'acao': a, 'antes': x, 'depois': d, 'por': p, 'em': e} for c, a, x, d, p, e in cur.fetchall()]


def comissionados(cur, exceto=None):
    """{CTRB: 'out/26 · Fulano'} de toda comissão aprovada (menos `exceto`) — o alerta "já comissionado"
    vale para qualquer mês e qualquer responsável."""
    cur.execute("SELECT id, responsavel, mes, conferencia FROM comissao_fechamento WHERE situacao = 'aprovada' AND id <> %s",
                (exceto or 0,))
    out = {}
    for fid, resp, mes, conf in cur.fetchall():
        rot = f'{rotulo_mes(mes)} · {resp}'
        for it in (conf or {}).get('linhas', []):
            if it.get('escolha') != 'fora':
                for k in {it.get('ctrb'), it.get('ctrb_usado')} - {None}:
                    out[k] = rot
        for it in (conf or {}).get('fora', []):
            if it.get('escolha') not in (None, 'fora'):
                out[it['ctrb']] = rot
    return out


# ── Excel para o financeiro ──────────────────────────────────────────────────

ROT_ESCOLHA = {'conferida': 'Conferida', 'planilha': 'Planilha', 'fora': 'Fora', 'incluir': 'Incluída'}


def exportar_xlsx(fech, conf):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    cols = ['Linha', 'Data', 'Cliente', 'Origem', 'Destino', 'Motorista', 'CTRB (planilha)', 'CTRB (SSW)',
            'Frete s/ ICMS (planilha)', 'Custo (planilha)', 'Resultado (planilha)', 'Comissão (planilha)',
            'Frete s/ ICMS (conferido)', 'Descarga (fora da conta)', 'Frete motorista', 'Vale-pedágio pago pela Rizza',
            'Resultado (conferido)', 'Comissão (conferida)', 'Decisão', 'A pagar', 'Situação', 'Motivo / alertas',
            'Pagamento do CTRB']
    wb = Workbook()
    ws = wb.active
    ws.title = 'Comissão'
    ws.cell(row=1, column=1, value=f"Comissão de frete — {fech['responsavel']} — {rotulo_mes(fech['mes'], True)}").font = Font(bold=True, size=13)
    t = conf['totais']
    sit = 'APROVADA' + (f" por {fech['aprovado_por']} em {fech['aprovado_em']}" if fech.get('aprovado_por') else '') \
        if fech['situacao'] == 'aprovada' else 'RASCUNHO (não aprovada)'
    ws.cell(row=2, column=1, value=f"A pagar: R$ {t['pagar']:,.2f} · Planilha pedia: R$ {t['planilha']:,.2f} · {sit}"
            .replace(',', 'X').replace('.', ',').replace('X', '.'))
    azul = PatternFill('solid', fgColor='1F4E78')
    for j, cab in enumerate(cols, start=1):
        c = ws.cell(row=4, column=j, value=cab)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = azul
        c.alignment = Alignment(wrap_text=True, vertical='center')
    sit_rot = {'confere': 'Confere', 'diverge': 'Valor diferente', 'sem_base': 'Sem base no SSW'}
    pag_rot = {'sem': 'Sem lançamento', 'pendente': 'Pendente', 'parcial': 'Parcial', 'pago': 'Pago'}
    i = 5
    for it in conf['linhas']:
        p, c = it['planilha'], it.get('conferido') or {}
        notas = list(it.get('motivos') or []) + [f"{a['texto']}: {a['detalhe']}" for a in it.get('alertas') or []]
        if it.get('motivo_escolha'):
            notas.append(f"Decisão: {it['motivo_escolha']}")
        vals = [p.get('lin'), _data(p.get('data')), p.get('cliente'), p.get('origem'), p.get('destino'), p.get('motorista'),
                p.get('ctrb'), it.get('ctrb_usado'), p.get('frete_s_icms'), p.get('total'), p.get('resultado'),
                p.get('comissao'), c.get('receita'), c.get('descarga') or None, c.get('frete_motorista'),
                (c.get('vale_pedagio') if c.get('vale_pago_rizza') else None), c.get('resultado'), c.get('comissao'),
                ROT_ESCOLHA.get(it['escolha'], it['escolha']), it['pagar'], sit_rot.get(it['status']),
                ' | '.join(notas) or None, pag_rot.get(it['pagamento']['sit'])]
        _linha_xlsx(ws, i, vals)
        i += 1
    for it in conf['fora']:
        if it['escolha'] == 'fora':
            continue
        c = it.get('conferido') or {}
        vals = [None, _data(it.get('emissao')), c.get('cliente'), it['rota'], None, it.get('motorista'), None, it['ctrb'],
                None, None, None, None, c.get('receita'), c.get('descarga') or None, c.get('frete_motorista'),
                (c.get('vale_pedagio') if c.get('vale_pago_rizza') else None), c.get('resultado'), c.get('comissao'),
                'Incluída (fora da planilha)', it['pagar'], None, ' | '.join(it.get('notas') or []) or None,
                pag_rot.get(it['pagamento']['sit'])]
        _linha_xlsx(ws, i, vals)
        i += 1
    ws.cell(row=i, column=19, value='TOTAL').font = Font(bold=True)
    c = ws.cell(row=i, column=20, value=t['pagar'])
    c.font = Font(bold=True)
    c.number_format = '#,##0.00'
    larg = {'Cliente': 22, 'Origem': 18, 'Destino': 18, 'Motorista': 28, 'Motivo / alertas': 70, 'Decisão': 14}
    for j, cab in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = larg.get(cab, 14)
    ws.row_dimensions[4].height = 42
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A4:{get_column_letter(len(cols))}{max(5, i - 1)}'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _data(iso):
    try:
        return date.fromisoformat(iso[:10]) if iso else None
    except ValueError:
        return None


def _linha_xlsx(ws, i, vals):
    for j, v in enumerate(vals, start=1):
        c = ws.cell(row=i, column=j, value=v)
        if isinstance(v, date):
            c.number_format = 'DD/MM/YYYY'
        elif isinstance(v, float):
            c.number_format = '#,##0.00'
