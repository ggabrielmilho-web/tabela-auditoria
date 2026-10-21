# -*- coding: utf-8 -*-
"""Descargas Nestlé (5100) — 5ª aba do CRM - Documentos: a descarga da Nestlé paga a terceiro até o
reembolso.

Cada linha é um lançamento do 477 no evento 5100 (CARGA E DESCARGA C/ TERCEIROS) desde DESDE cujo CT-e
é da NESTLÉ/Garoto (tomador do CT-e). O CRM usava o relatório da segunda-feira (relatorio_5100), com
todos os clientes, para ver o que já tinha comprovante; ela só cobra a Nestlé (pedido de 21/10/2026).

Como se acha o CT-e da carga (e por ele, que é da Nestlé):
  1. a NF escrita no histórico, com a régua do relatório (`relatorio_5100.escolher_cte`);
  2. sem NF, o nº da carga Nestlé (B10…: "CARREGAMENTO NESTLE B101234567 - PIX …") na observação dos
     CT-e da Nestlé — fica o CT-e principal da carga (não o complemento), o mais perto da despesa;
  3. a NF casou com CT-e de OUTRO cliente mas o histórico diz Nestlé (NF de mesmo número): procura de
     novo só entre os CT-e da Nestlé;
  4. o histórico cita Nestlé/Garoto e nada leva ao CT-e: entra "sem CT-e", para ela digitar.
Medido em 2026 (21/10): 269 lançamentos da Nestlé — 157 pela NF, 85 pelo B10, 27 só pelo nome.

O reembolso sai do próprio SSW, e só por ligação EXPLÍCITA no documento:
  * CT-e complementar cuja observação cita o CTRC da carga ("CTRC EMITIDO P/ COMPL. DO UDI123456-7");
  * nota de serviço de descarga do mesmo pagador que cita a NF da carga ("COBRANCA REF A DESCARGA DAS
    NFS 100001 / 100005 A 100008"), emitida depois da carga.
O VALOR não prova nada (a Nestlé reembolsa a 1,252 × a despesa, outros pelo custo); o casamento solto
(mesma NF + mesmo pagador) pegava CARGA FECHADA e SUBSTITUTO de outra perna.

O que só ela sabe vai para a marcação (crm_doc_marcacao, carteira 'descargas'): comprovante recebido,
ticket FSC e o status dele, OC, pedido de compras, doc. transporte, tickets rejeitados, CT-e à mão e
"não cobrável" — as colunas do controle "Custo Extra" que ela mandou de exemplo. O fato do SSW vence a
marcação: com o complemento emitido, a linha é "cobrada".
"""
import os
import re
import threading
import time
from datetime import date, datetime, timedelta

import crm_documentos as crm
import relatorio_5100 as r5

DESDE = date.fromisoformat(os.getenv('CRM_DESCARGAS_DESDE', '2026-01-01'))
EVENTO = r5.EVENTO
SITUACOES = {s['id']: s for s in crm.DESCARGAS['situacoes']}

_D = "'public consulta_despesas_477'"
_T = "'public conhecimentos_emitidos'"
COLS_477 = ('emissao', 'uni', 'numlancto', 'parcela', 'nome_fornecedor', 'nome_favorecido', 'vlr_final',
            'historico_despesa', 'nfiscal', 'vencimen', 'data_pgto', 'sit_des', 'data_importacao')

# "COMPL. DO UDI123456-7" / "COMPLEMENTO DO CTRC UDI 123456-7"
RE_COMPL = re.compile(r'COMPL\w*\.?\s+DO\s+(?:CTRC\s+|CT-?E\s+)?([A-Z]{3})\s?(\d{6})-?(\d)', re.I)
# "NFS 100001 / 100005 A 100008 / 100010" — a lista que vem depois de NF/NFS
RE_NF_LISTA = re.compile(r'\bNFS?\b[\s.:]*((?:\d{3,9}(?:\s*(?:/|,|\bE\b|\bA\b)\s*)?)+)', re.I)
# a carga da Nestlé (shipment): B + 9 dígitos, B101234567
RE_B10 = re.compile(r'\bB\s?(\d{9})\b', re.I)
RE_NESTLE = re.compile(r'NESTL|GAROTO', re.I)
NESTLE_RAIZES = set(crm.CARTEIRAS['nestle']['raizes'])


def _iso(v):
    return str(v)[:10] if v else None


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def eh_nestle(c):
    return str((c or {}).get('pag') or '')[:8] in NESTLE_RAIZES


def cargas_b10(texto):
    """Texto → as cargas Nestlé citadas ('B101234567'), sem repetir."""
    return list(dict.fromkeys('B' + n for n in RE_B10.findall(str(texto or '').upper())))


def chave_lancamento(d):
    """O lançamento do 477 → chave da marcação. uni + lançamento + parcela é único (1.947 de 1.947)."""
    n = str(d.get('numlancto') or '').strip()
    n = str(int(float(n))) if re.fullmatch(r'\d+(\.0+)?', n) else n
    p = str(d.get('parcela') or '').strip()
    p = str(int(float(p))) if re.fullmatch(r'\d+(\.0+)?', p) else p
    return f"L477-{str(d.get('uni') or '').strip()}-{n}-{p}"


def ctrc_complementado(obs):
    """Observação do CT-e → o CTRC que ele complementa, ou None. Substituto não é complemento."""
    s = str(obs or '').upper()
    if 'SUBSTITU' in s:
        return None
    m = RE_COMPL.search(s)
    return f'{m[1]}{m[2]}-{m[3]}' if m else None


def nfs_de_descarga(obs):
    """Observação da nota de serviço → as NFs cuja DESCARGA ela cobra (com as faixas "A"), ou []."""
    s = str(obs or '').upper()
    if 'DESCARG' not in s:
        return []
    out = []
    for grupo in RE_NF_LISTA.findall(s):
        toks = re.findall(r'\d{3,9}|\bA\b', grupo)
        i = 0
        while i < len(toks):
            if toks[i] == 'A' and out and i + 1 < len(toks) and toks[i + 1] != 'A':
                a, b = int(out[-1]), int(toks[i + 1])
                if 0 < b - a <= 50:                  # faixa curta; "A" solto não vira mil notas
                    out += [str(x) for x in range(a + 1, b + 1)]
                else:
                    out.append(toks[i + 1])
                i += 2
                continue
            if toks[i] != 'A':
                out.append(toks[i].lstrip('0'))
            i += 1
    return list(dict.fromkeys(out))


def indices(ce):
    """CT-e (todos os clientes) → índices: por CTRC, por NF (régua do relatório; e só os da Nestlé),
    por carga B10 (só Nestlé), complementos por CTRC complementado e notas de descarga por (NF, raiz)."""
    por_ctrc, comp, nfs, b10 = {}, {}, {}, {}
    for c in ce:
        k = crm.chave_ctrc(c.get('ctrc'))
        if not k:
            continue
        por_ctrc[k] = c
        alvo = ctrc_complementado(c.get('obs'))
        if alvo and alvo != k:
            comp.setdefault(alvo, []).append(c)
        raiz = str(c.get('pag') or '')[:8]
        for nf in nfs_de_descarga(c.get('obs')):
            nfs.setdefault((nf, raiz), []).append(c)
        if eh_nestle(c):
            for b in cargas_b10(c.get('obs')):
                b10.setdefault(b, []).append(c)
    return {'por_ctrc': por_ctrc, 'por_nf': r5.indexar_ctes(ce), 'comp': comp, 'nfs': nfs, 'b10': b10,
            'por_nf_nestle': r5.indexar_ctes([c for c in ce if eh_nestle(c)])}


def _escolha_b10(cands, emissao, base, b):
    """Os CT-e da Nestlé com a carga B10 -> o principal (não complemento nem nota), emitido antes ou no
    dia da despesa e o mais perto; na mesma forma do `relatorio_5100.escolher_cte`."""
    vistos = {crm.chave_ctrc(c['ctrc']): c for c in cands}
    princ = [c for c in vistos.values() if not str(c.get('doc') or '').upper().startswith(('COMPLEMENTAR', 'SUBC'))] \
        or list(vistos.values())
    dd = _iso(emissao) or '9999'
    antes = sorted([c for c in princ if (_iso(c.get('emissao')) or '') <= dd], key=lambda c: _iso(c.get('emissao')) or '')
    esc = antes[-1] if antes else min(princ, key=lambda c: _iso(c.get('emissao')) or '')
    k = crm.chave_ctrc(esc['ctrc'])
    # o complemento da carga não entra em "outros": ele já aparece como o reembolso
    outros = [crm.chave_ctrc(c['ctrc']) for c in princ if crm.chave_ctrc(c['ctrc']) != k]
    return dict(base, cte=k, cte_numero=esc.get('cte_num') or '', cte_remetente=esc.get('cli_rem') or '',
                cte_emissao=_iso(esc.get('emissao')) or '', cte_outros=', '.join(outros),
                cte_criterio=f'pela carga {b}' + (f' ({len(princ)} CT-e da Nestlé com ela)' if len(princ) > 1 else ''))


def cte_nestle(historico, emissao, idx):
    """Histórico do lançamento -> (escolha no formato de `escolher_cte`, é da Nestlé?). Ver o cabeçalho."""
    auto = r5.escolher_cte(historico, emissao, idx['por_nf'])
    c = idx['por_ctrc'].get(crm.chave_ctrc(auto['cte']))
    if c and eh_nestle(c):
        return auto, True
    for b in cargas_b10(historico):
        if idx['b10'].get(b):
            return _escolha_b10(idx['b10'][b], emissao, auto, b), True
    if RE_NESTLE.search(str(historico or '')):
        alt = r5.escolher_cte(historico, emissao, idx['por_nf_nestle'])
        if alt['cte']:
            return dict(alt, cte_criterio=alt['cte_criterio'] + ' — só entre os CT-e da Nestlé (a NF também é de outro cliente)'), True
        return dict(auto, cte='', cte_numero='', cte_outros='', cte_remetente='', cte_emissao='',
                    cte_criterio='o histórico cita a Nestlé, mas sem NF nem carga B10 que leve ao CT-e'), True
    return auto, False


def reembolsos_da_carga(ctrc, idx):
    """Os documentos que cobram a descarga desta carga (complemento que cita o CTRC, ou nota de
    descarga do mesmo pagador que cita a NF dela, emitida depois)."""
    if not ctrc:
        return []
    docs = list(idx['comp'].get(ctrc, []))
    o = idx['por_ctrc'].get(ctrc)
    if o:
        nf, raiz, em = str(o.get('nf') or '').lstrip('0'), str(o.get('pag') or '')[:8], _iso(o.get('emissao')) or ''
        docs += [c for c in idx['nfs'].get((nf, raiz), []) if (_iso(c.get('emissao')) or '') >= em]
    vistos, out = set(), []
    for c in docs:
        k = crm.chave_ctrc(c.get('ctrc'))
        if k and k != ctrc and k not in vistos:
            vistos.add(k)
            out.append(c)
    return sorted(out, key=lambda c: _iso(c.get('emissao')) or '')


def situacao_doc(doc, pagamentos, abertas, hoje):
    """Documento de reembolso -> pago / em fatura (vencida ou não) / só emitido. Mesmas fontes do CRM:
    441 das pagas, baixa LIQUIDADO do 455, carteira em aberto do 441 (com o pagamento parcial)."""
    k = crm.chave_ctrc(doc.get('ctrc'))
    pg = (pagamentos or {}).get(k)
    if pg:
        return {'situacao': 'pago', 'pago_em': max(p['pagamento'] for p in pg),
                'fatura': ', '.join(sorted({str(p['fatura']) for p in pg}))}
    if str(doc.get('baixa') or '').upper() == 'LIQUIDADO':
        return {'situacao': 'pago', 'pago_em': _iso(doc.get('liq')), 'fatura': ''}
    fas = (abertas or {}).get(k) or []
    if fas:
        f = min(fas, key=lambda x: x['vencimento'] or '9999-12-31')
        if k in (f.get('pagos_parcial') or []):
            return {'situacao': 'pago', 'pago_em': None, 'fatura': f['fatura']}
        vencida = bool(f['vencimento'] and f['vencimento'] < hoje.isoformat())
        return {'situacao': 'vencida' if vencida else 'fatura', 'fatura': f['fatura'], 'vencimento': f['vencimento']}
    return {'situacao': 'emitido'}


def situacao_linha(reembolsos, campos, tem_cte):
    """O fato do SSW primeiro; depois o que ela marcou, do fim do fluxo para o começo. Função pura."""
    if reembolsos:
        return 'reembolso_pago' if all(r['situacao'] == 'pago' for r in reembolsos) else 'cobrado'
    if campos.get('nao_cobravel'):
        return 'nao_cobravel'
    st = crm._sem_acento(str(campos.get('status_ticket') or '')).upper()
    if st.startswith('NEGAD'):
        return 'ticket_negado'
    if st.startswith('CONCLU'):
        return 'ticket_aprovado'
    if campos.get('ticket'):
        return 'ticket_aberto'
    if not tem_cte:
        return 'sem_cte'
    return 'cobrar' if campos.get('comprovante') else 'aguardando_comprovante'


def montar(dd, marc, pagamentos, abertas, hoje=None):
    """Lançamentos do 5100 da Nestlé + marcações -> linhas da tela (os nomes que a tela genérica usa:
    chave, ctrc, cte, nf, obs, frete = o valor pago ao terceiro, emissao, idade, aberto, campos…)."""
    hoje = hoje or crm.hoje_brt()
    idx = dd['idx']
    linhas = []
    for d in dd['desp']:
        chave = chave_lancamento(d)
        m = marc.get(chave) or {}
        campos = m.get('campos') or {}
        auto, da_nestle = cte_nestle(d.get('historico_despesa'), d.get('emissao'), idx)
        manual = crm.chave_ctrc(campos.get('cte_manual'))
        if not (da_nestle or manual):                # outro cliente: não é desta aba
            continue
        ctrc = manual or crm.chave_ctrc(auto['cte']) or None
        o = idx['por_ctrc'].get(ctrc) or {}
        reembs = []
        for c in reembolsos_da_carga(ctrc, idx):
            k = crm.chave_ctrc(c['ctrc'])
            reembs.append(dict(situacao_doc(c, pagamentos, abertas, hoje), ctrc=k, cte=c.get('cte_num'),
                               doc=c.get('doc'), frete=_num(c.get('frete')), emissao=_iso(c.get('emissao')),
                               carteira=crm.carteira_da_raiz(str(c.get('pag') or '')[:8]),
                               obs=str(c.get('obs') or '')[:160]))
        sit = situacao_linha(reembs, campos, bool(ctrc))
        em = _iso(d.get('emissao'))
        cte_num = str((o.get('cte_num') if manual else auto['cte_numero']) or '').strip()
        linhas.append({
            'chave': chave, 'emissao': em, 'idade': (hoje - date.fromisoformat(em)).days if em else None,
            'uni': d.get('uni'), 'lancto': f"{chave.split('-')[2]}/{chave.split('-')[3]}",
            'fornecedor': d.get('nome_fornecedor'), 'favorecido': d.get('nome_favorecido'),
            'frete': _num(d.get('vlr_final')), 'obs': d.get('historico_despesa'),
            'nf': auto['nf_historico'] or None, 'nf_despesa': d.get('nfiscal'),
            'vencimento_despesa': _iso(d.get('vencimen')), 'pago_terceiro': _iso(d.get('data_pgto')),
            'ctrc': ctrc, 'cte': int(cte_num) if cte_num.isdigit() else None, 'cte_manual': bool(manual),
            'cte_outros': '' if manual else auto['cte_outros'],
            'cte_criterio': 'digitado à mão' if manual else auto['cte_criterio'],
            'cte_emissao': _iso(o.get('emissao')) or auto['cte_emissao'] or None,
            'remetente': o.get('cli_rem') or auto['cte_remetente'] or None,
            'pagador': o.get('pagador'), 'cnpj': o.get('pag'),
            'destinatario': o.get('dest'), 'destino': '/'.join(x for x in (o.get('cid_d'), o.get('uf_d')) if x) or None,
            'transporte': ', '.join(cargas_b10(d.get('historico_despesa')) or cargas_b10(o.get('obs'))) or None,
            'carteira': crm.carteira_da_raiz(str(o.get('pag') or '')[:8]) if o else None,
            'reembolsos': reembs, 'valor_reembolso': round(sum(r['frete'] for r in reembs), 2),
            'situacao': sit, 'aberto': SITUACOES[sit]['aberto'],
            'campos': campos, 'atualizado': m.get('atualizado'), 'por': m.get('por'), 'editado': bool(m.get('editado')),
            # o que a tela das carteiras espera em toda linha
            'etapa': None, 'diverg': [], 'novo': False, 'manual': False, 'ssw': None,
        })
    linhas.sort(key=lambda x: (x['emissao'] or '', x['chave']), reverse=True)
    return linhas


def resumo(linhas):
    """Para a aba: o que está aberto e ainda sem complemento (do aguardando comprovante ao ticket aprovado)."""
    sem = [l for l in linhas if l['aberto'] and l['situacao'] != 'cobrado']
    return {'abertos': len(sem), 'valor': round(sum(l['frete'] for l in sem), 2), 'rotulo': 'sem cobrança'}


# ── Fonte (Power BI) ─────────────────────────────────────────────────────────

def _d(x):
    return f'DATE({x.year},{x.month},{x.day})'


def buscar(dax_dre, dax_op, hoje=None):
    """477 do 5100 (dataset do DRE) + 455 de TODOS os clientes (dataset principal), mês a mês: a NF do
    histórico pode cair em CT-e de qualquer cliente, e só o tomador diz se é Nestlé. A consulta do BI
    corta em 15 MiB e quem chama tem de levantar erro quando corta."""
    hoje = hoje or crm.hoje_brt()
    desp = dax_dre(f'EVALUATE SELECTCOLUMNS(FILTER(ALL({_D}), {_D}[evento]="{EVENTO}" && '
                   f'{_D}[emissao] >= {_d(DESDE)} && {_D}[emissao] < {_d(hoje + timedelta(days=1))}), '
                   + ', '.join(f'"{c}",{_D}[{c}]' for c in COLS_477) + ')')
    ce = []
    a = (DESDE - timedelta(days=r5.FOLGA_CTE_DIAS)).replace(day=1)
    while a <= hoje:
        b = (a.replace(day=28) + timedelta(days=4)).replace(day=1)
        ce += dax_op(f'EVALUATE SELECTCOLUMNS(FILTER(ALL({_T}), {_T}[data_emissao] >= {_d(a)} && '
                     f'{_T}[data_emissao] < {_d(b)}), "ctrc",{_T}[serie_numero_ctrc], "cte_num",{_T}[serie_numero_cte], '
                     f'"nf",{_T}[numero_nota_fiscal], "cli_rem",{_T}[cliente_remetente], "emissao",{_T}[data_emissao], '
                     f'"pag",{_T}[cnpj_pagador], "pagador",{_T}[cliente_pagador], "doc",{_T}[tipo_documento], '
                     f'"frete",{_T}[valor_frete], "baixa",{_T}[tipo_baixa], "liq",{_T}[data_liquidacao], '
                     f'"dest",{_T}[cliente_destinatario], "cid_d",{_T}[cidade_destinatario], "uf_d",{_T}[uf_destinatario], '
                     f'"obs",LEFT({_T}[observacao], 300))')
        a = b
    return {'desp': desp, 'idx': indices(ce), 'n_ce': len(ce),
            'lido_477': max((_iso(x.get('data_importacao')) for x in desp if x.get('data_importacao')), default=None),
            'lido_em': datetime.now(crm.BRT).strftime('%d/%m/%Y %H:%M')}


_CACHE = {'v': None, 'ts': 0.0}
_TRAVA = threading.Lock()
CACHE_S = 600


def dados(dax_dre, dax_op, forcar=False):
    """`buscar` com cache de 10 min (como o das carteiras)."""
    with _TRAVA:
        if not forcar and _CACHE['v'] is not None and time.time() - _CACHE['ts'] < CACHE_S:
            return _CACHE['v']
        v = buscar(dax_dre, dax_op)
        _CACHE.update(v=v, ts=time.time())
        return v


def cache_quente():
    """Os dados se já foram lidos há menos de 10 min (para o resumo da aba sem ir ao BI), senão None."""
    return _CACHE['v'] if _CACHE['v'] is not None and time.time() - _CACHE['ts'] < CACHE_S else None


# ── Excel ────────────────────────────────────────────────────────────────────

def exportar_xlsx(linhas):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    dt = lambda v: date.fromisoformat(v) if v and re.fullmatch(r'\d{4}-\d{2}-\d{2}', str(v)) else v
    cp = lambda k: (lambda l: l['campos'].get(k))
    cols = [('Emissão', lambda l: dt(l['emissao'])), ('Uni', lambda l: l['uni']), ('Lançamento', lambda l: l['lancto']),
            ('Fornecedor', lambda l: l['fornecedor']), ('Valor pago', lambda l: l['frete']),
            ('Histórico', lambda l: l['obs']), ('NF despesa', lambda l: l['nf_despesa']),
            ('NF (histórico)', lambda l: l['nf']), ('Transporte', lambda l: l['transporte']),
            ('CTRC', lambda l: l['ctrc']), ('Nº CT-e', lambda l: l['cte']), ('Tomador', lambda l: l['pagador']),
            ('Destinatário', lambda l: l['destinatario']), ('Destino', lambda l: l['destino']),
            ('Situação', lambda l: SITUACOES[l['situacao']]['rotulo']),
            ('Complemento', lambda l: ', '.join(r['ctrc'] for r in l['reembolsos'])),
            ('Valor do complemento', lambda l: l['valor_reembolso'] or None),
            ('Pago em', lambda l: dt(max((r.get('pago_em') or '' for r in l['reembolsos']), default='') or None)),
            ('Comprovante recebido', lambda l: dt(l['campos'].get('comprovante'))),
            ('Ticket FSC', cp('ticket')), ('Status do ticket', cp('status_ticket')), ('OC', cp('oc')),
            ('Pedido de compras', cp('pedido_compras')), ('Doc. transporte', cp('doc_transporte')),
            ('Tickets rejeitados', cp('tickets_rejeitados')), ('CT-e (à mão)', cp('cte_manual')),
            ('Não cobrável (motivo)', cp('nao_cobravel')), ('Anotação', cp('anotacao')),
            ('Como achou o CT-e', lambda l: l['cte_criterio'])]
    wb = Workbook()
    ws = wb.active
    ws.title = 'Descargas Nestle 5100'
    azul = PatternFill('solid', fgColor='1F4E78')
    for j, (cab, _) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=cab)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = azul
    for i, l in enumerate(linhas, start=2):
        for j, (cab, f) in enumerate(cols, start=1):
            v = f(l)
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, date):
                c.number_format = 'DD/MM/YYYY'
            elif cab in ('Valor pago', 'Valor do complemento'):
                c.number_format = '#,##0.00'
    for j, (cab, _) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = {'Histórico': 55, 'Fornecedor': 30, 'Destinatário': 30}.get(cab, 15)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(1, len(linhas) + 1)}'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
