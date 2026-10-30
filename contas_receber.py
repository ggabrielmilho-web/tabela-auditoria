# -*- coding: utf-8 -*-
"""Contas a Receber: o faturamento e a cobrança da Rizza numa tela só (29/10/2026).

Substitui as três planilhas da Laura (faturamento):

1. "Relatório de Faturamento" (o .xls do SSW) — os CT-e ainda sem fatura. Ela tirava linha à mão
   todo dia e pintava os casos-problema.
2. "PDD.xlsx" — os CT-e que NÃO vão ser faturados (substituído, recusado, cobrança indevida...),
   com o motivo escrito por cima da observação. Não é provisão contábil: a contadora não usa.
3. "VENCIMENTO DIÁRIO - LAURA.xlsx" — as faturas vencidas e não pagas, com um STATUS livre (uma aba
   por mês de vencimento, mais "FATURAS SEM BAIXA" e "MOTORISTA").

Duas filas, e cada campo tem um dono (como no CRM - Documentos):

- **Fato do SSW (só leitura), pelo Power BI:** o CT-e e a baixa vêm do 455 (`conhecimentos_emitidos`,
  dataset principal); a fatura, do 441 — em aberto (`faturas_441_abertas(_ctrcs)`, lido às 04h e às
  15h) e paga (`faturas_441(_ctrcs)`).
- **Marcação delas (editável):** a situação, o motivo, a previsão e a anotação, em `cr_marcacao`, com
  cada mudança em `cr_log`. A leitura do SSW NUNCA sobrescreve uma marcação.

**A faturar** = CT-e autorizado desde INICIO com baixa PENDENTE no 455 e fora de qualquer fatura do 441
(nem em aberto, nem paga). Conferido em 29/10/2026 contra o relatório dela de 28/10: 1.788 das 1.973
linhas desse universo estavam na lista; 148 das 185 restantes estavam na PDD.

Por que desde 2026: o 441 das pagas cobre desde 01/01/2026, e o 441 em aberto só lista CTRC recente.
Um CT-e de 2025 pago em 2025 não está em fatura nenhuma que a tela enxergue e a baixa do 455 fica
congelada no mês da autorização — ele pareceria "sem fatura" sem estar.

Os CT-e cancelados não chegam: o Power BI filtra CANCELADO e LIQU OCOR do 455. Um cancelamento depois
da virada do mês só chega pela releitura dos 3 meses anteriores (ssw_relatorio 455 releitura.py, 04:20).

Este módulo não fala Flask nem abre conexão: recebe as funções de consulta e o cursor.
"""

import io
import json
import re
import threading
import time
import unicodedata
from datetime import date, datetime, timedelta, timezone

BRT = timezone(timedelta(hours=-3))

INICIO = '2026-01-01'          # primeiro dia (autorização) da fila A faturar
CORTE_ANTIGAS = '2026-05-01'   # fatura vencida antes disto vai para o cartão "Vencidas antigas"
SALDO_MIN = 1.0                # resíduo de fatura (13 de R$ 0,01 em 29/10/2026) não é cobrança
PRAZO_CONFIRMA = 3             # dias para o 441 mostrar a fatura que ela marcou como emitida
RECEBIDAS_DIAS = 7


def hoje_brt():
    return datetime.now(BRT).date()


# ── Marcações ────────────────────────────────────────────────────────────────
# CT-e: a situação é a "etapa". Sem etapa = a faturar (o que ainda pede ação).
_S = lambda cod, rotulo, cor, desc: {'cod': cod, 'rotulo': rotulo, 'cor': cor, 'desc': desc}
SITUACOES_CTE = [
    _S('PENDENCIA', 'Pendência', 'laranja', 'falta algo para faturar (era a linha que ela pintava)'),
    _S('EMITIDA', 'Fatura emitida', 'azul', 'a fatura foi emitida no SSW; sai da fila e o 441 confirma'),
    _S('PDD', 'Não será faturado', 'roxo', 'o CT-e não vai ser faturado (era a planilha PDD)'),
]
# Os motivos da PDD dela (29/10/2026), agrupados pelo que estava escrito por cima da observação.
MOTIVOS_PDD = ['SUBSTITUÍDO', 'RECUSADO / DESACORDO', 'COBRANÇA INDEVIDA / EMITIDO ERRADO', 'QUITAÇÃO NESTLÉ',
               'SEM CUSTO / NÃO SERÁ PAGO', 'CONTRATO SINGULAR', 'SINISTRO / ACORDO', 'CANCELADO', 'OUTRO']

# Fatura: a situação da cobrança, tirada dos STATUS que ela escrevia (29/10/2026).
SITUACOES_FATURA = ['COBRADO', 'PREVISÃO DE PAGAMENTO', 'PENDENTE PROGRAMAÇÃO', 'EM TRATATIVA',
                    'AGUARDANDO COMPROVANTE', 'DESCONTO DO CLIENTE', 'SEM PREVISÃO', 'NEGOCIAÇÃO / JURÍDICO']
COR_SITUACAO_FATURA = {'COBRADO': 'azul', 'PREVISÃO DE PAGAMENTO': 'verde', 'PENDENTE PROGRAMAÇÃO': 'amarelo',
                       'EM TRATATIVA': 'laranja', 'AGUARDANDO COMPROVANTE': 'ciano', 'DESCONTO DO CLIENTE': 'roxo',
                       'SEM PREVISÃO': 'vermelho', 'NEGOCIAÇÃO / JURÍDICO': 'magenta'}
# As duas abas dela que não eram mês: o motorista que deve à Rizza e o resíduo de desconto do cliente.
CLASSES_FATURA = ['MOTORISTA', 'DESCONTO SEM BAIXA']

CAMPOS = {
    'cte': [
        {'id': 'motivo', 'rotulo': 'Motivo', 'tipo': 'lista', 'opcoes': MOTIVOS_PDD},
        {'id': 'anotacao', 'rotulo': 'Anotação', 'tipo': 'texto'},
    ],
    'fatura': [
        {'id': 'situacao', 'rotulo': 'Situação', 'tipo': 'lista', 'opcoes': SITUACOES_FATURA},
        {'id': 'previsao', 'rotulo': 'Previsão', 'tipo': 'data'},
        {'id': 'classe', 'rotulo': 'Classificação', 'tipo': 'lista', 'opcoes': CLASSES_FATURA},
        {'id': 'anotacao', 'rotulo': 'Anotação', 'tipo': 'texto'},
    ],
}
TIPOS = ('cte', 'fatura')


def config_publica():
    return {'situacoes_cte': SITUACOES_CTE, 'campos': CAMPOS, 'cores_fatura': COR_SITUACAO_FATURA,
            'inicio': INICIO, 'corte_antigas': CORTE_ANTIGAS, 'prazo_confirma': PRAZO_CONFIRMA,
            'recebidas_dias': RECEBIDAS_DIAS}


# ── Normalização ─────────────────────────────────────────────────────────────

def _sem_acento(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s or '') if unicodedata.category(c) != 'Mn')


def _chave_opcao(s):
    return re.sub(r'[^A-Z0-9]', '', _sem_acento(str(s)).upper())


def chave_ctrc(v):
    """'udi 123456-7 ' → 'UDI123456-7'; None se não tiver a forma de um CTRC do SSW."""
    s = re.sub(r'\s+', '', str(v or '')).upper()
    return s if re.fullmatch(r'[A-Z]{3}\d{6}-\d', s) else None


def chave_fatura(v):
    """'0039706-7 ' → '0039706-7' (o número como o 441 escreve)."""
    s = re.sub(r'\s+', '', str(v or '').replace('\xa0', ''))
    return s or None


def ler_data(s, hoje=None):
    """O que a pessoa digitou → ISO. Aceita 13/07, 13/7/26, 13/07/2026 e 2026-07-13. Sem ano: o ano que
    deixa a data mais perto de hoje (05/01 digitado em dezembro é o janeiro que vem; "previsão 25/08"
    lida em outubro é a de agosto que passou)."""
    s = (s or '').strip()
    if not s:
        return None
    hoje = hoje or hoje_brt()
    m = re.fullmatch(r'(\d{4})-(\d{2})-(\d{2})', s)
    try:
        if m:
            return date(int(m[1]), int(m[2]), int(m[3])).isoformat()
        m = re.fullmatch(r'(\d{1,2})/(\d{1,2})(?:/(\d{2}|\d{4}))?', s)
        if not m:
            raise ValueError
        d, mes, a = int(m[1]), int(m[2]), m[3]
        if a:
            return date(int(a) + (2000 if len(a) == 2 else 0), mes, d).isoformat()
        opcoes = []
        for ano in (hoje.year - 1, hoje.year, hoje.year + 1):
            try:
                opcoes.append(date(ano, mes, d))
            except ValueError:                          # 29/02 fora de ano bissexto
                pass
        if not opcoes:
            raise ValueError
        return min(opcoes, key=lambda x: abs((x - hoje).days)).isoformat()
    except ValueError:
        raise ValueError(f'"{s}" não é uma data — use dd/mm ou dd/mm/aaaa.')


def normalizar_entrada(tipo, campo, valor, hoje=None):
    """Valida e normaliza uma edição da tela. ValueError com a mensagem para a pessoa."""
    if tipo not in TIPOS:
        raise ValueError('Tipo desconhecido.')
    v = (valor or '').strip() if isinstance(valor, str) else ('' if valor is None else str(valor))
    if campo == 'etapa':
        if tipo != 'cte':
            raise ValueError('A fatura não tem etapa — use a Situação.')
        if not v:
            return None
        cods = [s['cod'] for s in SITUACOES_CTE]
        u = _chave_opcao(v)
        achou = next((s['cod'] for s in SITUACOES_CTE if _chave_opcao(s['cod']) == u or _chave_opcao(s['rotulo']) == u), None)
        if not achou:
            raise ValueError(f'Situação "{v}" não existe: use {", ".join(cods)}.')
        return achou
    c = next((x for x in CAMPOS[tipo] if x['id'] == campo), None)
    if not c:
        raise ValueError(f'Campo "{campo}" não existe.')
    if not v:
        return None
    if len(v) > 500:
        raise ValueError('Texto maior que 500 caracteres.')
    if c['tipo'] == 'data':
        return ler_data(v, hoje)
    if c['tipo'] == 'lista':
        o = next((o for o in c['opcoes'] if _chave_opcao(o) == _chave_opcao(v)), None)
        if o is None:
            raise ValueError(f'{c["rotulo"]} "{v}" não existe — use {", ".join(c["opcoes"])}.')
        return o
    return re.sub(r'[ \t]+', ' ', v)


# ── Fontes (SSW pelo Power BI) ───────────────────────────────────────────────

_T455 = "'public conhecimentos_emitidos'"
_DATA_455 = f'COALESCE({_T455}[data_autorizacao], {_T455}[data_emissao])'
_COLS_455 = (f'"ctrc", {_T455}[serie_numero_ctrc], "cte", {_T455}[serie_numero_cte], "aut", {_DATA_455}, '
             f'"emi", {_T455}[data_emissao], "doc", {_T455}[tipo_documento], "baixa", {_T455}[tipo_baixa], '
             f'"frete", {_T455}[valor_frete], "cnpj", {_T455}[cnpj_pagador], "pagador", {_T455}[cliente_pagador], '
             f'"unid", {_T455}[unidade_responsavel_pagador], "nf", {_T455}[numero_nota_fiscal], '
             f'"rem", {_T455}[cliente_remetente], "dest", {_T455}[cliente_destinatario], '
             f'"cid_d", {_T455}[cidade_destinatario], "uf_d", {_T455}[uf_destinatario], '
             f'"imp", {_T455}[data_importacao], "obs", LEFT({_T455}[observacao], 240), '
             f'"cav", {_T455}[placa_cavalo], "car", {_T455}[placa_carreta], "mf", {_T455}[ultimo_manifesto]')

# O motorista não está no 455: vem do último manifesto do CT-e (o 455 escreve "CAR 004648-5", a tabela
# de manifestos "CAR004648-5"). As placas ficam as do 455 — iguais às do manifesto e na grafia do SSW.
# O manifesto corta o nome em 30 letras (589 de 5.758 em 30/10/2026): o nome inteiro vem do cadastro de
# motoristas (047) pelo CPF; sem cadastro, fica o do manifesto.
_TMF = "'public manifestos'"
DAX_MOTORISTAS_047 = "EVALUATE SELECTCOLUMNS('public motoristas_047', \"cpf\", [cpf], \"nome\", [nome])"


def dax_manifestos(ini):
    return (f'EVALUATE SELECTCOLUMNS(FILTER({_TMF}, {_TMF}[data_emissao] >= DATE({ini.year},{ini.month},{ini.day})), '
            f'"mf", {_TMF}[CHAVE_MANIFESTO], "nome", {_TMF}[nome_motorista], "cpf", {_TMF}[cpf_motorista])')


def nomes_motoristas(manifestos, cadastro):
    """{manifesto: nome do motorista} — o do cadastro 047 pelo CPF, senão o (cortado) do manifesto."""
    so_dig = lambda v: re.sub(r'\D', '', str(v or ''))
    inteiro = {so_dig(r.get('cpf')): _txt(r.get('nome')) for r in cadastro if so_dig(r.get('cpf'))}
    return {chave_manifesto(r.get('mf')): inteiro.get(so_dig(r.get('cpf'))) or _txt(r.get('nome'))
            for r in manifestos if r.get('mf')}


def chave_manifesto(v):
    s = re.sub(r'\s+', '', str(v or '')).upper()
    return s or None


def dax_455(ini, fim):
    return (f'EVALUATE SELECTCOLUMNS(FILTER({_T455}, {_DATA_455} >= DATE({ini.year},{ini.month},{ini.day}) && '
            f'{_DATA_455} < DATE({fim.year},{fim.month},{fim.day})), {_COLS_455})')


DAX_PAGAS_CTRC = "EVALUATE SELECTCOLUMNS('public faturas_441_ctrcs', \"fatura\", [fatura], \"ctrc\", [ctrc])"
DAX_PAGAS_FAT = ("EVALUATE SELECTCOLUMNS('public faturas_441', \"fatura\", [fatura], \"cliente\", [cliente], "
                 "\"cnpj\", [cnpj_cpf], \"emissao\", [emissao], \"venc\", [vencimen], \"pagamento\", [pagamento], "
                 "\"vlr_fatur\", [vlr_fatur], \"vlr_pago\", [vlr_pago], \"imp\", [data_importacao])")
_COLS_ABERTAS = ('fatura', 'cnpj_cpf', 'cliente', 'emissao', 'venc_ori', 'vencimen', 'data_promessa', 'dias_atraso',
                 'vlr_ctrcs', 'vlr_fatur', 'vlr_pago', 'saldo', 'desconto', 'pgto_parcial', 'vlr_debito',
                 'vlr_credito', 'motivo_debito_credito', 'tipo_de_cobranca', 'banco_carteira', 'ddd', 'telefone',
                 'ddd_telefone_cobranca', 'telefone_cobranca', 'e_mail_cobranca', 'protestada_cartorio',
                 'spc_serasa', 'prorrogada', 'ultima_ocorrencia', 'unidade_responsavel', 'cidade', 'uf',
                 'vendedor', 'data_importacao')
DAX_ABERTAS = ("EVALUATE SELECTCOLUMNS('public faturas_441_abertas', "
               + ', '.join(f'"{c}", [{c}]' for c in _COLS_ABERTAS) + ')')
DAX_ABERTAS_CTRC = ("EVALUATE SELECTCOLUMNS('public faturas_441_abertas_ctrcs', \"fatura\", [fatura], "
                    "\"ctrc\", [ctrc], \"cte\", [cte], \"emi\", [data_autorizacao], \"frete\", [valor_frete])")


def _iso(v):
    return str(v)[:10] if v else None


def _num(v):
    try:
        return round(float(v or 0), 2)
    except (TypeError, ValueError):
        return 0.0


def _txt(v):
    s = re.sub(r'\s+', ' ', str(v or '').replace('\xa0', ' ')).strip()
    return s or None


def buscar_ssw(dax_op, dax_cont, hoje=None):
    """455 desde INICIO (um trimestre por consulta: o executeQueries corta em 15 MiB e `dax_op` tem de
    levantar erro quando isso acontece) + 441 em aberto e pago."""
    hoje = hoje or hoje_brt()
    ce, a = [], date.fromisoformat(INICIO)
    fim = hoje + timedelta(days=1)
    while a < fim:
        b = min(date(a.year + (a.month + 2) // 12, (a.month + 2) % 12 + 1, 1), fim)
        ce += dax_op(dax_455(a, b))
        a = b
    # Manifesto emitido até ~1 mês antes do CT-e mais antigo (carga que o CT-e herda de manifesto anterior).
    # Sem ele a tela segue, só sem o nome do motorista.
    try:
        mfs = dax_op(dax_manifestos(date.fromisoformat(INICIO) - timedelta(days=31)))
        try:
            cadastro = dax_op(DAX_MOTORISTAS_047)
        except Exception:
            cadastro = []
        motoristas = nomes_motoristas(mfs, cadastro)
    except Exception as e:
        print(f'[contas-receber] manifestos fora (sem motorista): {e}', flush=True)
        motoristas = {}
    abertas = dax_cont(DAX_ABERTAS)
    ab_ctrcs = dax_cont(DAX_ABERTAS_CTRC)
    pagas = dax_cont(DAX_PAGAS_FAT)
    pg_ctrcs = dax_cont(DAX_PAGAS_CTRC)
    lido_455 = max((_iso(r.get('imp')) for r in ce if r.get('imp')), default=None)
    lido_abertas = max((str(r.get('data_importacao'))[:16] for r in abertas if r.get('data_importacao')), default=None)
    lido_pagas = max((str(r.get('imp'))[:16] for r in pagas if r.get('imp')), default=None)
    return {'ce': ce, 'abertas': abertas, 'ab_ctrcs': ab_ctrcs, 'pagas': pagas, 'pg_ctrcs': pg_ctrcs,
            'motoristas': motoristas, 'lido_455': lido_455, 'lido_abertas': lido_abertas, 'lido_pagas': lido_pagas,
            'lido_em': datetime.now(BRT).strftime('%d/%m/%Y %H:%M')}


_CACHE = {'v': None, 'ts': 0.0}
_CACHE_TRAVA = threading.Lock()
CACHE_S = 600


def dados_ssw(dax_op, dax_cont, forcar=False):
    """`buscar_ssw` com cache de 10 min (o 455 é recarregado 8×/dia; o 441 em aberto, 2×)."""
    with _CACHE_TRAVA:
        if not forcar and _CACHE['v'] is not None and time.time() - _CACHE['ts'] < CACHE_S:
            return _CACHE['v']
        v = buscar_ssw(dax_op, dax_cont)
        _CACHE.update(v=v, ts=time.time())
        return v


# ── Banco ────────────────────────────────────────────────────────────────────

DDL = """
CREATE TABLE IF NOT EXISTS cr_marcacao (
    tipo           TEXT NOT NULL,             -- 'cte' (chave = CTRC) ou 'fatura' (chave = nº da fatura)
    chave          TEXT NOT NULL,
    etapa          TEXT,                      -- só CT-e: PENDENCIA, EMITIDA, PDD (vazio = a faturar)
    campos         JSONB NOT NULL DEFAULT '{}'::jsonb,
    ssw            JSONB,                     -- cópia do que o SSW dizia quando a marcação nasceu
    origem         TEXT NOT NULL DEFAULT 'tela',
    criado_em      TIMESTAMPTZ NOT NULL DEFAULT now(),
    atualizado_em  TIMESTAMPTZ NOT NULL DEFAULT now(),
    atualizado_por TEXT,
    PRIMARY KEY (tipo, chave)
);
CREATE TABLE IF NOT EXISTS cr_log (
    id     BIGSERIAL PRIMARY KEY,
    tipo   TEXT NOT NULL,
    chave  TEXT NOT NULL,
    campo  TEXT NOT NULL,
    antes  TEXT,
    depois TEXT,
    por    TEXT,
    em     TIMESTAMPTZ NOT NULL DEFAULT now()
);
CREATE INDEX IF NOT EXISTS cr_log_chave ON cr_log (tipo, chave, em);
CREATE INDEX IF NOT EXISTS cr_log_em ON cr_log (em);
"""


def garantir_tabelas(cur):
    cur.execute(DDL)


def marcacoes(cur, tipo):
    cur.execute("""SELECT chave, etapa, campos, ssw, origem,
                          to_char(atualizado_em AT TIME ZONE 'America/Sao_Paulo', 'DD/MM/YYYY HH24:MI'),
                          (atualizado_em AT TIME ZONE 'America/Sao_Paulo')::date,
                          atualizado_por,
                          (SELECT (max(em) AT TIME ZONE 'America/Sao_Paulo')::date FROM cr_log l
                            WHERE l.tipo = m.tipo AND l.chave = m.chave AND l.campo = 'etapa')
                     FROM cr_marcacao m WHERE tipo = %s""", (tipo,))
    out = {}
    for chave, etapa, campos, ssw, origem, quando, dia, por, dia_etapa in cur.fetchall():
        out[chave] = {'chave': chave, 'etapa': etapa, 'campos': campos or {}, 'ssw': ssw, 'origem': origem,
                      'atualizado': quando, 'dia': dia.isoformat() if dia else None, 'por': por,
                      'dia_etapa': (dia_etapa or dia).isoformat() if (dia_etapa or dia) else None}
    return out


def gravar(cur, tipo, chaves, campo, valor, por, ssw_por_chave=None, antes_out=None):
    """Grava `campo` = `valor` (já normalizado) nas `chaves`, com log de cada mudança. Linha sem
    marcação ainda ganha a sua (com a cópia do SSW). Devolve quantas mudaram; `antes_out` recebe o valor
    anterior de cada uma (é o que o Desfazer da tela devolve)."""
    ssw_por_chave = ssw_por_chave or {}
    mudou = 0
    for chave in chaves:
        cur.execute('SELECT etapa, campos FROM cr_marcacao WHERE tipo = %s AND chave = %s FOR UPDATE', (tipo, chave))
        r = cur.fetchone()
        if r is None:
            s = ssw_por_chave.get(chave)
            cur.execute("""INSERT INTO cr_marcacao (tipo, chave, ssw, atualizado_por) VALUES (%s, %s, %s::jsonb, %s)
                           ON CONFLICT (tipo, chave) DO NOTHING""",
                        (tipo, chave, json.dumps(s, ensure_ascii=False, default=str) if s else None, por))
            cur.execute('SELECT etapa, campos FROM cr_marcacao WHERE tipo = %s AND chave = %s FOR UPDATE', (tipo, chave))
            r = cur.fetchone()
        etapa, campos = r
        campos = campos or {}
        antes = etapa if campo == 'etapa' else campos.get(campo)
        if (antes or None) == (valor or None):
            continue
        if campo == 'etapa':
            cur.execute("""UPDATE cr_marcacao SET etapa = %s, atualizado_em = now(), atualizado_por = %s
                           WHERE tipo = %s AND chave = %s""", (valor, por, tipo, chave))
        else:
            if valor is None:
                campos.pop(campo, None)
            else:
                campos[campo] = valor
            cur.execute("""UPDATE cr_marcacao SET campos = %s::jsonb, atualizado_em = now(), atualizado_por = %s
                           WHERE tipo = %s AND chave = %s""", (json.dumps(campos, ensure_ascii=False), por, tipo, chave))
        cur.execute('INSERT INTO cr_log (tipo, chave, campo, antes, depois, por) VALUES (%s, %s, %s, %s, %s, %s)',
                    (tipo, chave, campo, antes, valor, por))
        if antes_out is not None:
            antes_out[chave] = antes
        mudou += 1
    return mudou


def ultima_mudanca(cur, tipo, chaves, campo):
    """{chave: {'por', 'em'}} da última mudança deste campo — para o aviso de conflito."""
    if not chaves:
        return {}
    cur.execute("""SELECT DISTINCT ON (chave) chave, por, to_char(em AT TIME ZONE 'America/Sao_Paulo', 'DD/MM HH24:MI')
                     FROM cr_log WHERE tipo = %s AND chave = ANY(%s) AND campo = %s
                    ORDER BY chave, em DESC, id DESC""", (tipo, list(chaves), campo))
    return {k: {'por': p, 'em': e} for k, p, e in cur.fetchall()}


def historico(cur, tipo, chave):
    cur.execute("""SELECT campo, antes, depois, por, to_char(em AT TIME ZONE 'America/Sao_Paulo', 'DD/MM/YYYY HH24:MI')
                     FROM cr_log WHERE tipo = %s AND chave = %s ORDER BY em DESC, id DESC LIMIT 200""", (tipo, chave))
    return [{'campo': c, 'antes': a, 'depois': d, 'por': p, 'em': e} for c, a, d, p, e in cur.fetchall()]


def mudancas_desde(cur, desde, eu):
    """[(tipo, chave, por)] mudadas por OUTRA pessoa depois de `desde` (a importação conta)."""
    cur.execute("""SELECT tipo, chave, por FROM cr_log WHERE em > %s::timestamptz AND por IS DISTINCT FROM %s""",
                (desde, eu))
    return cur.fetchall()


# ── As linhas da tela ────────────────────────────────────────────────────────

def em_fatura(dados):
    """{CTRC: ('aberta'|'paga', fatura)} — onde cada CT-e está no 441. Aberta vence a paga (a fatura
    excluída e refeita fica com o número novo)."""
    out = {}
    for r in dados['pg_ctrcs']:
        k = chave_ctrc(r.get('ctrc'))
        if k:
            out[k] = ('paga', chave_fatura(r.get('fatura')))
    for r in dados['ab_ctrcs']:
        k = chave_ctrc(r.get('ctrc'))
        if k:
            out[k] = ('aberta', chave_fatura(r.get('fatura')))
    return out


def base_cte(r, motoristas=None):
    """Linha do 455 → o que a tela mostra (e o que se copia para a marcação). `motoristas` = {manifesto: nome}."""
    mf = _txt(r.get('mf'))
    return {'ctrc': chave_ctrc(r.get('ctrc')), 'cte': r.get('cte'), 'emissao': _iso(r.get('aut')),
            'emissao_455': _iso(r.get('emi')), 'doc': _txt(r.get('doc')), 'frete': _num(r.get('frete')),
            'cnpj': _txt(r.get('cnpj')), 'pagador': _txt(r.get('pagador')), 'unid': _txt(r.get('unid')),
            'nf': _txt(r.get('nf')), 'remetente': _txt(r.get('rem')), 'destinatario': _txt(r.get('dest')),
            'destino': '/'.join(x for x in (_txt(r.get('cid_d')), _txt(r.get('uf_d'))) if x) or None,
            'obs': _txt(r.get('obs')), 'baixa': _txt(r.get('baixa')),
            'cavalo': _txt(r.get('cav')), 'carreta': _txt(r.get('car')), 'manifesto': mf,
            'motorista': (motoristas or {}).get(chave_manifesto(mf)) if mf else None}


COPIA_CTE = ('ctrc', 'cte', 'emissao', 'doc', 'frete', 'cnpj', 'pagador', 'unid', 'nf', 'remetente',
             'destinatario', 'destino', 'obs', 'cavalo', 'carreta', 'manifesto', 'motorista')


def montar_ctes(dados, marc, crm_etapas=None, hoje=None):
    """A fila A faturar e o que está marcado → linhas da tela.

    `fila` de cada linha (é por ela que os cartões contam):
    - 'faturar'  — sem fatura e ninguém marcou (ou marcou e voltou);
    - 'pendencia'— sem fatura, com pendência marcada;
    - 'emitida'  — ela marcou Fatura emitida e o 441 ainda não mostra (até PRAZO_CONFIRMA dias);
    - 'pdd'      — não será faturado;
    - 'conferir' — algo não fecha: marcada emitida há mais de PRAZO_CONFIRMA dias e o 441 não confirma;
                   ou o 455 diz FATURADO e o 441 não tem a fatura (fatura excluída?).
    Marcada como Pendência ou PDD e que entrou numa fatura sai da tela: o SSW resolveu."""
    hoje = hoje or hoje_brt()
    crm_etapas = crm_etapas or {}
    fat = em_fatura(dados)
    ce = {}
    for r in dados['ce']:
        k = chave_ctrc(r.get('ctrc'))
        if k:
            ce[k] = r
    linhas = []
    for k in sorted(set(ce) | set(marc)):
        r, m = ce.get(k), marc.get(k)
        etapa = (m or {}).get('etapa')
        onde = fat.get(k)
        if r is not None:
            base = base_cte(r, dados.get('motoristas'))
            baixa = (base['baixa'] or '').upper()
        elif m and m.get('ssw'):
            base = dict({c: None for c in COPIA_CTE}, **m['ssw'], baixa=None, emissao_455=None)
            baixa = ''
        else:
            continue
        if onde:
            continue                                   # está em fatura: o SSW resolveu
        if r is not None and baixa not in ('PENDENTE', 'FATURADO'):
            continue                                   # liquidado sem fatura no 441: não é fila
        if r is None and etapa is None:
            continue
        conferir = None
        if etapa == 'EMITIDA':
            # O 455 lê 8×/dia e o 441 em aberto 2×: a baixa FATURADO chegar antes da fatura é normal aqui.
            dias = (hoje - date.fromisoformat(m['dia_etapa'])).days if m.get('dia_etapa') else 0
            if dias > PRAZO_CONFIRMA:
                conferir = (f'marcada como Fatura emitida há {dias} dias e o 441 (lido às 04h e às 15h) '
                            'ainda não tem o CT-e em fatura nenhuma')
            fila = 'conferir' if conferir else 'emitida'
        else:
            if baixa == 'FATURADO' and etapa != 'PDD':
                conferir = ('o 455 leu a baixa FATURADO, mas o 441 não tem o CT-e em fatura nenhuma — '
                            'fatura excluída no SSW?')
            fila = 'conferir' if conferir else {'PENDENCIA': 'pendencia', 'PDD': 'pdd'}.get(etapa, 'faturar')
        em = base.get('emissao')
        idade = (hoje - date.fromisoformat(em)).days if em else None
        crm = crm_etapas.get(k)
        linhas.append(dict(base, chave=k, etapa=etapa, campos=(m or {}).get('campos') or {}, fila=fila,
                           conferir=conferir, idade=idade, fora_455=r is None,
                           crm=crm, atualizado=(m or {}).get('atualizado'), por=(m or {}).get('por')))
    linhas.sort(key=lambda x: (x['emissao'] or '', x['chave']))
    return linhas


def _tel(ddd, tel):
    t = re.sub(r'\D', '', str(tel or ''))
    d = re.sub(r'\D', '', str(ddd or ''))
    if not t or set(t) == {'0'}:
        return None
    return f'({d}) {t}' if d else t


def montar_faturas(dados, marc, hoje=None):
    """Faturas em aberto do 441 (saldo ≥ SALDO_MIN) e as pagas nos últimos RECEBIDAS_DIAS → linhas.

    `grupo`: 'a_vencer', 'vencida' (vencida desde CORTE_ANTIGAS), 'antiga' (antes dele) ou 'recebida'."""
    hoje = hoje or hoje_brt()
    h = hoje.isoformat()
    # veículo e motorista de cada CT-e da fatura, pelo 455 (o 441 não traz)
    no_455 = {}
    for r in dados['ce']:
        k = chave_ctrc(r.get('ctrc'))
        if k:
            no_455[k] = r
    # A lista de CTRCs escreve a fatura sem o dígito ("0030426") e a fatura em aberto com ele ("0030426-0"):
    # casa pelo número sem o dígito, como o CRM - Documentos (até 30/10/2026 a lista nunca casava).
    sem_dv = lambda f: f.split('-')[0]
    ctrcs = {}
    for r in dados['ab_ctrcs']:
        f = chave_fatura(r.get('fatura'))
        f = sem_dv(f) if f else None
        k = chave_ctrc(r.get('ctrc'))
        if f and k:
            b = base_cte(no_455[k], dados.get('motoristas')) if k in no_455 else {}
            ctrcs.setdefault(f, []).append({'ctrc': k, 'cte': r.get('cte'), 'emissao': _iso(r.get('emi')),
                                            'frete': _num(r.get('frete')), 'cavalo': b.get('cavalo'),
                                            'carreta': b.get('carreta'), 'motorista': b.get('motorista')})
    linhas = []
    for a in dados['abertas']:
        f = chave_fatura(a.get('fatura'))
        saldo = _num(a.get('saldo'))
        if not f or saldo < SALDO_MIN:
            continue
        m = marc.get(f) or {}
        campos = m.get('campos') or {}
        venc = _iso(a.get('vencimen'))
        grupo = 'a_vencer' if not venc or venc >= h else ('vencida' if venc >= CORTE_ANTIGAS else 'antiga')
        atraso = (hoje - date.fromisoformat(venc)).days if venc and venc < h else 0
        prev = campos.get('previsao')
        cs = ctrcs.get(sem_dv(f), [])
        linhas.append({
            'chave': f, 'fatura': f, 'cliente': _txt(a.get('cliente')), 'cnpj': _txt(a.get('cnpj_cpf')),
            'emissao': _iso(a.get('emissao')), 'vencimento': venc, 'venc_original': _iso(a.get('venc_ori')),
            'valor': _num(a.get('vlr_fatur')), 'pago': _num(a.get('vlr_pago')), 'saldo': saldo,
            'desconto': _num(a.get('desconto')), 'parcial': _num(a.get('pgto_parcial')),
            'debito': _num(a.get('vlr_debito')), 'credito': _num(a.get('vlr_credito')),
            'motivo_deb_cred': _txt(a.get('motivo_debito_credito')),
            'promessa_ssw': _iso(a.get('data_promessa')), 'atraso': atraso,
            'cobranca': _txt(a.get('tipo_de_cobranca')), 'banco': _txt(a.get('banco_carteira')),
            'telefone': _tel(a.get('ddd'), a.get('telefone')),
            'telefone_cobranca': _tel(a.get('ddd_telefone_cobranca'), a.get('telefone_cobranca')),
            'email': _txt(a.get('e_mail_cobranca')), 'protesto': _txt(a.get('protestada_cartorio')),
            'spc': _txt(a.get('spc_serasa')), 'prorrogada': _txt(a.get('prorrogada')),
            'ocorrencia': _txt(a.get('ultima_ocorrencia')), 'unid': _txt(a.get('unidade_responsavel')),
            'cidade': '/'.join(x for x in (_txt(a.get('cidade')), _txt(a.get('uf'))) if x) or None,
            'vendedor': _txt(a.get('vendedor')),
            'ctrcs': cs[:60], 'n_ctrcs': len(cs), 'grupo': grupo, 'campos': campos,
            'previsao_vencida': bool(prev and prev <= h),
            'sem_tratativa': grupo != 'a_vencer' and not campos.get('situacao') and not prev,
            'atualizado': m.get('atualizado'), 'por': m.get('por')})
    desde = (hoje - timedelta(days=RECEBIDAS_DIAS)).isoformat()
    vistas = {x['chave'] for x in linhas}
    for p in dados['pagas']:
        f = chave_fatura(p.get('fatura'))
        pg = _iso(p.get('pagamento'))
        if not f or not pg or pg < desde or f in vistas:
            continue
        vistas.add(f)
        m = marc.get(f) or {}
        linhas.append({'chave': f, 'fatura': f, 'cliente': _txt(p.get('cliente')), 'cnpj': _txt(p.get('cnpj')),
                       'emissao': _iso(p.get('emissao')), 'vencimento': _iso(p.get('venc')), 'pago_em': pg,
                       'valor': _num(p.get('vlr_fatur')), 'pago': _num(p.get('vlr_pago')), 'saldo': 0.0,
                       'grupo': 'recebida', 'ctrcs': [], 'n_ctrcs': 0, 'campos': m.get('campos') or {},
                       'atraso': 0, 'previsao_vencida': False, 'sem_tratativa': False,
                       'atualizado': m.get('atualizado'), 'por': m.get('por')})
    linhas.sort(key=lambda x: (x['vencimento'] or '9999', x['chave']))
    return linhas


def copia_ssw(tipo, linha):
    """O que se guarda na marcação quando ela nasce: o bastante para a linha continuar dizendo o que
    era quando sair do SSW."""
    if tipo == 'cte':
        return {k: linha.get(k) for k in COPIA_CTE}
    return {k: linha.get(k) for k in ('fatura', 'cliente', 'cnpj', 'emissao', 'vencimento', 'valor', 'saldo')}


# ── Importação das planilhas da Laura ───────────────────────────────────────

def motivo_pdd(texto):
    """O motivo escrito por cima da observação (PDD.xlsx) → um dos MOTIVOS_PDD."""
    u = _sem_acento(str(texto or '')).upper()
    for motivo, pats in [('SUBSTITUÍDO', ['SUBSTITU']), ('RECUSADO / DESACORDO', ['RECUSAD', 'DESACORD', 'REJEITAD']),
                         ('COBRANÇA INDEVIDA / EMITIDO ERRADO', ['INDEVID', 'ERRAD', 'RPS']),
                         ('QUITAÇÃO NESTLÉ', ['QUITA']),
                         ('SEM CUSTO / NÃO SERÁ PAGO', ['NAO SERA PAGO', 'NAO TEM CUSTO', 'SEM CUSTO', 'SEM INFORMA']),
                         ('CONTRATO SINGULAR', ['SINGULAR']), ('SINISTRO / ACORDO', ['SINISTRO', 'ACORDO']),
                         ('CANCELADO', ['CANCEL'])]:
        if any(p in u for p in pats):
            return motivo
    return 'OUTRO'


def situacao_do_status(texto):
    """O STATUS livre da planilha de vencimento → uma das SITUACOES_FATURA (ou None)."""
    u = _sem_acento(str(texto or '')).upper()
    for sit, pats in [('DESCONTO DO CLIENTE', ['DESCONTO']), ('NEGOCIAÇÃO / JURÍDICO', ['NEGOCIA', 'RECUPERACAO JUDICIAL', 'JURIDIC']),
                      ('SEM PREVISÃO', ['SEM PREVISAO']), ('PREVISÃO DE PAGAMENTO', ['PREVISAO']),
                      ('PENDENTE PROGRAMAÇÃO', ['PENDENTE PROGRAMA']),
                      ('AGUARDANDO COMPROVANTE', ['COMP FISIC', 'COMPROVANTE']),
                      ('EM TRATATIVA', ['TRATATIVA', 'GLPI', 'CHAMADO', 'AGUARDANDO', 'PENDENTE', 'VALIDAR']),
                      ('COBRADO', ['COBRADO'])]:
        if any(p in u for p in pats):
            return sit
    return None


_RE_PREVISAO = re.compile(r'PREVIS[AÃ]O\s*(?:DIA\s*)?(\d{1,2}/\d{1,2}(?:/\d{2,4})?)', re.I)


def previsao_do_status(texto, hoje=None):
    m = _RE_PREVISAO.search(_sem_acento(str(texto or '')).upper())
    if not m:
        return None
    try:
        return ler_data(m[1], hoje)
    except ValueError:
        return None


def _abrir_xlsx(arq):
    import openpyxl
    return openpyxl.load_workbook(arq, data_only=True)


def ler_pdd(arq):
    """PDD.xlsx → [registro]: abas PDD e SINGULAR, uma linha por CTRC. A da SINGULAR leva o motivo
    Contrato Singular (é o que a aba separava)."""
    wb = _abrir_xlsx(arq)
    regs, vistos = [], set()
    for ws in wb.worksheets:
        cab = [_txt(c.value) or '' for c in ws[1]]
        ix = {c.upper(): i for i, c in enumerate(cab)}
        i_ctrc = ix.get('CTRC')
        if i_ctrc is None:
            continue
        g = lambda row, nome: row[ix[nome]] if nome in ix and ix[nome] < len(row) else None
        singular = 'SINGULAR' in ws.title.upper()
        for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            k = chave_ctrc(row[i_ctrc] if i_ctrc < len(row) else None)
            if not k or k in vistos:
                continue
            vistos.add(k)
            obs = _txt(g(row, 'OBSERVAÇÃO')) or _txt(g(row, 'OBSERVACAO'))
            emi = g(row, 'EMISSAO')
            frete = g(row, 'FRETE')
            regs.append({'tipo': 'cte', 'chave': k, 'etapa': 'PDD',
                         'campos': {'motivo': 'CONTRATO SINGULAR' if singular else motivo_pdd(obs),
                                    **({'anotacao': obs[:500]} if obs else {})},
                         'ssw': {'ctrc': k, 'cte': g(row, 'CT-E'), 'frete': frete if isinstance(frete, (int, float)) else None,
                                 'emissao': emi.date().isoformat() if isinstance(emi, datetime) else None,
                                 'cnpj': _txt(g(row, 'CNPJ PAGADOR')), 'pagador': _txt(g(row, 'CLIENTE PAGADOR')),
                                 'nf': _txt(g(row, 'NOTA FISCAL')), 'obs': obs},
                         'origem': f'{ws.title}, linha {n}'})
    return regs


def ler_vencimento(arq, hoje=None):
    """VENCIMENTO DIÁRIO.xlsx → [registro]: uma linha por fatura com STATUS. A aba vira Classificação
    (MOTORISTA, FATURAS SEM BAIXA); as abas de mês não classificam."""
    wb = _abrir_xlsx(arq)
    regs, vistos = [], set()
    for ws in wb.worksheets:
        cab = [(_txt(c.value) or '').upper() for c in ws[1]]
        if 'FATURA' not in cab:
            continue
        i_f, i_s = cab.index('FATURA'), (cab.index('STATUS') if 'STATUS' in cab else None)
        t = _sem_acento(ws.title).upper()
        classe = 'MOTORISTA' if 'MOTORISTA' in t else 'DESCONTO SEM BAIXA' if 'SEM BAIXA' in t else None
        for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            f = chave_fatura(row[i_f] if i_f < len(row) else None)
            if not f or f in vistos:
                continue
            vistos.add(f)
            st = _txt(row[i_s]) if i_s is not None and i_s < len(row) else None
            campos = {}
            if st:
                campos['anotacao'] = st[:500]
                sit = situacao_do_status(st)
                if sit:
                    campos['situacao'] = sit
                prev = previsao_do_status(st, hoje)
                if prev:
                    campos['previsao'] = prev
            if classe:
                campos['classe'] = classe
            if campos:
                regs.append({'tipo': 'fatura', 'chave': f, 'etapa': None, 'campos': campos, 'ssw': None,
                             'origem': f'{ws.title}, linha {n}'})
    return regs


def editadas_na_tela(cur, regs):
    """Chaves dos registros que alguém já mexeu na tela (a importação não passa por cima)."""
    out = set()
    for tipo in TIPOS:
        ks = [g['chave'] for g in regs if g['tipo'] == tipo]
        if ks:
            cur.execute("""SELECT DISTINCT chave FROM cr_log WHERE tipo = %s AND chave = ANY(%s)
                             AND campo <> '*importacao*'""", (tipo, ks))
            out |= {(tipo, k) for (k,) in cur.fetchall()}
    return out


def gravar_importacao(cur, regs, por):
    """Grava as marcações das planilhas. Linha já mexida na tela fica como está. → (gravadas, preservadas)."""
    pular = editadas_na_tela(cur, regs)
    n = 0
    for g in regs:
        if (g['tipo'], g['chave']) in pular:
            continue
        cur.execute("""INSERT INTO cr_marcacao (tipo, chave, etapa, campos, ssw, origem, atualizado_por)
                       VALUES (%s, %s, %s, %s::jsonb, %s::jsonb, 'planilha', %s)
                       ON CONFLICT (tipo, chave) DO UPDATE SET etapa = EXCLUDED.etapa, campos = EXCLUDED.campos,
                           ssw = COALESCE(cr_marcacao.ssw, EXCLUDED.ssw), origem = 'planilha',
                           atualizado_em = now(), atualizado_por = EXCLUDED.atualizado_por""",
                    (g['tipo'], g['chave'], g['etapa'], json.dumps(g['campos'], ensure_ascii=False),
                     json.dumps(g['ssw'], ensure_ascii=False, default=str) if g.get('ssw') else None, por))
        cur.execute('INSERT INTO cr_log (tipo, chave, campo, antes, depois, por) VALUES (%s, %s, %s, %s, %s, %s)',
                    (g['tipo'], g['chave'], '*importacao*', None,
                     json.dumps({'etapa': g['etapa'], **g['campos'], 'de': g['origem']}, ensure_ascii=False)[:2000], por))
        n += 1
    return n, len([g for g in regs if (g['tipo'], g['chave']) in pular])


def previa_importacao(regs, dados, marc_cte, marc_fat):
    """O que a importação faria, contra o SSW de agora — para ela conferir antes de gravar."""
    fat = em_fatura(dados)
    ce = {chave_ctrc(r.get('ctrc')) for r in dados['ce']}
    abertas = {chave_fatura(a.get('fatura')) for a in dados['abertas']}
    pdd = [g for g in regs if g['tipo'] == 'cte']
    fats = [g for g in regs if g['tipo'] == 'fatura']
    from collections import Counter
    return {
        'pdd': {'linhas': len(pdd), 'na_fila': sum(1 for g in pdd if g['chave'] in ce and g['chave'] not in fat),
                'ja_faturados': sum(1 for g in pdd if g['chave'] in fat),
                'antes_de_2026': sum(1 for g in pdd if g['chave'] not in ce and g['chave'] not in fat),
                'ja_marcados': sum(1 for g in pdd if g['chave'] in marc_cte),
                'motivos': Counter(g['campos']['motivo'] for g in pdd).most_common()},
        'faturas': {'linhas': len(fats), 'em_aberto': sum(1 for g in fats if g['chave'] in abertas),
                    'ja_pagas': sum(1 for g in fats if g['chave'] not in abertas),
                    'ja_marcadas': sum(1 for g in fats if g['chave'] in marc_fat),
                    'com_previsao': sum(1 for g in fats if g['campos'].get('previsao')),
                    'situacoes': Counter(g['campos'].get('situacao') or '(só anotação)' for g in fats).most_common()},
    }


# ── Excel ────────────────────────────────────────────────────────────────────

def _d(iso):
    try:
        return date.fromisoformat(iso) if iso else None
    except ValueError:
        return iso


ROT_FILA = {'faturar': 'A FATURAR', 'pendencia': 'PENDÊNCIA', 'emitida': 'FATURA EMITIDA (AGUARDANDO 441)',
            'pdd': 'NÃO SERÁ FATURADO', 'conferir': 'CONFERIR'}
ROT_GRUPO = {'a_vencer': 'A VENCER', 'vencida': 'VENCIDA', 'antiga': 'VENCIDA (ANTES DE MAI/26)', 'recebida': 'RECEBIDA'}


def exportar_xlsx(tipo, linhas):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    if tipo == 'cte':
        cols = [('Emissão', lambda l: _d(l['emissao'])), ('CTRC', lambda l: l['ctrc']), ('CT-e', lambda l: l['cte']),
                ('Unidade', lambda l: l.get('unid')), ('CNPJ Pagador', lambda l: l['cnpj']),
                ('Cliente Pagador', lambda l: l['pagador']), ('Documento', lambda l: l['doc']), ('NF', lambda l: l['nf']),
                ('Frete', lambda l: l['frete']), ('Idade (dias)', lambda l: l['idade']),
                ('Fila', lambda l: ROT_FILA.get(l['fila'], l['fila'])),
                ('Motivo', lambda l: l['campos'].get('motivo')), ('Anotação', lambda l: l['campos'].get('anotacao')),
                ('Etapa no CRM', lambda l: (l.get('crm') or {}).get('rotulo')),
                ('Conferir', lambda l: l.get('conferir')), ('Observação do CT-e', lambda l: l['obs'])]
    else:
        cols = [('Fatura', lambda l: l['fatura']), ('CNPJ/CPF', lambda l: l['cnpj']), ('Cliente', lambda l: l['cliente']),
                ('Emissão', lambda l: _d(l['emissao'])), ('Vencimento', lambda l: _d(l['vencimento'])),
                ('Dias de atraso', lambda l: l['atraso'] or None), ('Valor', lambda l: l['valor']),
                ('Pago', lambda l: l['pago']), ('Saldo', lambda l: l['saldo']),
                ('Grupo', lambda l: ROT_GRUPO.get(l['grupo'], l['grupo'])),
                ('Situação', lambda l: l['campos'].get('situacao')), ('Previsão', lambda l: _d(l['campos'].get('previsao'))),
                ('Classificação', lambda l: l['campos'].get('classe')), ('Anotação', lambda l: l['campos'].get('anotacao')),
                ('Telefone', lambda l: l.get('telefone_cobranca') or l.get('telefone')), ('E-mail', lambda l: l.get('email')),
                ('Unidade', lambda l: l.get('unid')), ('Última ocorrência SSW', lambda l: l.get('ocorrencia'))]
    wb = Workbook()
    ws = wb.active
    ws.title = 'A faturar' if tipo == 'cte' else 'Cobrança'
    azul = PatternFill('solid', fgColor='1F4E78')
    for j, (cab, _) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=cab)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = azul
    for i, l in enumerate(linhas, start=2):
        for j, (cab, f) in enumerate(cols, start=1):
            v = f(l)
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, date):
                c.number_format = 'DD/MM/YYYY'
            elif cab in ('Frete', 'Valor', 'Pago', 'Saldo'):
                c.number_format = '#,##0.00'
    for j, (cab, _) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = {'Cliente Pagador': 34, 'Cliente': 34, 'Anotação': 45,
                                                            'Observação do CT-e': 50, 'Conferir': 40}.get(cab, 14)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(1, len(linhas) + 1)}'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
