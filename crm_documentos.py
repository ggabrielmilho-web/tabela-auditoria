# -*- coding: utf-8 -*-
"""CRM - Documentos: acompanhamento dos CT-e dos clientes que pagam por portal.

O CRM controlava isto numa planilha ("Relatório de Documentos Nestle"): extraía o 455,
colava os CT-e e marcava à mão em que etapa do portal do cliente cada um estava (PQD, NBS,
RZA, AUTORIZADO, FATURADO...). Depois que o ciclo fecha, o faturamento segue na mesma
planilha. Esta aba faz a parte mecânica sozinha e deixa para elas só o que o SSW não sabe.

Duas fontes, e cada campo tem um dono:

- **Fato do SSW (só leitura):** o CT-e, o valor e a baixa vêm do 455
  (`conhecimentos_emitidos`); o pagamento vem do 441 (`faturas_441` + `faturas_441_ctrcs`).
- **Marcação delas (editável):** etapa, tipo de frete, canal, datas, anotações — em
  `crm_doc_marcacao`, com cada mudança em `crm_doc_log`. A leitura do SSW NUNCA sobrescreve
  uma marcação.

Por que o pagamento vem do 441 e não da baixa do 455 (medido em 20/10/2026): o robô do 455
pede o mês da AUTORIZAÇÃO do CT-e e regrava só esse período, então a baixa de um CT-e fica
como estava no fim do mês em que ele foi autorizado. Das 3.102 linhas da planilha que o 441
mostra pagas, o 455 ainda dizia PENDENTE em 2.141. O 441 é pedido pela data de LIQUIDAÇÃO,
do dia 1º/01 até ontem, e enxerga o pagamento de qualquer CT-e. A baixa do 455 só vale para
o que ele já tinha visto pago (LIQUIDADO não volta atrás).

Quando a marcação contradiz o SSW, as duas aparecem e a linha vai para "divergência" — a
linha fica onde ELA colocou. É o caso do CT-e substituído: o SSW o mostra a receber para
sempre (83 linhas, R$ 405 mil na planilha de 15/09/2026), e só a marcação dela sabe que ele
nunca vai ser pago.

Este módulo não fala Flask nem abre conexão: recebe as funções de consulta e o cursor.
"""

import hashlib
import io
import json
import re
import threading
import time
import unicodedata
from datetime import date, datetime, timedelta, timezone

# O servidor de produção roda em UTC; a tela é de Brasília (sem horário de verão desde 2019).
BRT = timezone(timedelta(hours=-3))


def hoje_brt():
    return datetime.now(BRT).date()


# ── Carteiras ────────────────────────────────────────────────────────────────
# Uma por aba da planilha; `desde` é o primeiro CT-e de cada aba (decisão de 20/10/2026: cada
# cliente começa onde a planilha dela começa; a aba Anterior, Nestlé 2023–24, não entra).
# A Garoto fica DENTRO da Nestlé, como na planilha delas (decisão de 20/10/2026) — sem mexer no
# GRUPOS_ECONOMICOS do server, que vale para o Faturamento por Tomador e a Análise por Veículo.
# Raízes: 60409075 Nestlé Brasil, 08334818 Nestlé Nordeste, 28053619 Chocolates Garoto. Até 22/10/2026
# a 08334818 estava aqui como "Garoto" e a Garoto de verdade ficava fora: os 40 CT-e dela não vinham
# do SSW e os da planilha apareciam como "sem CT-e no SSW" (UDI417858-1).
#
# Etapas: os códigos, a ordem e as cores são os da legenda e das listas de validação de cada
# aba da planilha. Cores: as da formatação condicional dela (RZA vermelho, NBS laranja, PQD
# amarelo, AUTORIZADO verde, FATURADO azul, DIGITADO rosa, SUBST roxo). O código é o da planilha
# (é o que se grava, importa e exporta); a tela mostra o `rotulo`, por extenso — pedido de
# 21/10/2026: "o que tiver muito abreviado, escrever mais completo" (SUBST era o "Substituído"
# que ela não achava). SUBMETIDO e REJEITADO entraram na Nestlé na mesma data (legenda dela).
#
# Campos: as colunas que elas preenchem, na ordem da planilha. `col` é o cabeçalho da
# planilha (importação e Excel). Tipos: `lista` (sugestões, aceita outro valor), `data_texto`
# (data quando dá para ler, senão o texto como foi escrito — a planilha tem "PQD 13/7" no meio
# das datas) e `texto`.

_E = lambda cod, desc, cor, rotulo: {'cod': cod, 'desc': desc, 'cor': cor, 'rotulo': rotulo}

TIPOS_FRETE = ['PRODUTO', 'PALLETS', 'DESCARGA', 'CARGA DESCARGA/PALLET', 'INDENIZADO',
               'INDENIZADO/PRODUTO BOM', 'ESTADIA DESTINO', 'ESTADIA', 'TROCA DE NOTA', 'PEDAGIO',
               'CHOCOLATES GAROTO', 'MAT PRIMA', 'RECALL', 'DEVOLUÇÃO', 'DEVOLUÇÃO PARCIAL',
               'DEVOLUÇÃO TOTAL', 'COMPLEMENTO', 'CARTA BÔNUS', 'NOTA DE DEBITO', 'ANULAÇÃO', 'CUSTO EXTRA',
               'DESLOCAMENTO VAZIO PARA REMONTAR PALLET', 'DIARIA', 'REENTREGA', 'ARMAZENAGEM',
               'REPALETIZAÇÃO', 'BEBIDAS', 'SUBSTITUTO']

CANAIS = ['MIDAS', 'B2B', 'GLL', 'PIPEFY', 'PORTAL OLIMPO', 'GALILEU', 'E-MAIL']

_ANOTACAO = {'id': 'anotacao', 'rotulo': 'Anotação', 'tipo': 'texto', 'col': None}

# Cancelado à mão (pedido de 29/10/2026): a linha da planilha sem CTRC não tem como o SSW dizer que foi
# cancelada — a Química escrevia "CANCELADO" no Vencimento. Fora do em aberto, como o SUBST; o cartão
# "Cancelado" da tela junta estas e as que o SSW cancelou (`ssw == 'cancelado'`).
ETAPA_CANCELADO = _E('CANCELADO', 'CT-e cancelado', 'chumbo', 'Cancelado')

# Custo extra da Nestlé (pedido de 26/10/2026): descarga, estadia, devolução, pallet e pedágio que ela
# cobra por ticket FSC — as colunas da planilha dela. Tudo manual: o complemento que a operação emite
# nem sempre é o do ticket dela; o SSW só sugere. Um CT-e pode ter mais de um (em 2025–26, 46 cargas
# com 2 complementos e 4 com 3), então o campo é uma LISTA, um item por custo. No pallet o valor já
# vem no frete: ela deixa o Valor em branco. Pedido de 27/10/2026: sai o Transporte (a carga B10 já é
# coluna da tabela) e entra o Pipefy (nº do ticket). Item que não está aqui cai ao regravar o custo —
# o Transporte de um custo antigo fica só no histórico (o log guarda a lista inteira).
TIPOS_CUSTO = ['DESCARGA', 'ESTADIA', 'DEVOLUÇÃO TOTAL', 'DEVOLUÇÃO PARCIAL', 'PALLET', 'PEDÁGIO']
STATUS_CUSTO = ['EM ANDAMENTO', 'CONCLUÍDO', 'NEGADO']
ITENS_CUSTO = [   # (id, rótulo = cabeçalho da planilha dela, tipo), na ordem da planilha
    ('status', 'Status', 'lista'), ('oc', 'OC', 'texto'),
    ('valor', 'Valor', 'moeda'), ('tipo', 'Tipo', 'lista'), ('ticket', 'Ticket', 'texto'),
    ('pipefy', 'Pipefy', 'texto'), ('doc_transporte', 'Doc transporte', 'texto'), ('pedido_compras', 'Pedido Compras', 'texto'),
    ('cte_compl', 'CT-e complementar', 'ctrc'),
]
CAMPO_CUSTOS = {'id': 'custos', 'rotulo': 'Custo extra', 'tipo': 'custos', 'col': None,
                'tipos': TIPOS_CUSTO, 'status': STATUS_CUSTO,
                'itens': [{'id': i, 'rotulo': r, 'tipo': t} for i, r, t in ITENS_CUSTO]}

CARTEIRAS = {
    'nestle': {
        'nome': 'Nestlé', 'raizes': ('60409075', '08334818', '28053619'), 'desde': '2025-01-02',
        'col_transporte': True,             # a carga B10 da observação, coluna depois do Frete (pedido de 26/10/2026)
        'aba': 'NESTLE 2025', 'col_etapa': 'STATUS', 'col_cte': 'CT-e/ NFS',
        'etapas': [
            _E('AG. COMPROVANTE', 'Aguardando retorno do comprovante de entrega', 'ciano', 'Aguardando comprovante'),
            _E('AG. AUTORIZAR', 'Aguardando autorizar', 'cinza', 'Aguardando autorizar'),
            _E('SUBMETIDO', 'Submetido no portal, aguardando aprovação da Nestlé', 'indigo', 'Submetido'),
            _E('REJEITADO', 'Rejeitado pela Nestlé no portal', 'magenta', 'Rejeitado'),
            _E('PIPEFY', 'Pendência a ser tratada pela Nestlé (Pipefy)', 'rosa', 'Pendência no Pipefy'),
            _E('NBS', 'Aguardando liberação Nestlé', 'laranja', 'Liberação Nestlé (NBS)'),
            _E('PQD', 'Parqueado: aguardando a Nestlé criar o pedido', 'amarelo', 'Parqueado (PQD)'),
            _E('RZA', 'Pendência aguardando ação Rizza', 'vermelho', 'Pendência Rizza (RZA)'),
            _E('AUTORIZADO', 'Aguardando antecipar', 'verde', 'Autorizado'),
            _E('FATURADO', 'Fatura emitida', 'azul', 'Faturado'),
            _E('SUBST', 'CT-e substituído por outro', 'roxo', 'Substituído'),
            ETAPA_CANCELADO,
        ],
        'campos': [
            {'id': 'tipo', 'rotulo': 'Tipo frete', 'tipo': 'lista', 'opcoes': TIPOS_FRETE, 'col': 'TIPO FRETE'},
            {'id': 'conferencia', 'rotulo': 'Ticket B2B/Transp', 'tipo': 'texto', 'col': 'TICKET B2B/TRANSP'},
            {'id': 'canal', 'rotulo': 'Canal', 'tipo': 'lista', 'opcoes': CANAIS, 'col': 'MIDAS/B2B/PIPEFY'},
            # no lugar da Antecipação (pedido de 21/10/2026): os AUTORIZADO/ANTECIPADO que vieram da coluna
            # MIDAS/B2B/PIPEFY (1.576) ficam no banco como `antecipacao`, sem coluna — a etapa já diz isso
            {'id': 'data_solicitacao', 'rotulo': 'Data solicitação', 'tipo': 'data_texto', 'col': None},
            {'id': 'ticket_pipefy', 'rotulo': 'Ticket Pipefy', 'tipo': 'texto', 'col': None},
            {'id': 'data_envio', 'rotulo': 'Data Pipefy', 'tipo': 'data_texto', 'col': 'DATA ENVIO/ data PQD'},
            {'id': 'autorizado_gll', 'rotulo': 'Autorizado GLL', 'tipo': 'data_texto', 'col': 'AUTORIZADOS GLL'},
            {'id': 'motorista', 'rotulo': 'Motorista', 'tipo': 'texto', 'col': 'MOTORISTA'},
            _ANOTACAO,
            CAMPO_CUSTOS,
        ],
    },
    'quimica_amparo': {
        'nome': 'Química Amparo', 'raizes': ('43461789',), 'desde': '2025-12-06',
        'col_vencimento': 'campo',          # manual: o que ela digita, nada do SSW
        'aba': 'QUIMICA AMPARO', 'col_etapa': 'STATUS', 'col_cte': 'CT-e',
        'etapas': [
            _E('AG. COMPROVANTE', 'Aguardando retorno do comprovante de entrega', 'ciano', 'Aguardando comprovante'),
            _E('DIGITADO', 'Digitado', 'rosa', 'Digitado'),
            _E('NBS', 'Aguardando liberação', 'laranja', 'Liberação (NBS)'),
            # na lista da aba desde a planilha de 20/10/2026 (2 linhas); a legenda dela diz "parqueado
            # aguardando criação de pedido", no lugar do PQD — o PQD segue na lista de validação
            _E('AG. LANCAMENTO', 'Parqueado: aguardando o lançamento do pedido', 'indigo', 'Aguardando lançamento'),
            _E('PQD', 'Parqueado: aguardando a criação do pedido', 'amarelo', 'Parqueado (PQD)'),
            _E('RZA', 'Pendência aguardando ação Rizza', 'vermelho', 'Pendência Rizza (RZA)'),
            _E('AUTORIZADO', 'Aguardando antecipar', 'verde', 'Autorizado'),
            _E('FATURADO', 'Fatura emitida', 'azul', 'Faturado'),
            _E('SUBST', 'CT-e substituído por outro', 'roxo', 'Substituído'),
            ETAPA_CANCELADO,
        ],
        'campos': [
            {'id': 'tipo', 'rotulo': 'Tipo frete', 'tipo': 'lista', 'opcoes': TIPOS_FRETE, 'col': 'TIPO FRETE'},
            # digitado por ela (1.429 de 1.692 em 21/10/2026: data, ou "PROGRAMADO"); é a coluna
            # Vencimento entre a Etapa e o SSW (`col_vencimento: 'campo'`, 28/10/2026)
            {'id': 'vencimento', 'rotulo': 'Vencimento', 'tipo': 'data_texto', 'col': 'VENCIMENTO'},
            # Situação, Programado, Pagamento e Arquivado saíram da tela a pedido dela (28/10/2026); o
            # que estava preenchido segue no banco (`campos`) e na linha original da planilha (gaveta)
            _ANOTACAO,
        ],
    },
    'unilever': {
        'nome': 'Unilever', 'raizes': ('61068276',), 'desde': '2025-08-31',
        'col_vencimento': 'ssw',            # o vencimento da fatura do 441 em que o CT-e está
        'aba': 'UNILEVER 2025', 'col_etapa': 'Farol', 'col_cte': 'CT-e',
        # a emissão está na 1ª coluna, que a tabela do Excel chama de "Colunas2"; a coluna S, sem nome
        # (14 números de 10 dígitos em 21/10/2026), vai para a anotação para não se perder
        'col_emissao': 'Colunas2', 'extras_anotacao': {'Colunas1': 'col. S da planilha'},
        'etapas': [
            _E('AG. COMPROVANTE', 'Aguardando retorno do comprovante de entrega', 'ciano', 'Aguardando comprovante'),
            _E('DIGITADO', 'Aguardando liberação Unilever', 'rosa', 'Digitado'),
            _E('UNI', 'Aguardando liberação', 'laranja', 'Liberação Unilever (UNI)'),
            _E('PIPEFY', 'Pendência a ser tratada pelo cliente', 'cinza', 'Pendência no Pipefy'),
            _E('RZA', 'Pendência aguardando ação Rizza', 'vermelho', 'Pendência Rizza (RZA)'),
            _E('AUTORIZADO', 'Aguardando antecipar', 'verde', 'Autorizado'),
            _E('FATURADO', 'Fatura emitida', 'azul', 'Faturado'),
            _E('SUBST', 'CT-e substituído por outro', 'roxo', 'Substituído'),
            ETAPA_CANCELADO,
        ],
        'campos': [
            {'id': 'tipo', 'rotulo': 'Tipo frete', 'tipo': 'lista', 'opcoes': TIPOS_FRETE, 'col': 'TIPO FRETE'},
            {'id': 'situacao', 'rotulo': 'Situação', 'tipo': 'texto', 'col': 'SITUAÇÃO'},
            {'id': 'multi', 'rotulo': 'Multi', 'tipo': 'data_texto', 'col': 'MULTI'},
            {'id': 'digitacao', 'rotulo': 'Dta digitação', 'tipo': 'data_texto', 'col': 'DTA DIGITAÇÃO'},
            {'id': 'confirming', 'rotulo': 'Confirming', 'tipo': 'texto', 'col': 'CONFIRMING'},
            {'id': 'antecipacao_santander', 'rotulo': 'Antecipação Santander', 'tipo': 'data_texto', 'col': 'ANTECIPAÇÃO SANTANDER'},
            _ANOTACAO,
        ],
    },
    'ambev': {
        'nome': 'Ambev', 'raizes': ('07526557',), 'desde': '2026-01-17',
        'aba': 'AMBEV', 'col_etapa': 'STATUS', 'col_cte': 'CT-e',
        # A planilha calcula o vencimento como emissão + 110 dias; pagamento sempre no dia 4.
        'prazo_dias': 110, 'col_vencimento': 'prazo',
        # Sem FATURADO automático: a fatura da Ambev sai junto com o CT-e (os 10 em aberto em
        # 20/10/2026 já estavam em fatura) e as etapas dela — programação, AG. PAGAMENTO, PGTO
        # AGENDADO — vêm depois; virar tudo FATURADO apagaria o que a aba acompanha.
        'auto_faturado': False,
        'extras_anotacao': {'Colunas1': 'col. R da planilha'},
        'etapas': [
            _E('PENDENTE', 'Aguardando programação de pagamento', 'amarelo', 'Pendente'),
            _E('AMBEV', 'Aguardando liberação', 'laranja', 'Liberação Ambev'),
            _E('RZA', 'Pendência aguardando ação Rizza', 'vermelho', 'Pendência Rizza (RZA)'),
            _E('AG. PAGAMENTO', 'Aguardando pagamento', 'verde', 'Aguardando pagamento'),
            _E('PGTO AGENDADO', 'Pagamento agendado (sempre dia 4)', 'verde', 'Pagamento agendado'),
            _E('FATURADO', 'Fatura emitida', 'azul', 'Faturado'),
            _E('SUBST', 'CT-e substituído por outro', 'roxo', 'Substituído'),
            ETAPA_CANCELADO,
        ],
        'campos': [
            {'id': 'tipo', 'rotulo': 'Tipo frete', 'tipo': 'lista', 'opcoes': TIPOS_FRETE, 'col': 'TIPO FRETE'},
            {'id': 'situacao', 'rotulo': 'Situação', 'tipo': 'texto', 'col': 'SITUAÇÃO'},
            {'id': 'programado', 'rotulo': 'Programado', 'tipo': 'data_texto', 'col': 'PROGRAMADO'},
            {'id': 'xml', 'rotulo': 'XML', 'tipo': 'texto', 'col': 'XML'},
            # a planilha calcula emissão + 110; onde ela digitou outra data (4 em 21/10/2026, o dia 4 do
            # pagamento), a data dela vale no lugar da conta — só esses vêm da importação
            {'id': 'vencimento', 'rotulo': 'Vencimento ajustado', 'tipo': 'data_texto', 'col': 'VENCIMENTO'},
            _ANOTACAO,
        ],
    },
}
ORDEM = ['nestle', 'quimica_amparo', 'unilever', 'ambev']

# Etapas que não estão "em aberto": o substituído e o cancelado nunca vão ser pagos.
ETAPAS_FORA_DO_ABERTO = {'SUBST', 'CANCELADO'}

# A 5ª aba: a descarga da NESTLÉ paga a terceiro (477, evento 5100) — a lógica mora em crm_descargas.py.
# Só Nestlé/Garoto, pelo tomador do CT-e (21/10/2026: o relatório de segunda ia com todos os clientes).
# É a conferência do relatório, nada mais (pedido de 26/10/2026): ela valida cada lançamento e o farol
# passa de vermelho para verde. A cobrança (ticket FSC, OC, complemento) é o Custo extra da aba Nestlé.
# As marcações ficam na mesma crm_doc_marcacao, carteira 'descargas', chave L477-<uni>-<lançamento>-<parcela>.
DESCARGAS = {
    'id': 'descargas', 'nome': 'Descargas Nestlé (5100)', 'modo': 'descargas', 'etapas': [],
    'campos': [
        {'id': 'validado', 'rotulo': 'Validação', 'tipo': 'farol', 'col': None},
        {'id': 'cte_manual', 'rotulo': 'CT-e (à mão)', 'tipo': 'texto', 'col': None},
        _ANOTACAO,
    ],
    # id → rótulo e cor; a ordem é a dos cartões. `aberto`: ainda falta ela validar.
    'situacoes': [
        {'id': 'pendente', 'rotulo': 'A validar', 'cor': 'vermelho', 'aberto': True,
         'desc': 'farol vermelho: ela ainda não validou o lançamento'},
        {'id': 'sem_cte', 'rotulo': 'Sem CT-e identificado', 'cor': 'cinza', 'aberto': True,
         'desc': 'o histórico cita a Nestlé mas não traz NF nem B10 que leve ao CT-e — preencha o CT-e à mão'},
        {'id': 'validado', 'rotulo': 'Validado', 'cor': 'verde', 'aberto': False,
         'desc': 'farol verde: validado por ela'},
    ],
    'prazo_dias': None, 'fora_do_aberto': [],
}
TODAS = list(ORDEM) + [DESCARGAS['id']]


def cfg_de(carteira):
    """Configuração de uma das abas (as 4 carteiras ou a de descargas); KeyError se não existir."""
    return CARTEIRAS[carteira] if carteira in CARTEIRAS else {DESCARGAS['id']: DESCARGAS}[carteira]


def carteira_da_raiz(raiz):
    for cid in ORDEM:
        if raiz in CARTEIRAS[cid]['raizes']:
            return cid
    return None


# O número da nota de serviço (UDS…): o SSW não traz (o CT-e vem vazio e nenhuma coluna do 455 tem a
# NFS), então ela digita — na própria coluna do CT-e, e só onde o SSW não preencheu. Não é coluna à
# parte: sem `col`, fora da lista de campos da carteira.
CAMPO_NFS = {'id': 'nfs', 'rotulo': 'NFS', 'tipo': 'texto', 'col': None}


def campo_def(carteira, campo_id):
    if campo_id == 'nfs' and carteira in CARTEIRAS:
        return CAMPO_NFS
    return next((c for c in cfg_de(carteira)['campos'] if c['id'] == campo_id), None)


def config_publica():
    """O que a tela precisa saber das abas (sem as raízes de CNPJ): as 4 carteiras e a de descargas."""
    return [{'id': cid, 'nome': CARTEIRAS[cid]['nome'], 'etapas': CARTEIRAS[cid]['etapas'],
             'campos': CARTEIRAS[cid]['campos'], 'prazo_dias': CARTEIRAS[cid].get('prazo_dias'),
             'transporte': bool(CARTEIRAS[cid].get('col_transporte')),
             # coluna Vencimento entre a Etapa e o SSW (28/10/2026): 'prazo' (Ambev, emissão + prazo ou o
             # ajustado), 'ssw' (Unilever, a fatura do 441), 'campo' (Química, só o que ela digita) ou
             # None (Nestlé: não usam)
             'col_vencimento': CARTEIRAS[cid].get('col_vencimento'),
             'fora_do_aberto': sorted(ETAPAS_FORA_DO_ABERTO)}
            for cid in ORDEM] + [DESCARGAS]


# ── Normalização ─────────────────────────────────────────────────────────────

def _sem_acento(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s or '') if unicodedata.category(c) != 'Mn')


def chave_ctrc(v):
    """'udi 123456-7 ' → 'UDI123456-7'; None se não tiver a forma de um CTRC do SSW."""
    s = re.sub(r'\s+', '', str(v or '')).upper()
    return s if re.fullmatch(r'[A-Z]{3}\d{6}-\d', s) else None


def normalizar_etapa(v, carteira):
    """O texto da coluna de status → código da etapa da carteira (ou None)."""
    s = _sem_acento(re.sub(r'\s+', ' ', str(v or '').strip().upper()))
    if not s:
        return None
    if s.startswith('SUBSTITU'):
        s = 'SUBST'
    return s


_TIPO_DE_PARA = {
    'PRODUTO1': 'PRODUTO', 'PRDUTO': 'PRODUTO',
    'INDENIZADOS': 'INDENIZADO', 'INDIZADO': 'INDENIZADO',
    'INDENIZADO/ PRODUTO BOM': 'INDENIZADO/PRODUTO BOM', 'INDE/PROD BOM': 'INDENIZADO/PRODUTO BOM',
    'PRODUTO BOM': 'INDENIZADO/PRODUTO BOM', 'INDENIZADO PRODUTO': 'INDENIZADO/PRODUTO BOM',
    'DEV PARCIAL': 'DEVOLUÇÃO PARCIAL', 'DEVOL PARCIAL': 'DEVOLUÇÃO PARCIAL', 'DEVL PARCIAL': 'DEVOLUÇÃO PARCIAL',
    'DEV TOTAL': 'DEVOLUÇÃO TOTAL', 'DEVL TOTAL': 'DEVOLUÇÃO TOTAL', 'DEV TOTAL DO FRETE': 'DEVOLUÇÃO TOTAL',
    'DEVOL': 'DEVOLUÇÃO', 'FRETE DEVOLUÇÃO': 'DEVOLUÇÃO', 'DEVOLUÇÃ': 'DEVOLUÇÃO',
    'COMPL': 'COMPLEMENTO', 'COMPLEMENTAR': 'COMPLEMENTO', 'COMPL DEVOL': 'COMPLEMENTO',
    'COMPL DEV': 'COMPLEMENTO', 'COMPL TROCA DE NOTA': 'COMPLEMENTO',
    'RECALL1': 'RECALL',
    'CHOCOLATE/GAROTO': 'CHOCOLATES GAROTO', 'TROCA DE NOTA CHOCOLATE GAROTO': 'TROCA DE NOTA',
    'FRETE DE PALLETS': 'PALLETS',
}


def normalizar_tipo(v):
    """Tipo de frete da planilha → (tipo da lista ou None, texto que vai para a anotação).

    A coluna é texto livre: 52 grafias na Nestlé ("prODUTO", "DEV PARCIAL" × "DEVOL
    PARCIAL"), e na Química ela virou campo de anotação ("e-mail enviado em 26/6"). O que é
    tipo vira o tipo; o que é recado vai para a anotação, sem perder nada."""
    bruto = re.sub(r'\s+', ' ', str(v or '').strip())
    if not bruto:
        return None, None
    s = bruto.upper()
    if s in TIPOS_FRETE:
        return s, None
    if s in _TIPO_DE_PARA:
        return _TIPO_DE_PARA[s], None
    if s.startswith('COMPL AO'):
        return 'COMPLEMENTO', bruto
    if s.startswith('SUBSTITUTO'):
        return 'SUBSTITUTO', (bruto if s != 'SUBSTITUTO' else None)
    return None, bruto


def separar_canal(v):
    """A coluna MIDAS/B2B/PIPEFY da Nestlé mistura três coisas: o canal, a situação da
    antecipação (AUTORIZADO/ANTECIPADO) e o número do ticket do Pipefy (10 dígitos). O que não
    é nenhum dos três (275 grafias, quase todas recado curto) fica no canal como ela escreveu."""
    if v is None or str(v).strip() == '':
        return {}
    if isinstance(v, (datetime, date)):
        return {'anotacao': f'[canal] {v:%d/%m/%Y}'}
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        v = str(int(v))
    t = re.sub(r'\s+', ' ', str(v).strip())
    u = _sem_acento(t).upper()
    if re.fullmatch(r'\d{10}', t):
        return {'ticket_pipefy': t}
    if u in ('AUTORIZADO', 'ANTECIPADO'):
        return {'antecipacao': u}
    if u in CANAIS:
        return {'canal': u}
    if u == 'ENVIADO MIDAS':
        return {'canal': 'MIDAS'}
    return {'canal': t}


def _excel_serial(n):
    """Número de série do Excel (45663 = 06/01/2025) → date, se estiver numa faixa plausível."""
    if 40000 <= n <= 50000:
        return date(1899, 12, 30) + timedelta(days=int(n))
    return None


def valor_campo(v, tipo):
    """Valor de célula da planilha → o que se grava no campo (sempre texto; data em ISO)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if tipo == 'data_texto':
        if isinstance(v, datetime):
            return v.date().isoformat()
        if isinstance(v, date):
            return v.isoformat()
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            d = _excel_serial(v)
            return d.isoformat() if d else str(int(v))
    if isinstance(v, datetime):
        return v.strftime('%d/%m/%Y')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return re.sub(r'[ \t]+', ' ', str(v).strip())


def ler_data_texto(s, hoje=None):
    """O que a pessoa digitou num campo de data → ISO quando dá para ler, senão o texto.

    Aceita 13/07, 13/7/26, 13/07/2026 e 2026-07-13. Sem ano: o ano corrente, ou o anterior se a
    data cairia mais de 60 dias no futuro (é registro do que aconteceu, não agenda)."""
    s = (s or '').strip()
    if not s:
        return None
    hoje = hoje or hoje_brt()
    m = re.fullmatch(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m:
        try:
            return date(int(m[1]), int(m[2]), int(m[3])).isoformat()
        except ValueError:
            return s
    m = re.fullmatch(r'(\d{1,2})/(\d{1,2})(?:/(\d{2}|\d{4}))?', s)
    if not m:
        return s
    d, mes, a = int(m[1]), int(m[2]), m[3]
    try:
        if a:
            ano = int(a) + (2000 if len(a) == 2 else 0)
            return date(ano, mes, d).isoformat()
        x = date(hoje.year, mes, d)
        if x > hoje + timedelta(days=60):
            x = date(hoje.year - 1, mes, d)
        return x.isoformat()
    except ValueError:
        return s


def _chave_opcao(s):
    return re.sub(r'[^A-Z0-9]', '', _sem_acento(str(s)).upper())


def ler_moeda(v):
    """'1.234,56', '1234.56', 'R$ 570' ou número → float com 2 casas (None se vazio); ValueError se não
    for número."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return round(float(v), 2)
    s = re.sub(r'[R$\s]', '', str(v or ''))
    if not s:
        return None
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    elif re.fullmatch(r'\d{1,3}(\.\d{3})+', s):      # 1.234 é mil duzentos e trinta e quatro
        s = s.replace('.', '')
    try:
        return round(float(s), 2)
    except ValueError:
        raise ValueError(f'Valor "{v}" não é um número — use 1.234,56.')


def normalizar_custos(valor):
    """A lista de custos extras de um CT-e (da tela: lista ou JSON) → lista limpa, ou None se vazia.
    Item sem nada preenchido sai; item com algo e sem Tipo é erro (os cartões contam pelo tipo)."""
    if isinstance(valor, str):
        if not valor.strip():
            return None
        try:
            valor = json.loads(valor)
        except ValueError:
            raise ValueError('Custo extra em formato inválido.')
    if not valor:
        return None
    if not isinstance(valor, list) or len(valor) > 10:
        raise ValueError('Custo extra: até 10 custos por CT-e.')
    out = []
    for i, it in enumerate(valor, start=1):
        if not isinstance(it, dict):
            raise ValueError('Custo extra em formato inválido.')
        x = {}
        for cid, rot, tp in ITENS_CUSTO:
            v = it.get(cid)
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            if tp == 'moeda':
                v = ler_moeda(v)
                if v is None:
                    continue
            elif tp == 'ctrc':
                k = chave_ctrc(v)
                if not k:
                    raise ValueError(f'Custo {i}: "{v}" não é um CTRC do SSW — use o formato UDI123456-7.')
                v = k
            else:
                v = re.sub(r'\s+', ' ', str(v).strip())
                if len(v) > 200:
                    raise ValueError(f'Custo {i}: {rot} com mais de 200 caracteres.')
                if tp == 'lista':
                    opcoes = TIPOS_CUSTO if cid == 'tipo' else STATUS_CUSTO
                    o = next((o for o in opcoes if _chave_opcao(o) == _chave_opcao(v)), None)
                    if o is None:
                        raise ValueError(f'Custo {i}: {rot} "{v}" não existe — use {", ".join(opcoes)}.')
                    v = o
            x[cid] = v
        if x and 'tipo' not in x:
            raise ValueError(f'Custo {i}: escolha o Tipo.')
        if x:
            out.append(x)
    return out or None


def normalizar_entrada(carteira, campo, valor, hoje=None):
    """Valida e normaliza uma edição da tela. Levanta ValueError com a mensagem para a pessoa."""
    cfg = cfg_de(carteira)
    if campo == 'custos' and campo_def(carteira, campo):
        return normalizar_custos(valor)
    v = (valor or '').strip() if isinstance(valor, str) else ('' if valor is None else str(valor))
    if campo == 'etapa':
        if not v:
            return None
        cods = [e['cod'] for e in cfg['etapas']]
        u = normalizar_etapa(v, carteira)
        if u not in cods:
            raise ValueError(f'Etapa "{v}" não existe em {cfg["nome"]}: use {", ".join(cods) or "nenhuma"}.')
        return u
    c = campo_def(carteira, campo)
    if not c:
        raise ValueError(f'Campo "{campo}" não existe em {cfg["nome"]}.')
    if not v:
        return None
    if campo == 'cte_manual':                       # o CT-e da carga, quando o histórico não leva a ele
        k = chave_ctrc(v)
        if not k:
            raise ValueError(f'"{v}" não é um CTRC do SSW — use o formato UDI123456-7.')
        return k
    if len(v) > 500:
        raise ValueError('Texto maior que 500 caracteres.')
    if c['tipo'] == 'data_texto':
        return ler_data_texto(v, hoje)
    if c['tipo'] == 'lista':
        # casa com a opção ignorando acento, caixa, espaço e pontuação: "email" e "e mail" viram E-MAIL,
        # "carga descarga pallet" vira CARGA DESCARGA/PALLET; o que não casa fica como foi escrito
        return next((o for o in c.get('opcoes', []) if _chave_opcao(o) == _chave_opcao(v)), v)
    if c['tipo'] == 'farol':                        # verde = validado por ela; vermelho = vazio
        u = _chave_opcao(v)
        if u in ('VALIDADO', 'VERDE', 'OK', 'SIM', 'APROVADO'):
            return 'VALIDADO'
        if u in ('PENDENTE', 'VERMELHO', 'NAO'):
            return None
        raise ValueError(f'"{v}": use VALIDADO (verde) ou PENDENTE (vermelho).')
    return v


def chave_manual(carteira, base):
    """Chave estável de uma linha da planilha sem CT-e no SSW (nota de serviço, cancelado)."""
    s = json.dumps([carteira, base.get('ctrc'), base.get('cte'), base.get('emissao'),
                    base.get('frete'), base.get('nf'), base.get('linha')], sort_keys=True, default=str)
    return 'MAN-' + hashlib.sha1(s.encode()).hexdigest()[:12]


# ── Fontes (SSW pelo Power BI) ───────────────────────────────────────────────

_T455 = "'public conhecimentos_emitidos'"
# A data do CT-e na tela é a da AUTORIZAÇÃO. A `data_emissao` do 455 difere dela em 760 de 8.859 CT-e
# das carteiras (22/10/2026), quase sempre DEPOIS (742; complemento e carga fechada; até ~50 dias); a
# planilha do CRM seguia a autorização (Ambev 100%). Sem autorização, vale a emissão.
_DATA_455 = f'COALESCE({_T455}[data_autorizacao], {_T455}[data_emissao])'
_COLS_455 = (f'"ctrc", {_T455}[serie_numero_ctrc], "cte", {_T455}[serie_numero_cte], "emissao", {_DATA_455}, '
             f'"emi_ssw", {_T455}[data_emissao], '
             f'"pag", {_T455}[cnpj_pagador], "pagador", {_T455}[cliente_pagador], "doc", {_T455}[tipo_documento], '
             f'"frete", {_T455}[valor_frete], "nf", {_T455}[numero_nota_fiscal], "uf_o", {_T455}[uf_origem_prestacao], '
             f'"rem", {_T455}[cliente_remetente], "dest", {_T455}[cliente_destinatario], '
             f'"cid_d", {_T455}[cidade_destinatario], "uf_d", {_T455}[uf_destinatario], '
             f'"baixa", {_T455}[tipo_baixa], "liq", {_T455}[data_liquidacao], "imp", {_T455}[data_importacao], '
             f'"obs", LEFT({_T455}[observacao], 300)')


def dax_455(raizes, ini, fim):
    lista = '{' + ','.join(f'"{r}"' for r in raizes) + '}'
    T = _T455
    return (f'EVALUATE SELECTCOLUMNS(FILTER({T}, LEFT({T}[cnpj_pagador], 8) IN {lista} && '
            f'{_DATA_455} >= DATE({ini.year},{ini.month},{ini.day}) && '
            f'{_DATA_455} < DATE({fim.year},{fim.month},{fim.day})), {_COLS_455})')


def dax_455_ctrcs(ctrcs):
    lista = '{' + ','.join(f'"{k}"' for k in ctrcs) + '}'
    return f'EVALUATE SELECTCOLUMNS(FILTER({_T455}, {_T455}[serie_numero_ctrc] IN {lista}), {_COLS_455})'


DAX_441_CTRC = ("EVALUATE SELECTCOLUMNS('public faturas_441_ctrcs', \"fatura\", [fatura], "
                "\"ctrc\", [ctrc], \"frete\", [valor_frete])")
DAX_441_FAT = ("EVALUATE SELECTCOLUMNS('public faturas_441', \"fatura\", [fatura], \"emissao\", [emissao], "
               "\"venc\", [vencimen], \"pagamento\", [pagamento], \"vlr_pago\", [vlr_pago], "
               "\"imp\", [data_importacao], \"ini\", [periodo_ini])")


# Carteira em aberto do 441 (Tipo C, por data de EMISSÃO, situação P), gravada pelo
# `ssw_relatorio 441.py --abertas` em faturas_441_abertas(_ctrcs). Enquanto as tabelas não
# estiverem no dataset da contábil, a consulta falha e a tela segue sem elas (como antes).
DAX_ABERTAS_FAT = ("EVALUATE SELECTCOLUMNS('public faturas_441_abertas', \"fatura\", [fatura], "
                   "\"emissao\", [emissao], \"venc\", [vencimen], \"vlr_fatur\", [vlr_fatur], "
                   "\"saldo\", [saldo], \"situacao\", [liquidada_atrasada], \"imp\", [data_importacao])")
DAX_ABERTAS_CTRC = ("EVALUATE SELECTCOLUMNS('public faturas_441_abertas_ctrcs', \"fatura\", [fatura], "
                    "\"ctrc\", [ctrc], \"frete\", [valor_frete], \"emi\", [data_emissao])")


def _iso(v):
    return str(v)[:10] if v else None


def _pagos_no_parcial(ctrcs, pago):
    """Pagamento parcial de fatura: o SSW registra só o valor ("PGTO PARCIAL - VALOR CTRCS
    APONTADOS: 13.800,00"), não quais CTRCs. Se UMA combinação de CTRCs da fatura soma o valor
    pago, são eles; se nenhuma ou mais de uma, não dá para saber (None).

    Medido em 20/10/2026: 49 das 372 faturas em aberto tinham pagamento parcial; a combinação
    é única em 14."""
    from itertools import combinations
    if pago <= 0.005 or not ctrcs or len(ctrcs) > 18:
        return None
    achadas = []
    for n in range(1, len(ctrcs) + 1):
        for comb in combinations(range(len(ctrcs)), n):
            if abs(sum(ctrcs[i][1] for i in comb) - pago) < 0.02:
                achadas.append(frozenset(ctrcs[i][0] for i in comb))
                if len(achadas) > 1:
                    return None
    return set(achadas[0]) if len(achadas) == 1 else None


def montar_abertas(fat_rows, ctrc_rows):
    """Linhas do BI → ({ctrc: [fatura em aberto, ...]}, corte), cada fatura com o que se sabe do parcial.

    `corte` = emissão do CTRC mais antigo que o 441 listou. O SSW só lista CTRC recente (na 1ª
    rodada, 20/10/2026: emitidos desde 09/06/2025); fatura com CTRC mais antigo vem sem a lista
    (as 317 em aberto de 2023–2024). Antes do corte, "não está em fatura em aberto" não prova nada."""
    fat = {}
    for f in fat_rows:
        k = str(f.get('fatura') or '').split('-')[0].strip()
        if k:
            vf, sal = float(f.get('vlr_fatur') or 0), float(f.get('saldo') or 0)
            fat[k] = {'fatura': f.get('fatura'), 'emissao': _iso(f.get('emissao')), 'vencimento': _iso(f.get('venc')),
                      'valor': vf, 'saldo': sal, 'pago': round(vf - sal, 2), 'situacao': f.get('situacao'),
                      'ctrcs': []}
    for r in ctrc_rows:
        f = fat.get(str(r.get('fatura') or '').strip())
        k = chave_ctrc(r.get('ctrc'))
        if f and k:
            f['ctrcs'].append((k, float(r.get('frete') or 0)))
    por_ctrc = {}
    for f in fat.values():
        f['pagos_parcial'] = sorted(_pagos_no_parcial(f['ctrcs'], f['pago']) or []) if f['pago'] > 0.005 else []
        f['parcial'] = f['pago'] > 0.005
        for k, _ in f['ctrcs']:
            por_ctrc.setdefault(k, []).append(f)
    corte = min((_iso(r.get('emi')) for r in ctrc_rows if r.get('emi')), default=None)
    return por_ctrc, corte


def buscar_abertas(dax_cont):
    """→ ({ctrc: [faturas]}, data da leitura, corte) ou (None, motivo, None) se não estiver no BI."""
    try:
        fats = dax_cont(DAX_ABERTAS_FAT)
        ctrcs = dax_cont(DAX_ABERTAS_CTRC)
    except Exception as e:
        return None, f'carteira em aberto do 441 indisponível no BI ({str(e)[:80]})', None
    lido = max((_iso(f.get('imp')) for f in fats if f.get('imp')), default=None)
    por_ctrc, corte = montar_abertas(fats, ctrcs)
    return por_ctrc, lido, corte


def buscar_ssw(dax_op, dax_cont, hoje=None):
    """455 das carteiras + 441 → {'ce': [...], 'pag': {ctrc: [...]}, 'lido_455', 'lido_441'}.

    O 455 vem um ano por vez: as consultas do executeQueries são cortadas em 15 MiB e quem as
    chama (`dax_op`) tem de levantar erro quando isso acontece — resposta cortada parece
    resposta."""
    hoje = hoje or hoje_brt()
    raizes = sorted({r for c in CARTEIRAS.values() for r in c['raizes']})
    ini = date.fromisoformat(inicio_455())
    ce = []
    a = ini
    while a <= hoje:
        b = min(date(a.year + 1, 1, 1), hoje + timedelta(days=1))
        ce += dax_op(dax_455(raizes, a, b))
        a = b
    fat = {str(f.get('fatura') or '').split('-')[0].strip(): f for f in dax_cont(DAX_441_FAT)}
    pag = {}
    for r in dax_cont(DAX_441_CTRC):
        k = chave_ctrc(r.get('ctrc'))
        f = fat.get(str(r.get('fatura') or '').strip())
        if k and f and f.get('pagamento'):
            pag.setdefault(k, []).append({'fatura': f.get('fatura'), 'pagamento': _iso(f.get('pagamento')),
                                          'vencimento': _iso(f.get('venc')), 'emissao': _iso(f.get('emissao')),
                                          'vlr_pago': f.get('vlr_pago')})
    lido_455 = max((_iso(r.get('imp')) for r in ce if r.get('imp')), default=None)
    lido_441 = max((_iso(f.get('imp')) for f in fat.values() if f.get('imp')), default=None)
    # Desde quando o 441 das pagas cobre (a carga apaga só dentro da janela do ano, então acumula).
    inicio_441 = min((_iso(f.get('ini')) for f in fat.values() if f.get('ini')), default=None)
    abertas, lido_ab, corte = buscar_abertas(dax_cont)
    return {'ce': ce, 'pag': pag, 'lido_455': lido_455, 'lido_441': lido_441, 'inicio_441': inicio_441,
            'abertas': abertas, 'lido_abertas': lido_ab if abertas is not None else None, 'abertas_corte': corte,
            'abertas_aviso': None if abertas is not None else lido_ab,
            'lido_em': datetime.now(BRT).strftime('%d/%m/%Y %H:%M')}


def inicio_455():
    """Primeiro dia que a tela lê do 455: o início da carteira mais antiga (02/01/2025, que é também
    onde o 455 do dataset principal começa)."""
    return min(c['desde'] for c in CARTEIRAS.values())


def ctrcs_do_historico(marc, ce_chaves):
    """CTRCs marcados (vindos da planilha) fora do 455 lido — os CT-e de 2023–2024 que ficaram na aba
    da Nestlé. O 455 do dataset principal começa em `inicio_455`; o do DRE vai até 2021. Não se filtra
    pela emissão da planilha: ela pode estar errada (uma nota de serviço de 26/12/2024 estava como
    16/01/2025). O que for de 2025 em diante e não estiver no 455 lido (nota cancelada) volta None."""
    return sorted(k for k in marc if chave_ctrc(k) == k and k not in ce_chaves)


def completar_historico(dados, marc, dax_hist):
    """Busca no 455 histórico (`dax_hist`: o dataset do DRE) os `ctrcs_do_historico` que esta leitura
    do cache ainda não buscou, e guarda em dados['ce_historico'] = {ctrc: linha do 455 ou None}. Some
    junto com o cache (10 min). Erro no BI deixa as linhas como estavam — manuais — até a próxima."""
    if 'ce_chaves' not in dados:
        dados['ce_chaves'] = {chave_ctrc(r.get('ctrc')) for r in dados['ce']}
    hist = dados.setdefault('ce_historico', {})
    faltam = [k for k in ctrcs_do_historico(marc, dados['ce_chaves']) if k not in hist]
    for i in range(0, len(faltam), 300):
        lote = faltam[i:i + 300]
        try:
            achou = {chave_ctrc(r.get('ctrc')): r for r in dax_hist(dax_455_ctrcs(lote))}
        except Exception:
            achou = {}
        for k in lote:
            hist[k] = achou.get(k)
    return hist


_CACHE = {'v': None, 'ts': 0.0}
_CACHE_TRAVA = threading.Lock()
CACHE_S = 600


def dados_ssw(dax_op, dax_cont, forcar=False):
    """`buscar_ssw` com cache de 10 min (o 455 é recarregado 8×/dia; o 441, 1×/dia)."""
    with _CACHE_TRAVA:
        if not forcar and _CACHE['v'] is not None and time.time() - _CACHE['ts'] < CACHE_S:
            return _CACHE['v']
        v = buscar_ssw(dax_op, dax_cont)
        _CACHE.update(v=v, ts=time.time())
        return v


# ── Banco ────────────────────────────────────────────────────────────────────

DDL = """
CREATE TABLE IF NOT EXISTS crm_doc_marcacao (
    chave          TEXT PRIMARY KEY,          -- CTRC do SSW (UDI123456-7) ou MAN-<hash> (linha sem CT-e)
    carteira       TEXT NOT NULL,
    ctrc           TEXT,
    cte            BIGINT,
    etapa          TEXT,
    campos         JSONB NOT NULL DEFAULT '{}'::jsonb,
    planilha       JSONB,                     -- a linha como veio da planilha (nunca muda depois)
    origem         TEXT NOT NULL DEFAULT 'tela',
    criado_em      TIMESTAMPTZ NOT NULL DEFAULT now(),
    atualizado_em  TIMESTAMPTZ NOT NULL DEFAULT now(),
    atualizado_por TEXT
);
CREATE INDEX IF NOT EXISTS crm_doc_marcacao_carteira ON crm_doc_marcacao (carteira);
CREATE TABLE IF NOT EXISTS crm_doc_log (
    id     BIGSERIAL PRIMARY KEY,
    chave  TEXT NOT NULL,
    campo  TEXT NOT NULL,
    antes  TEXT,
    depois TEXT,
    por    TEXT,
    em     TIMESTAMPTZ NOT NULL DEFAULT now()
);
CREATE INDEX IF NOT EXISTS crm_doc_log_chave ON crm_doc_log (chave, em);
-- Mudança feita pela tela a partir do SSW (FATURADO automático), não por uma pessoa. Só altera se
-- falta a coluna: o ALTER, mesmo com IF NOT EXISTS, pede trava exclusiva e espera qualquer
-- transação aberta na tabela (travou o teste local em 21/10/2026).
DO $$ BEGIN
    IF NOT EXISTS (SELECT 1 FROM information_schema.columns WHERE table_schema = current_schema()
                     AND table_name = 'crm_doc_log' AND column_name = 'auto') THEN
        ALTER TABLE crm_doc_log ADD COLUMN auto BOOLEAN NOT NULL DEFAULT false;
    END IF;
END $$;
-- Cópia do que o 455 dizia do CT-e (emissão, valor, pagador…) quando a marcação nasceu da tela: o 455
-- não traz CT-e cancelado, e sem ela a linha cancelada ficava só com o CTRC (29/10/2026).
DO $$ BEGIN
    IF NOT EXISTS (SELECT 1 FROM information_schema.columns WHERE table_schema = current_schema()
                     AND table_name = 'crm_doc_marcacao' AND column_name = 'ssw') THEN
        ALTER TABLE crm_doc_marcacao ADD COLUMN ssw JSONB;
    END IF;
END $$;
"""


def garantir_tabelas(cur):
    cur.execute(DDL)


def marcacoes(cur, carteira):
    cur.execute("""SELECT chave, ctrc, cte, etapa, campos, planilha, origem, ssw,
                          to_char(atualizado_em AT TIME ZONE 'America/Sao_Paulo', 'DD/MM/YYYY HH24:MI'),
                          atualizado_por,
                          EXISTS (SELECT 1 FROM crm_doc_log l WHERE l.chave = m.chave
                                     AND l.campo <> '*importacao*' AND NOT l.auto),
                          COALESCE((SELECT max(em) FROM crm_doc_log l WHERE l.chave = m.chave
                                       AND l.campo = 'etapa' AND NOT l.auto)
                                 > (SELECT max(em) FROM crm_doc_log l WHERE l.chave = m.chave
                                       AND l.campo = 'etapa' AND l.auto), false)
                     FROM crm_doc_marcacao m WHERE carteira = %s""", (carteira,))
    out = {}
    for (chave, ctrc, cte, etapa, campos, planilha, origem, ssw, quando, por, editado, mao) in cur.fetchall():
        out[chave] = {'chave': chave, 'ctrc': ctrc, 'cte': cte, 'etapa': etapa, 'campos': campos or {},
                      'planilha': planilha, 'origem': origem, 'ssw': ssw, 'atualizado': quando, 'por': por,
                      'editado': editado, 'mao_depois_do_auto': mao}
    return out


# O que se copia do 455 para a marcação (coluna `ssw`): o bastante para a linha cancelada no SSW
# continuar dizendo o que era.
CAMPOS_COPIA_SSW = ('emissao', 'emissao_455', 'frete', 'nf', 'uf_o', 'doc', 'pagador', 'cnpj',
                    'remetente', 'destinatario', 'destino')


def base_ssw(linha):
    """A linha da tela → o que `gravar` recebe em `ssw_por_chave` (CTRC, CT-e e a cópia do 455)."""
    return {k: linha.get(k) for k in ('ctrc', 'cte') + CAMPOS_COPIA_SSW}


def copias_faltando(linhas, marc):
    """Marcações nascidas da tela, com o CT-e ainda no 455, que não têm a cópia → [(chave, cópia)].
    Preenche uma vez as que já existiam antes da coluna `ssw` (29/10/2026)."""
    out = []
    for l in linhas:
        m = marc.get(l['chave'])
        if m and not l['manual'] and m.get('planilha') is None and m.get('ssw') is None:
            out.append((l['chave'], {k: l.get(k) for k in CAMPOS_COPIA_SSW}))
    return out


def gravar_copias(cur, itens):
    """Grava a cópia do 455 (sem log e sem mexer em `atualizado_em`: não é mudança de ninguém)."""
    for chave, copia in itens:
        cur.execute('UPDATE crm_doc_marcacao SET ssw = %s::jsonb WHERE chave = %s AND ssw IS NULL',
                    (json.dumps(copia, ensure_ascii=False, default=str), chave))
    return len(itens)


def gravar(cur, carteira, chaves, campo, valor, por, ssw_por_chave=None, auto=False, antes_out=None):
    """Grava `campo` = `valor` (já normalizado) nas `chaves`, com log de cada mudança.
    Chave sem marcação ainda (CT-e "novo") ganha a linha. Devolve quantas mudaram.
    `auto`: mudança da tela a partir do SSW, não de uma pessoa (fica marcada no log).
    `antes_out`, se dado, recebe {chave: valor anterior} das que mudaram — é o que o Desfazer da
    tela devolve (o valor do banco, não o que a tela tinha na memória)."""
    ssw_por_chave = ssw_por_chave or {}
    mudou = 0
    for chave in chaves:
        cur.execute('SELECT etapa, campos, carteira FROM crm_doc_marcacao WHERE chave = %s FOR UPDATE', (chave,))
        r = cur.fetchone()
        if r is None:
            # Linha sem marcação ainda. Duas pessoas nela no mesmo instante: o ON CONFLICT faz a segunda
            # esperar a primeira e seguir com a linha que ela criou (antes, a segunda dava erro de chave).
            s = ssw_por_chave.get(chave) or {}
            cte = s.get('cte')
            copia = {k: s[k] for k in CAMPOS_COPIA_SSW if k in s} or None
            cur.execute("""INSERT INTO crm_doc_marcacao (chave, carteira, ctrc, cte, ssw, atualizado_por)
                           VALUES (%s, %s, %s, %s, %s::jsonb, %s) ON CONFLICT (chave) DO NOTHING""",
                        (chave, carteira, s.get('ctrc'), int(cte) if str(cte or '').isdigit() else None,
                         json.dumps(copia, ensure_ascii=False, default=str) if copia else None, por))
            cur.execute('SELECT etapa, campos, carteira FROM crm_doc_marcacao WHERE chave = %s FOR UPDATE', (chave,))
            r = cur.fetchone()
        etapa, campos, cart = r
        campos = campos or {}
        if cart != carteira:
            raise ValueError(f'{chave} é de outra carteira.')
        antes = etapa if campo == 'etapa' else campos.get(campo)
        if (antes or None) == (valor or None):
            continue
        if campo == 'etapa':
            cur.execute("""UPDATE crm_doc_marcacao SET etapa = %s, atualizado_em = now(), atualizado_por = %s
                           WHERE chave = %s""", (valor, por, chave))
        else:
            if valor is None:
                campos.pop(campo, None)
            else:
                campos[campo] = valor
            cur.execute("""UPDATE crm_doc_marcacao SET campos = %s::jsonb, atualizado_em = now(), atualizado_por = %s
                           WHERE chave = %s""", (json.dumps(campos, ensure_ascii=False), por, chave))
        txt = lambda x: json.dumps(x, ensure_ascii=False) if isinstance(x, (list, dict)) else x   # custos: lista
        cur.execute('INSERT INTO crm_doc_log (chave, campo, antes, depois, por, auto) VALUES (%s, %s, %s, %s, %s, %s)',
                    (chave, campo, txt(antes), txt(valor), por, auto))
        if antes_out is not None:
            antes_out[chave] = antes
        mudou += 1
    return mudou


# ── FATURADO pelo SSW ────────────────────────────────────────────────────────

AUTO_POR = 'Automático (SSW)'


def _br(iso):
    return f'{iso[8:10]}/{iso[5:7]}/{iso[:4]}' if iso else '?'


def faturados_pelo_ssw(carteira, linhas, marc):
    """Linhas que o SSW prova faturadas e ainda não estão em FATURADO → [(linha, motivo)].

    Prova de fatura: o CT-e está numa fatura em aberto ou paga do 441, ou o 455 leu a baixa
    FATURADO (que só existe com a fatura emitida). LIQUIDADO do 455 sozinho não prova fatura.
    Fica de fora: SUBST e CANCELADO (em fatura é divergência para alguém olhar), linha que
    alguém tirou do FATURADO à mão depois da mudança automática (a mão vale; a mudança não
    volta) e carteira com `auto_faturado` False."""
    if not CARTEIRAS[carteira].get('auto_faturado', True):
        return []
    out = []
    for l in linhas:
        if l['etapa'] in ('FATURADO', 'SUBST', 'CANCELADO') or (marc.get(l['chave']) or {}).get('mao_depois_do_auto'):
            continue
        fa = l.get('fatura_aberta') or {}
        if l['ssw'] == 'fatura' and l.get('fonte') == '441 abertas':
            motivo = f"fatura {fa.get('fatura')} em aberto no 441"
        elif l['ssw'] == 'fatura' and l.get('fonte') == '455':
            motivo = f"baixa FATURADO no 455 (lida em {_br(l.get('lido_455'))})"
        elif l['ssw'] == 'pago' and l.get('fonte') == '441':
            motivo = f"fatura {', '.join(l.get('faturas') or [])} paga em {_br(l.get('pago_em'))} (441)"
        elif l['ssw'] == 'pago' and l.get('fonte') == '441 parcial':
            motivo = f"pagamento parcial da fatura {fa.get('fatura')} (441)"
        else:
            continue
        out.append((l, motivo))
    return out


def gravar_faturados(cur, carteira, itens):
    """Grava FATURADO nas linhas de `faturados_pelo_ssw`, com a prova no log. Devolve quantas mudaram."""
    n = 0
    for l, motivo in itens:
        n += gravar(cur, carteira, [l['chave']], 'etapa', 'FATURADO', f'{AUTO_POR}: {motivo}',
                    {l['chave']: base_ssw(l)}, auto=True)
    return n


# ── Fatura excluída: o FATURADO automático volta atrás ──────────────────────
# Pedido de 28/10/2026: para antecipar, ela emite a fatura no SSW e depois a EXCLUI. O automático
# via a fatura em aberto, gravava FATURADO e ficava assim (26/10: 18 CT-e da Nestlé, faturas
# 0039887–0039897, que somem do 441 inteiro — a exclusão não deixa rastro). Buracos na numeração de
# 2026: 576 faturas em 140 blocos, 9 a 19 de uma vez a cada 2–3 semanas.

def candidatos_desfaturar(linhas):
    """Linhas em FATURADO que o SSW mostra sem fatura nenhuma — só estas podem voltar."""
    return [l for l in linhas if l['etapa'] == 'FATURADO' and l['ssw'] == 'aberto']


def ultimas_etapas(cur, chaves):
    """{chave: última mudança de etapa} — a importação da planilha conta, porque também grava etapa."""
    if not chaves:
        return {}
    cur.execute("""SELECT DISTINCT ON (chave) chave, campo, por, antes, depois, auto,
                          (em AT TIME ZONE 'America/Sao_Paulo')::date
                     FROM crm_doc_log WHERE chave = ANY(%s) AND campo IN ('etapa', '*importacao*')
                    ORDER BY chave, em DESC, id DESC""", (list(chaves),))
    return {k: {'campo': c, 'por': p, 'antes': a, 'depois': d, 'auto': au, 'dia': dia.isoformat()}
            for k, c, p, a, d, au, dia in cur.fetchall()}


_RE_FATURA_AUTO = re.compile(r'fatura (\S+) em aberto')


def desfaturados_pelo_ssw(carteira, linhas, ultimas, lido_abertas):
    """FATURADO posto pelo automático cuja fatura sumiu do SSW → [(linha, etapa de antes, motivo)].

    Volta para a etapa que a linha tinha antes do automático (`antes` do log), só quando:
    - a última mudança de etapa foi o FATURADO automático — se uma pessoa ou a importação mexeu
      depois, a etapa é dela;
    - o SSW mostra o CT-e fora de qualquer fatura (`aberto`: a carteira em aberto do 441 cobre o
      CT-e e ele não está em fatura aberta nem paga);
    - a foto da carteira em aberto (`lido_abertas`) é de um dia DEPOIS do automático: a leitura que
      diz "sumiu" é mais nova que a que provou a fatura. Sem foto, nada volta.
    A volta é automática (não conta como mão): se a fatura de verdade sair depois, o FATURADO volta."""
    if not CARTEIRAS[carteira].get('auto_faturado', True) or not lido_abertas:
        return []
    out = []
    for l in candidatos_desfaturar(linhas):
        u = ultimas.get(l['chave'])
        if not u or u['campo'] != 'etapa' or not u['auto'] or u['depois'] != 'FATURADO' or u['dia'] >= lido_abertas:
            continue
        m = _RE_FATURA_AUTO.search(u['por'] or '')
        de = f"fatura {m[1]}" if m else 'a fatura'
        antes = u['antes'] or None
        out.append((l, antes, f"{de} excluída no SSW (o 441 não tem mais o CT-e em fatura nenhuma) — "
                              f"{'volta para ' + antes if antes else 'volta sem etapa'}"))
    return out


def gravar_desfaturados(cur, carteira, itens):
    """Volta as linhas de `desfaturados_pelo_ssw` para a etapa de antes, com o motivo no log."""
    n = 0
    for l, antes, motivo in itens:
        n += gravar(cur, carteira, [l['chave']], 'etapa', antes, f'{AUTO_POR}: {motivo}', auto=True)
    return n


def ultima_mudanca(cur, chaves, campo):
    """{chave: {'por', 'em'}} da última mudança deste campo nestas linhas — para o aviso de conflito."""
    if not chaves:
        return {}
    cur.execute("""SELECT DISTINCT ON (chave) chave, por,
                          to_char(em AT TIME ZONE 'America/Sao_Paulo', 'DD/MM HH24:MI')
                     FROM crm_doc_log WHERE chave = ANY(%s) AND campo = %s
                    ORDER BY chave, em DESC, id DESC""", (list(chaves), campo))
    return {k: {'por': por, 'em': em} for k, por, em in cur.fetchall()}


def historico(cur, chave):
    cur.execute("""SELECT campo, antes, depois, por,
                          to_char(em AT TIME ZONE 'America/Sao_Paulo', 'DD/MM/YYYY HH24:MI')
                     FROM crm_doc_log WHERE chave = %s ORDER BY em DESC, id DESC LIMIT 200""", (chave,))
    return [{'campo': c, 'antes': a, 'depois': d, 'por': p, 'em': e} for c, a, d, p, e in cur.fetchall()]


# ── A linha da tela ──────────────────────────────────────────────────────────

def _resumo_fatura(f):
    return {k: f[k] for k in ('fatura', 'emissao', 'vencimento', 'valor', 'saldo', 'pago', 'parcial',
                              'pagos_parcial', 'situacao')} | {'ctrcs': f['ctrcs'][:40], 'n_ctrcs': len(f['ctrcs'])}


def situacao_ssw(ce, pagamentos, abertas=None, hoje=None, ctrc=None, cobre=False, inicio_441=None):
    """O que o SSW diz do recebimento — só o que dá para provar:

    1. `pago` pelo 441 da contábil (fatura liquidada);
    2. fatura EM ABERTO pelo 441 --abertas: número, vencimento, vencida — ou `pago` quando um
       pagamento parcial da fatura aponta este CTRC (o valor identifica os CTRCs pagos);
    3. `pago` pela baixa LIQUIDADO do 455 (pagamento não "despaga");
    4. `aberto` (sem fatura) só quando a carteira em aberto COBRE o CT-e (`cobre`: emitido depois
       do CTRC mais antigo que o 441 lista) e ele não está em fatura nenhuma;
    5. `fatura` pela baixa FATURADO do 455 só se a leitura é posterior ao início do 441 das pagas:
       pagamento antes da leitura teria virado LIQUIDADO, e depois dela está no 441 — então é
       fatura emitida e não paga (todas as leituras do 455 são de 2026, e o 441 cobre desde
       01/01/2026, conferido em 20/10/2026);
    6. senão `indefinido`. A baixa PENDENTE do 455 não prova nada: o CT-e pode ter sido faturado
       depois da leitura (ela fica congelada no mês da autorização)."""
    if pagamentos:
        ult = max(pagamentos, key=lambda p: p['pagamento'])
        return {'ssw': 'pago', 'pago_em': ult['pagamento'], 'fonte': '441',
                'faturas': sorted({p['fatura'] for p in pagamentos}), 'venc_fatura': ult.get('vencimento')}
    if abertas:
        hoje = hoje or hoje_brt()
        parcial = [f for f in abertas if ctrc in f['pagos_parcial']]
        if parcial:
            f = parcial[0]
            return {'ssw': 'pago', 'pago_em': None, 'fonte': '441 parcial', 'faturas': [f['fatura']],
                    'fatura_aberta': _resumo_fatura(f), 'venc_fatura': f['vencimento']}
        f = min(abertas, key=lambda x: x['vencimento'] or '9999-12-31')
        return {'ssw': 'fatura', 'fonte': '441 abertas', 'fatura_aberta': _resumo_fatura(f),
                'vencida': bool(f['vencimento'] and f['vencimento'] < hoje.isoformat()), 'venc_fatura': f['vencimento']}
    if ce is None:
        return {'ssw': 'sem_ssw'}
    b = str(ce.get('baixa') or '').upper()
    if b == 'LIQUIDADO':
        return {'ssw': 'pago', 'pago_em': _iso(ce.get('liq')), 'fonte': '455', 'faturas': []}
    if cobre:
        return {'ssw': 'aberto', 'fonte': '441 abertas'}
    lido = _iso(ce.get('imp'))
    if b == 'FATURADO' and inicio_441 and lido and lido >= inicio_441:
        return {'ssw': 'fatura', 'fonte': '455'}
    return {'ssw': 'indefinido', 'fonte': '455'}


# Complemento da Nestlé: "CTRC EMITIDO P/ COMPL. DO UDI123456-7 … TRANSP-12345678 PO-7200000000 - MOTIVO:
# estadia". Em 2025–26, 687 de 689 citam o CTRC da carga; o motivo é descarga, estadia ou complementar
# (pallet, pedágio…); TRANSP e PO vêm em ~45%. É só sugestão para o Custo extra — ela confirma.
_RE_COMPL_DO = re.compile(r'COMPL\w*\.?\s+DO\s+(?:CTRC\s+|CT-?E\s+)?([A-Z]{3})\s?(\d{6})-?(\d)', re.I)
_RE_MOTIVO = re.compile(r'MOTIVO\s*:?\s*([A-Za-zÀ-ÿ]+)', re.I)
_RE_TRANSP = re.compile(r'TRANSP\s*-?\s*(\d{6,})', re.I)
_RE_PO = re.compile(r'\bPO\s*-?\s*(\d{6,})', re.I)


# A carga da Nestlé (transporte, shipment): B10 + 7 dígitos na observação do CT-e ("B101234567 - PEDIDO…").
# Em 2025–26, 91% das linhas da Nestlé (98% das cargas fechadas; devolução e subcontrato não têm). Aceita
# o B10 digitado com um dígito a menos ou a mais e o número que a vírgula do BI partiu ao meio.
_RE_TRANSPORTE = re.compile(r'\bB\s?(10\d{6,8})\b', re.I)


def transporte_da_obs(obs):
    """Observação do CT-e → 'B101234567' (ou 'B10…, B10…' com mais de uma carga), ou None."""
    s = str(obs or '').upper()
    achou = _RE_TRANSPORTE.findall(s) or _RE_TRANSPORTE.findall(s.replace(',', ''))
    return ', '.join(dict.fromkeys('B' + n for n in achou)) or None


def complementos_ssw(ce_rows):
    """{CTRC da carga: [complementos do SSW que o citam]} — a sugestão do Custo extra."""
    out = {}
    for r in ce_rows:
        if 'COMPLEMENT' not in str(r.get('doc') or '').upper():
            continue
        obs = re.sub(r'\s+', ' ', str(r.get('obs') or '').replace(',', ''))   # o BI quebra a observação com vírgulas
        m = _RE_COMPL_DO.search(obs)
        k = chave_ctrc(r.get('ctrc'))
        if not m or not k:
            continue
        mot = _RE_MOTIVO.search(obs)
        mot = mot[1].lower() if mot else None
        u = _sem_acento(obs).upper()
        tipo = ('DESCARGA' if mot == 'descarga' else 'ESTADIA' if mot == 'estadia'
                else 'PEDÁGIO' if 'PEDAGIO' in u else 'PALLET' if re.search(r'PALL?ET|PALETE', u) else None)
        t, po = _RE_TRANSP.search(obs), _RE_PO.search(obs)
        out.setdefault(f'{m[1].upper()}{m[2]}-{m[3]}', []).append({
            'ctrc': k, 'emissao': _iso(r.get('emissao')), 'frete': r.get('frete'), 'motivo': mot,
            'tipo': tipo, 'transporte': t[1] if t else None, 'po': po[1] if po else None})
    for v in out.values():
        v.sort(key=lambda x: x['emissao'] or '')
    return out


# CT-e substituto: "CTRC EMITIDO PARA SUBSTITUIR O CTRC UDI 404842-3,CT-e 002 …". O original é
# cancelado e sai do 455; o substituto fica, e é por ele que a linha cancelada aponta para onde foi.
_RE_SUBSTITUI = re.compile(r'SUBSTITUIR\s+O\s+CTRC\s+([A-Z]{3})\s?(\d{6})-?(\d)', re.I)


def substitutos_ssw(ce_rows):
    """{CTRC original: CTRC do substituto} pelas observações do 455."""
    out = {}
    for r in ce_rows:
        obs = str(r.get('obs') or '')
        if 'SUBSTITU' not in obs.upper():
            continue
        m = _RE_SUBSTITUI.search(re.sub(r'\s+', ' ', obs.replace(',', ' ')))
        k = chave_ctrc(r.get('ctrc'))
        if m and k:
            out.setdefault(f'{m[1].upper()}{m[2]}-{m[3]}', k)
    return out


def divergencias(etapa, sit):
    """Marcação × SSW, só com prova do 441. A linha fica onde ela marcou; a divergência só aparece.

    Nada aqui sai da baixa PENDENTE/FATURADO do 455: congelada no mês da autorização, ela
    acusava de "sem fatura" CT-e que a Química já tinha faturado e até recebido (medido em
    20/10/2026) e de "a receber" substituído que ninguém mais cobraria."""
    out = []
    fa = sit.get('fatura_aberta')
    if etapa == 'SUBST' and sit.get('ssw') == 'fatura' and fa:
        out.append({'cod': 'subst_em_fatura',
                    'texto': f"Substituído aqui, mas está na fatura {fa['fatura']} em aberto no SSW."})
    if etapa == 'CANCELADO' and sit.get('ssw') == 'fatura' and fa:
        out.append({'cod': 'cancelado_em_fatura',
                    'texto': f"Cancelado aqui, mas está na fatura {fa['fatura']} em aberto no SSW."})
    if etapa == 'FATURADO' and sit.get('ssw') == 'aberto':
        out.append({'cod': 'faturado_sem_fatura',
                    'texto': 'FATURADO aqui, e o 441 não tem este CT-e em fatura nenhuma (nem paga, nem em aberto) — conferir.'})
    return out


def montar_linhas(carteira, ce_rows, pagamentos, marc, hoje=None, abertas=None, abertas_corte=None,
                  inicio_441=None, ce_historico=None):
    """CT-e do SSW da carteira (desde o início dela) + as marcações → linhas da tela.
    `abertas` = {ctrc: [faturas em aberto]} do 441 --abertas, ou None se não estiver no BI;
    `abertas_corte` = emissão do CTRC mais antigo que o 441 lista (antes dele a carteira não prova
    falta de fatura); `ce_historico` = {ctrc: linha} do 455 do DRE (`completar_historico`)."""
    hoje = hoje or hoje_brt()
    abertas_ok = abertas is not None
    cfg = CARTEIRAS[carteira]
    desde = cfg['desde']
    compl = complementos_ssw(ce_rows) if any(c['tipo'] == 'custos' for c in cfg['campos']) else {}
    subst = substitutos_ssw(ce_rows)
    por_chave, antes = {}, {}
    for r in ce_rows:
        k = chave_ctrc(r.get('ctrc'))
        if not k or str(r.get('pag') or '')[:8] not in cfg['raizes']:
            continue
        (antes if (_iso(r.get('emissao')) or '') < desde else por_chave)[k] = r
    # Linha marcada com emissão antes do início da carteira: o CT-e vem do 455 mesmo assim — do lido
    # (antes do `desde`) ou do histórico do DRE. O início vale só para o CT-e que ninguém marcou.
    do_historico = set()
    for k in marc:
        if k in por_chave:
            continue
        r = antes.get(k)
        if r is None and (ce_historico or {}).get(k):
            r = ce_historico[k]
            if str(r.get('pag') or '')[:8] in cfg['raizes']:
                do_historico.add(k)
            else:
                r = None
        if r is not None:
            por_chave[k] = r
    linhas = []
    for k in sorted(set(por_chave) | set(marc)):
        ce, m = por_chave.get(k), marc.get(k)
        p = (m or {}).get('planilha') or {}
        if ce is not None:
            base = {'ctrc': k, 'cte': ce.get('cte'), 'emissao': _iso(ce.get('emissao')),
                    'emissao_455': _iso(ce.get('emi_ssw')), 'uf_o': ce.get('uf_o'), 'nf': ce.get('nf'), 'frete': ce.get('frete'),
                    'doc': ce.get('doc'), 'pagador': ce.get('pagador'), 'cnpj': ce.get('pag'),
                    'remetente': ce.get('rem'), 'destinatario': ce.get('dest'),
                    'destino': '/'.join(x for x in (ce.get('cid_d'), ce.get('uf_d')) if x),
                    'obs': ce.get('obs'), 'lido_455': _iso(ce.get('imp'))}
        else:
            c = {} if p else ((m or {}).get('ssw') or {})   # nascida da tela: a cópia do que o 455 dizia
            base = {'ctrc': (m or {}).get('ctrc') or p.get('ctrc'), 'cte': (m or {}).get('cte') or p.get('cte'),
                    'emissao': p.get('emissao') or c.get('emissao'), 'emissao_455': c.get('emissao_455'),
                    'uf_o': p.get('uf_o') or c.get('uf_o'), 'nf': p.get('nf') or c.get('nf'),
                    'frete': p.get('frete') or c.get('frete'), 'doc': c.get('doc'),
                    'pagador': p.get('cliente') or c.get('pagador'), 'cnpj': p.get('cnpj') or c.get('cnpj'),
                    'remetente': c.get('remetente'), 'destinatario': c.get('destinatario'),
                    'destino': c.get('destino'), 'obs': p.get('obs'), 'lido_455': None}
        # Linha manual (fora do 455) com CTRC de verdade — nota de serviço de descarga, CT-e
        # antigo — ainda pode ter sido paga numa fatura do 441.
        ck = k if chave_ctrc(k) else chave_ctrc(base['ctrc'])
        etapa = (m or {}).get('etapa')
        em = base['emissao']
        cobre = bool(abertas_ok and em and (abertas_corte is None or em >= abertas_corte))
        s = situacao_ssw(ce, pagamentos.get(ck) if ck else None,
                         (abertas or {}).get(ck) if ck else None, hoje, ck, cobre, inicio_441)
        # Nasceu do SSW (marcada na tela, sem linha da planilha) e saiu do 455: o relatório não traz CT-e
        # cancelado (29/10/2026). Só quando o 441 também não o conhece — fatura ou pagamento provam que existe.
        if ce is None and not p and chave_ctrc(k) == k and s['ssw'] == 'sem_ssw':
            s = {'ssw': 'cancelado', 'fonte': '455', 'substituto': subst.get(k)}
        idade = (hoje - date.fromisoformat(em)).days if em else None
        venc = None
        if cfg.get('prazo_dias') and em:
            venc = (date.fromisoformat(em) + timedelta(days=cfg['prazo_dias'])).isoformat()
            ajuste = ((m or {}).get('campos') or {}).get('vencimento')
            if ajuste and re.fullmatch(r'\d{4}-\d{2}-\d{2}', str(ajuste)):
                venc = ajuste                       # data digitada por ela vale no lugar da conta
        # A receber: o SSW não mostra pagamento e ela não o deu por encerrado. Onde o SSW não tem
        # resposta confiável (CT-e que ele não conhece, ou situação indefinida), vale a etapa dela,
        # e FATURADO encerra.
        aberto = (s['ssw'] not in ('pago', 'cancelado') and etapa not in ETAPAS_FORA_DO_ABERTO
                  and not (s['ssw'] in ('sem_ssw', 'indefinido') and etapa == 'FATURADO'))
        # NFS: só onde o SSW não deu o número do CT-e. O que ela digitou na tela vale; senão, o que estava
        # na coluna do CT-e da planilha (na Nestlé ela se chamava "CT-e/ NFS": 124 das 130 UDS).
        campos_m = (m or {}).get('campos') or {}
        nfs_editavel = not base['cte']
        nfs = None
        if nfs_editavel:
            nfs = campos_m.get('nfs') or ((p.get('original') or {}).get(cfg['col_cte']))
            nfs = str(nfs).strip() if nfs not in (None, '') else None
        linha = dict(base, chave=k, etapa=etapa, campos=campos_m, nfs=nfs, nfs_editavel=nfs_editavel,
                     manual=ce is None, nota_linha=p.get('nota_linha') if ce is None else None,
                     do_historico=k in do_historico, compl_ssw=compl.get(k) or [],
                     transporte=transporte_da_obs(base['obs']) if cfg.get('col_transporte') else None,
                     atualizado=(m or {}).get('atualizado'), por=(m or {}).get('por'),
                     editado=bool((m or {}).get('editado')), idade=idade, vencimento=venc,
                     aberto=aberto, novo=etapa is None and s['ssw'] not in ('pago', 'cancelado'),
                     abertas_cobre=cobre, diverg=divergencias(etapa, s), **s)
        linha.setdefault('vencida', False)
        linha.setdefault('fatura_aberta', None)
        linhas.append(linha)
    linhas.sort(key=lambda x: (x['emissao'] or '', x['chave']), reverse=True)
    return linhas


# ── Importação da planilha ───────────────────────────────────────────────────

def _cab(s):
    return re.sub(r'\s+', ' ', str(s or '').strip())


def ler_planilha(caminho):
    """→ {carteira: [linha]}, cada linha com os campos base e as colunas dela, cruas."""
    import openpyxl
    from openpyxl.utils import range_boundaries
    wb = openpyxl.load_workbook(caminho, data_only=True)
    out = {}
    for cid in ORDEM:
        cfg = CARTEIRAS[cid]
        if cfg['aba'] not in wb.sheetnames:
            continue
        ws = wb[cfg['aba']]
        tabs = list(ws.tables.values())
        if not tabs:
            continue
        c1, r1, c2, r2 = range_boundaries(tabs[0].ref)
        cab = [_cab(ws.cell(r1, c).value) for c in range(c1, c2 + 1)]
        linhas = []
        # Lê também o que foi digitado ABAIXO da tabela do Excel (ela nem sempre cresce sozinha: na
        # Unilever, em 21/10/2026, 24 CT-e de set–out estavam fora dela) — só linha com CTRC ou CT-e,
        # para não trazer total ou rascunho.
        for r in range(r1 + 1, ws.max_row + 1):
            vals = {cab[i]: ws.cell(r, c1 + i).value for i in range(len(cab))}
            if all(v in (None, '') for v in vals.values()):
                continue
            # Sem CTRC e sem CT-e não é documento: abaixo da tabela, nunca; dentro dela, só se tiver a
            # emissão. A linha de total da Ambev (frete R$ 93 mil, vencimento e mês por fórmula) entrava
            # como CT-e manual a receber (21/10/2026).
            tem_doc = vals.get('CTRC') not in (None, '') or vals.get(cfg['col_cte']) not in (None, '')
            if not tem_doc and (r > r2 or vals.get(cfg.get('col_emissao', 'Emissao')) in (None, '')):
                continue
            vals['_linha'] = r
            linhas.append(vals)
        out[cid] = linhas
    return out


def preparar_importacao(planilha, ce_rows):
    """Planilha lida + 455 → (registros a gravar, relatório). Não toca no banco.

    Casamento pela chave do SSW: o CTRC; sem ele (a Química entra sem CTRC desde 03/08/2026),
    o número do CT-e. O que não casa (nota de serviço de descarga, CT-e cancelado) entra como
    linha manual, com os dados da própria planilha."""
    # O nº do CT-e sozinho não identifica: as filiais CAR e NOD da Unilever usam a mesma série 3
    # começando do 1. Número que aponta para mais de um CTRC não serve de chave.
    por_cte = {}
    ctrcs = set()
    for r in ce_rows:
        k = chave_ctrc(r.get('ctrc'))
        if k:
            ctrcs.add(k)
            if r.get('cte'):
                por_cte.setdefault(int(r['cte']), set()).add(k)
    por_cte = {n: next(iter(ks)) for n, ks in por_cte.items() if len(ks) == 1}
    regs, rel = [], {}
    for cid, linhas in planilha.items():
        cfg = CARTEIRAS[cid]
        st = rel.setdefault(cid, {'linhas': 0, 'pelo_ctrc': 0, 'pelo_cte': 0, 'manual': 0, 'repetidas': 0,
                                  'etapas': {}, 'tipos_para_anotacao': {}, 'etapa_desconhecida': {}})
        vistas = set()
        cods = {e['cod'] for e in cfg['etapas']}
        for v in linhas:
            st['linhas'] += 1
            ctrc_txt = v.get('CTRC')
            cte_v = v.get(cfg['col_cte'])
            ctrc = chave_ctrc(ctrc_txt)
            cte = int(cte_v) if isinstance(cte_v, (int, float)) and not isinstance(cte_v, bool) else None
            if ctrc is None and isinstance(ctrc_txt, (int, float)) and not isinstance(ctrc_txt, bool):
                cte = cte or int(ctrc_txt)          # CT-e digitado na coluna do CTRC
            emissao = v.get(cfg.get('col_emissao', 'Emissao'))
            base = {'linha': v['_linha'], 'aba': cfg['aba'],
                    'emissao': emissao.date().isoformat() if isinstance(emissao, datetime) else None,
                    'ctrc': ctrc or (str(ctrc_txt).strip() if ctrc_txt else None), 'cte': cte,
                    'cnpj': valor_campo(v.get('CNPJ Pagador'), 'texto'),
                    'cliente': valor_campo(v.get('Cliente Pagador'), 'texto'),
                    'uf_o': valor_campo(v.get('UF O'), 'texto'), 'nf': valor_campo(v.get('NF'), 'texto'),
                    'frete': v.get('Frete') if isinstance(v.get('Frete'), (int, float)) else None,
                    'obs': valor_campo(v.get('Observacao'), 'texto')}
            chave = None
            if ctrc and ctrc in ctrcs:
                chave, como = ctrc, 'pelo_ctrc'
            elif cte and cte in por_cte:
                chave, como = por_cte[cte], 'pelo_cte'
            elif ctrc:
                chave, como = ctrc, 'manual'       # CTRC de verdade, mas fora do 455 (NFS, cancelado, antigo)
            if chave is None or chave in vistas:
                if chave in vistas:
                    st['repetidas'] += 1
                    base['nota_linha'] = f'linha repetida na planilha ({chave})'
                chave, como = chave_manual(cid, base), 'manual'
            if como == 'manual' and not base.get('nota_linha'):
                base['nota_linha'] = 'sem CT-e no SSW'
            vistas.add(chave)
            st[como] += 1
            etapa = normalizar_etapa(v.get(cfg['col_etapa']), cid)
            if etapa and etapa not in cods:
                st['etapa_desconhecida'][etapa] = st['etapa_desconhecida'].get(etapa, 0) + 1
            st['etapas'][etapa or '(vazio)'] = st['etapas'].get(etapa or '(vazio)', 0) + 1
            campos, notas = {}, []
            for c in cfg['campos']:
                if not c['col']:
                    continue
                bruto = v.get(c['col'])
                if c['id'] == 'tipo':
                    tipo, sobra = normalizar_tipo(bruto)
                    if tipo:
                        campos['tipo'] = tipo
                    if sobra:
                        notas.append(f'[tipo frete] {sobra}')
                        st['tipos_para_anotacao'][sobra.upper()] = st['tipos_para_anotacao'].get(sobra.upper(), 0) + 1
                elif c['id'] == 'canal':
                    for k2, v2 in separar_canal(bruto).items():
                        if k2 == 'anotacao':
                            notas.append(v2)
                        else:
                            campos[k2] = v2
                else:
                    x = valor_campo(bruto, c['tipo'])
                    if (c['id'] == 'vencimento' and cfg.get('prazo_dias') and base['emissao'] and x ==
                            (date.fromisoformat(base['emissao']) + timedelta(days=cfg['prazo_dias'])).isoformat()):
                        x = None                    # é a conta da própria planilha, não um ajuste
                    if x is not None:
                        campos[c['id']] = x
            for col, rot in cfg.get('extras_anotacao', {}).items():
                x = valor_campo(v.get(col), 'data_texto')
                if x is not None:
                    notas.append(f'[{rot}] {_dt(x).strftime("%d/%m/%Y") if isinstance(_dt(x), date) else x}')
            if notas:
                campos['anotacao'] = ' · '.join(notas)
            original = {k2: (v2.isoformat() if isinstance(v2, (datetime, date)) else v2)
                        for k2, v2 in v.items() if v2 not in (None, '')}
            regs.append({'chave': chave, 'carteira': cid, 'ctrc': ctrc, 'cte': cte, 'etapa': etapa,
                         'campos': campos, 'planilha': dict(base, original=original)})
    return regs, rel


def orfas_da_planilha(cur, regs, ssw_chaves):
    """Linhas que vieram de uma importação anterior e não existem mais na planilha: a chave da linha
    manual leva o nº da linha da planilha, e ela muda quando se insere linha ou se preenche o CTRC
    (21/10/2026: 9 duplicadas — a Nestlé 3744 virou RIO004579-9 e a antiga, sem CTRC, ficou). Só
    linha manual (MAN-… ou CTRC que o 455 não conhece) e nunca editada na tela; CT-e do SSW fica.
    """
    carts = sorted({g['carteira'] for g in regs})
    cur.execute("""SELECT m.chave FROM crm_doc_marcacao m
                    WHERE m.origem = 'planilha' AND m.carteira = ANY(%s) AND NOT (m.chave = ANY(%s))
                      AND NOT EXISTS (SELECT 1 FROM crm_doc_log l WHERE l.chave = m.chave
                                         AND l.campo <> '*importacao*' AND NOT l.auto)""",
                (carts, [g['chave'] for g in regs]))
    return sorted(k for (k,) in cur.fetchall() if k.startswith('MAN-') or k not in ssw_chaves)


def gravar_importacao(cur, regs, por='Importação da planilha', ssw_chaves=None):
    """Grava as marcações da planilha. Linha já editada na tela não é tocada (o que ela mudou
    na tela vale mais que a planilha antiga); o FATURADO automático não conta como edição —
    a planilha entra e a tela volta a pôr FATURADO na próxima leitura, se o SSW provar. Com
    `ssw_chaves` (os CTRCs do 455), apaga as órfãs da importação anterior (`orfas_da_planilha`).
    Devolve (novas, atualizadas, preservadas, removidas)."""
    novas = atual = preserv = 0
    for g in regs:
        cur.execute("""SELECT EXISTS (SELECT 1 FROM crm_doc_log WHERE chave = %s AND campo <> '*importacao*'
                                         AND NOT auto),
                              EXISTS (SELECT 1 FROM crm_doc_marcacao WHERE chave = %s)""", (g['chave'], g['chave']))
        editada, existe = cur.fetchone()
        if editada:
            preserv += 1
            continue
        cur.execute("""INSERT INTO crm_doc_marcacao (chave, carteira, ctrc, cte, etapa, campos, planilha, origem, atualizado_por)
                       VALUES (%s, %s, %s, %s, %s, %s::jsonb, %s::jsonb, 'planilha', %s)
                       ON CONFLICT (chave) DO UPDATE SET carteira = EXCLUDED.carteira, ctrc = EXCLUDED.ctrc,
                           cte = EXCLUDED.cte, etapa = EXCLUDED.etapa, campos = EXCLUDED.campos,
                           planilha = EXCLUDED.planilha, origem = 'planilha',
                           atualizado_em = now(), atualizado_por = EXCLUDED.atualizado_por""",
                    (g['chave'], g['carteira'], g['ctrc'], g['cte'], g['etapa'],
                     json.dumps(g['campos'], ensure_ascii=False),
                     json.dumps(g['planilha'], ensure_ascii=False, default=str), por))
        cur.execute('INSERT INTO crm_doc_log (chave, campo, antes, depois, por) VALUES (%s, %s, %s, %s, %s)',
                    (g['chave'], '*importacao*', None,
                     json.dumps({'etapa': g['etapa'], **g['campos']}, ensure_ascii=False)[:2000], por))
        if existe:
            atual += 1
        else:
            novas += 1
    removidas = orfas_da_planilha(cur, regs, ssw_chaves) if ssw_chaves is not None else []
    if removidas:
        cur.execute('DELETE FROM crm_doc_log WHERE chave = ANY(%s)', (removidas,))
        cur.execute('DELETE FROM crm_doc_marcacao WHERE chave = ANY(%s)', (removidas,))
    return novas, atual, preserv, removidas


# ── Excel no layout da planilha ──────────────────────────────────────────────

def _dt(iso):
    if not iso:
        return None
    try:
        return date.fromisoformat(iso)
    except ValueError:
        return iso


def exportar_xlsx(carteira, linhas):
    """As linhas da carteira no layout da aba da planilha delas, mais o que o SSW diz no fim."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    cfg = CARTEIRAS[carteira]
    cols = [('Emissao', lambda l: _dt(l['emissao'])), ('Mês', lambda l: _dt(l['emissao']).month if l['emissao'] else None),
            ('Ano', lambda l: _dt(l['emissao']).year if l['emissao'] else None),
            ('CTRC', lambda l: l['ctrc']), (cfg['col_cte'], lambda l: l['cte'] or l.get('nfs')),
            ('CNPJ Pagador', lambda l: l['cnpj']), ('Cliente Pagador', lambda l: l['pagador']),
            ('UF O', lambda l: l['uf_o']), ('NF', lambda l: l['nf']), ('Frete', lambda l: l['frete'])]
    if cfg.get('col_transporte'):
        cols.append(('TRANSPORTE', lambda l: l.get('transporte')))
    for c in cfg['campos']:
        if c['id'] == 'tipo':
            cols.append(('TIPO FRETE', lambda l: l['campos'].get('tipo')))
    cols.append(('Observacao', lambda l: l['obs']))
    cols.append((cfg['col_etapa'], lambda l: l['etapa']))
    for c in cfg['campos']:
        if c['id'] == 'tipo' or (c['id'] == 'vencimento' and cfg.get('prazo_dias')) or c['tipo'] == 'custos':
            continue
        cols.append((c['col'] or c['rotulo'],
                     (lambda cid, tp: (lambda l: _dt(l['campos'].get(cid)) if tp == 'data_texto' else l['campos'].get(cid)))(c['id'], c['tipo'])))
    if cfg.get('prazo_dias'):
        cols.append(('VENCIMENTO', lambda l: _dt(l['vencimento'])))
    rotulo_ssw = {'pago': 'PAGO', 'fatura': 'FATURA EMITIDA', 'aberto': 'SEM FATURA', 'sem_ssw': 'SEM CT-E NO SSW',
                  'cancelado': 'CANCELADO NO SSW',
                  'indefinido': 'SEM INFORMAÇÃO'}
    fa = lambda l, k: (l.get('fatura_aberta') or {}).get(k)
    cols += [('SSW', lambda l: ('FATURA VENCIDA' if l.get('vencida') else
                                'PAGO (PARCIAL DA FATURA)' if l.get('fonte') == '441 parcial' else
                                rotulo_ssw.get(l['ssw'], l['ssw']))),
             ('PAGO EM', lambda l: _dt(l.get('pago_em'))),
             ('FATURA', lambda l: fa(l, 'fatura') or ', '.join(l.get('faturas') or [])),
             ('VENC. FATURA', lambda l: _dt(l.get('venc_fatura') or fa(l, 'vencimento'))),
             ('SALDO FATURA', lambda l: fa(l, 'saldo')),
             ('DIVERGÊNCIA', lambda l: ' | '.join(d['texto'] for d in l['diverg']))]
    wb = Workbook()
    ws = wb.active
    ws.title = cfg['aba'][:31]
    azul = PatternFill('solid', fgColor='1F4E78')
    for j, (cab, _) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=cab)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = azul
    for i, l in enumerate(linhas, start=2):
        for j, (cab, f) in enumerate(cols, start=1):
            v = f(l)
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, date):
                c.number_format = 'DD/MM/YYYY'
            elif cab == 'Frete':
                c.number_format = '#,##0.00'
    for j, (cab, _) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = {'Observacao': 50, 'DIVERGÊNCIA': 45}.get(cab, 15)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(1, len(linhas) + 1)}'
    if any(c['tipo'] == 'custos' for c in cfg['campos']):
        # Custo extra: uma linha por custo, nas colunas da planilha dela, com o CT-e da carga na frente
        wc = wb.create_sheet('Custo extra')
        cab = ['CTRC', 'CT-e', 'Emissao'] + [r for _, r, _ in ITENS_CUSTO]
        for j, t in enumerate(cab, start=1):
            c = wc.cell(row=1, column=j, value=t)
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = azul
        i = 1
        for l in linhas:
            for x in (l['campos'].get('custos') or []):
                i += 1
                vals = [l['ctrc'], l['cte'] or l.get('nfs'), _dt(l['emissao'])] + [x.get(cid) for cid, _, _ in ITENS_CUSTO]
                for j, v in enumerate(vals, start=1):
                    c = wc.cell(row=i, column=j, value=v)
                    if isinstance(v, date):
                        c.number_format = 'DD/MM/YYYY'
                    elif cab[j - 1] == 'Valor':
                        c.number_format = '#,##0.00'
        for j in range(1, len(cab) + 1):
            wc.column_dimensions[get_column_letter(j)].width = 16
        wc.freeze_panes = 'A2'
        wc.auto_filter.ref = f'A1:{get_column_letter(len(cab))}{max(1, i)}'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
